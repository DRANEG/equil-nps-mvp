# -*- coding: utf-8 -*-
"""
Test automat: construieste o versiune redusa a workbook-ului, o populeaza cu date
de test si verifica valorile calculate de Excel fata de valorile calculate in Python.

Ruleaza: python3 tools/test_workbook.py
"""
import os
import subprocess
import sys
import datetime as dt

SMALL = dict(N_CLIENTI="8", N_PRODUSE="4", N_OAMENI="3", N_VANZARI="40",
             N_TINTE="8", N_NPS="10", N_OPP="6", N_CONTACTE="5",
             OUT="/tmp/equil_test.xlsx")

env = dict(os.environ, **SMALL)
subprocess.run([sys.executable, "tools/build_equil_workbook.py"], env=env, check=True)

import openpyxl
path = SMALL["OUT"]
wb = openpyxl.load_workbook(path)

CUR = (dt.date(2026, 1, 1), dt.date(2026, 6, 30))
CMP = (dt.date(2025, 1, 1), dt.date(2025, 6, 30))

clienti = [
    ("C001", "Alfa SRL", "Key Account", "Wholesale", "București", "RO", "S001", dt.date(2021, 3, 15), "Activ", 30, 50000, ""),
    ("C002", "Beta SA", "Mid-market", "Retail", "Cluj", "RO", "S002", dt.date(2023, 6, 1), "Activ", 15, 20000, ""),
    ("C003", "Gama SRL", "Small", "Online", "Iași", "RO", "S002", dt.date(2024, 1, 10), "Activ", 15, 10000, ""),
    ("C004", "Delta SRL", "Small", "Retail", "Timiș", "RO", "S001", dt.date(2022, 5, 5), "Activ", 30, 10000, ""),
]
produse = [
    ("P001", "Produs X", "Categoria A", "", 120, 84, "Nu", None, "Activ"),
    ("P002", "Serviciu Y", "Servicii", "", 1200, 600, "Da", None, "Activ"),
]
oameni = [
    ("S001", "Ana P.", "KAM", "B2B", 1, dt.date(2023, 2, 1), None, 96000, "Da"),
    ("S002", "Bogdan M.", "Sales", "B2B", 0.5, dt.date(2024, 9, 1), None, 72000, "Da"),
]
# (data, tip, client, produs, persoana, unitati, venit, discount, cogs)
vanzari = [
    (dt.date(2025, 2, 28), "Factură", "C001", "P001", "S001", 100, 10000, 0, 7000),
    (dt.date(2025, 3, 31), "Factură", "C002", "P002", "S002", 10, 12000, 0, 6000),
    (dt.date(2025, 4, 30), "Factură", "C003", "P001", "S002", 20, 2000, 0, 1500),
    (dt.date(2025, 5, 31), "Factură", "C004", "P001", "S001", 30, 3000, 0, 2600),
    (dt.date(2026, 2, 28), "Factură", "C001", "P001", "S001", 60, 6000, 200, 4500),
    (dt.date(2026, 3, 31), "Factură", "C002", "P002", "S002", 12, 14400, 0, 7200),
    (dt.date(2026, 3, 31), "Factură", "C002", "P001", "S002", 5, 500, 0, 350),
    (dt.date(2026, 4, 30), "Factură", "C004", "P001", "S001", 35, 3500, 100, 3100),
]
tinte = [(dt.date(2026, 3, 31), "Companie", "", "Venit", 15000, ""),
         (dt.date(2026, 6, 30), "Companie", "", "Venit", 15000, ""),
         (dt.date(2026, 6, 30), "Persoană", "S001", "Venit", 12000, "")]
nps = [("R1", dt.date(2026, 2, 10), "C001", "Mihai I.", "Decident", 9, None, "ok", "Relație", "Email", "T1", "S001", "Neînceput"),
       ("R2", dt.date(2026, 2, 12), "C002", "Ioana R.", "User", 5, None, "întârzieri", "Livrare", "Email", "T1", "S002", "În lucru"),
       ("R3", dt.date(2026, 3, 2), "C003", "Radu T.", "User", 8, None, "ok", "Produs", "Email", "T1", "S002", "Neînceput"),
       ("R4", dt.date(2026, 3, 5), "C004", "Elena V.", "Decident", 10, None, "excelent", "Relație", "Email", "T1", "S001", "Neînceput")]


def fill(sheet, rows, first=5, cols_map=None):
    ws = wb[sheet]
    for i, row in enumerate(rows):
        for j, v in enumerate(row, start=1):
            col = cols_map[j - 1] if cols_map else j
            if col and v is not None:
                ws.cell(row=first + i, column=col, value=v)


for s in ("10_CLIENTI", "11_PRODUSE", "12_OAMENI", "13_VANZARI", "14_TINTE", "15_NPS"):
    ws = wb[s]
    for r in range(5, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and not v.startswith("="):
                ws.cell(row=r, column=c, value=None)

fill("10_CLIENTI", clienti)
fill("11_PRODUSE", produse)
fill("12_OAMENI", oameni)
# 13_VANZARI: A data, C tip, D client, E produs, F persoana, G unitati, H venit, I discount, J cogs
fill("13_VANZARI", vanzari, cols_map=[1, 3, 4, 5, 6, 7, 8, 9, 10])
fill("14_TINTE", tinte)
fill("15_NPS", nps)
wb.save(path)

# ------------------------------------------------ valori asteptate (Python)
cur = [v for v in vanzari if CUR[0] <= v[0] <= CUR[1]]
cmpp = [v for v in vanzari if CMP[0] <= v[0] <= CMP[1]]
rev = lambda rows: sum(r[6] for r in rows)
gp = lambda rows: sum(r[6] - r[8] for r in rows)
exp = {
    "Venit net": (rev(cur), rev(cmpp)),
    "Profit brut": (gp(cur), gp(cmpp)),
    "Marjă brută %": (gp(cur) / rev(cur), gp(cmpp) / rev(cmpp)),
    "Unități vândute": (sum(r[5] for r in cur), sum(r[5] for r in cmpp)),
    "Discount acordat": (sum(r[7] for r in cur), sum(r[7] for r in cmpp)),
    "Clienți activi": (len({r[2] for r in cur}), len({r[2] for r in cmpp})),
    "Clienți noi": (0, 4),
    "Număr de linii / facturi": (len(cur), len(cmpp)),
    "FTE activ": (1.5, 1.5),  # ambii angajați erau activi în ambele perioade
    "Țintă de venit": (30000, 0),
    "NPS": ((2 - 1) / 4 * 100, ""),
    "% Promotori": (0.5, ""),
    "% Detractori": (0.25, ""),
}

# Motorul de evaluare 'formulas' nu accepta paranteze drepte in literalii de text
# (limitare a bibliotecii, nu a Excel-ului). Doar pentru test le inlocuim.
wb2 = openpyxl.load_workbook(path)
for _ws in wb2.worksheets:
    for _row in _ws.iter_rows():
        for _c in _row:
            if isinstance(_c.value, str) and _c.value.startswith("="):
                _c.value = _c.value.replace("[", "(").replace("]", ")")
    _ws.data_validations.dataValidation = []
    _ws.conditional_formatting = type(_ws.conditional_formatting)()
ascii_path = path.replace(".xlsx", "_ascii.xlsx")
wb2.save(ascii_path)
path = ascii_path

print("\n== Se calculeaza workbook-ul cu motorul 'formulas' (poate dura) ==")
import formulas
xl = formulas.ExcelModel().loads(path).finish()
sol = xl.calculate()


BOOK = os.path.basename(path)


def val(sheet, cell):
    v = sol[f"'[{BOOK}]{sheet}'!{cell}"].value[0, 0]
    return v


kpi = wb["20_KPI"]
rows = {kpi.cell(row=r, column=1).value: r for r in range(5, kpi.max_row + 1)}
ok = bad = 0
print("\n%-34s %14s %14s %14s %14s" % ("KPI", "Excel curent", "Aștept curent", "Excel comp.", "Aștept comp."))
for name, (ec, ep) in exp.items():
    r = rows[name]
    gc, gpv = val("20_KPI", f"B{r}"), val("20_KPI", f"C{r}")
    def cmp_(a, b):
        if isinstance(b, str):
            return True
        try:
            return abs(float(a) - float(b)) < 0.01
        except Exception:
            return False
    good = cmp_(gc, ec) and cmp_(gpv, ep)
    ok, bad = (ok + 1, bad) if good else (ok, bad + 1)
    print("%-34s %14s %14s %14s %14s  %s" % (name, gc, ec, gpv, ep, "OK" if good else "<<< DIFERIT"))

print("\n== Semnale pe clienți (21_ANALIZA_CLIENTI) ==")
for i in range(4):
    r = 5 + i
    print("  %-10s venit=%-9s Δ=%-9s semnal=%-22s oport.=%-22s potențial=%s" % (
        val("21_ANALIZA_CLIENTI", f"A{r}"), val("21_ANALIZA_CLIENTI", f"E{r}"),
        val("21_ANALIZA_CLIENTI", f"G{r}"), val("21_ANALIZA_CLIENTI", f"O{r}"),
        val("21_ANALIZA_CLIENTI", f"P{r}"), val("21_ANALIZA_CLIENTI", f"Q{r}")))

print("\n== Semnale pe produse (22) ==")
for i in range(2):
    r = 5 + i
    print("  %-8s venit=%-9s marjă=%-8s semnal=%-28s oport.=%s" % (
        val("22_ANALIZA_PRODUSE", f"A{r}"), val("22_ANALIZA_PRODUSE", f"D{r}"),
        val("22_ANALIZA_PRODUSE", f"H{r}"), val("22_ANALIZA_PRODUSE", f"M{r}"),
        val("22_ANALIZA_PRODUSE", f"N{r}")))

print("\n== Semnale pe oameni (23) ==")
for i in range(2):
    r = 5 + i
    print("  %-8s venit=%-9s țintă=%-9s atingere=%-10s semnal=%s" % (
        val("23_ANALIZA_OAMENI", f"A{r}"), val("23_ANALIZA_OAMENI", f"F{r}"),
        val("23_ANALIZA_OAMENI", f"M{r}"), val("23_ANALIZA_OAMENI", f"N{r}"),
        val("23_ANALIZA_OAMENI", f"O{r}")))

rap = wb["40_RAPORT"]
print("\n== Concluzii automate (40_RAPORT) ==")
concl_start = next(r for r in range(1, rap.max_row + 1)
                   if str(rap.cell(row=r, column=1).value or "").startswith("CONCLUZII"))
for r in range(concl_start + 1, concl_start + 9):
    lab = rap.cell(row=r, column=1).value
    print(f"  {lab:16s} {val('40_RAPORT', f'B{r}')}")

print("\n== Zona de calcul a graficelor (41_GRAFICE) ==")
print("  ultimele 3 luni:")
for rr in (70, 71, 72):
    print("    %s venit=%-10s profit=%-9s țintă=%-9s marjă=%s" % (
        val("41_GRAFICE", f"B{rr}"), val("41_GRAFICE", f"C{rr}"),
        val("41_GRAFICE", f"D{rr}"), val("41_GRAFICE", f"E{rr}"),
        val("41_GRAFICE", f"F{rr}")))
print("  top clienți:", [(val("41_GRAFICE", f"H{61+i}"), val("41_GRAFICE", f"I{61+i}")) for i in range(4)])
print("  top produse:", [(val("41_GRAFICE", f"K{61+i}"), val("41_GRAFICE", f"L{61+i}")) for i in range(3)])
print("  NPS:", [(val("41_GRAFICE", f"N{61+i}"), val("41_GRAFICE", f"O{61+i}")) for i in range(3)])
print("  oportunități:", [(val("41_GRAFICE", f"Q{61+i}"), val("41_GRAFICE", f"R{61+i}")) for i in range(3)])

print("\n== Raport (40_RAPORT) — top oportunități ==")
start = next(r for r in range(1, rap.max_row + 1)
             if str(rap.cell(row=r, column=1).value or "").startswith("TOP 5"))
for r in range(start + 2, start + 7):
    print("  %s | %-34s | scor=%s" % (val("40_RAPORT", f"A{r}"),
                                      val("40_RAPORT", f"B{r}"),
                                      val("40_RAPORT", f"C{r}")))

print(f"\nRezultat: {ok} KPI corecți, {bad} diferiți.")
sys.exit(1 if bad else 0)
