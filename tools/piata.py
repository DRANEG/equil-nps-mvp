# -*- coding: utf-8 -*-
"""Modulul de piata si perceptie: concurenta, recenzii publice, parerea specialistului.

Ideea de baza: multe firme nu au NPS. Dar au o nota pe Google, au concurenti
cu note vizibile, si exista cineva (consultantul) care poate da o parere
structurata. Astea trei se pot aduna intr-o dupa-amiaza, fara sa intrebi
niciun client, si tin locul NPS-ului pana cand firma incepe sa masoare.

Scorul de perceptie NU este NPS si nu se amesteca cu el: sunt afisate separat,
iar foaia arata explicit la ce nivel e firma si care e urmatorul pas.

Folosit de build_equil_workbook.py si de build_equil_light.py.
"""
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from stil import (INK, TEAL, TEAL_LIGHT, INPUT_FILL, CALC_FILL, WHITE, GREY_TXT,
                  F_MONEY, F_PCT, F_DATE, F_INT, F_NUM, BORDER,
                  style_header, title_block, add_dv, mon, pct)

HR = 4
R0 = HR + 1

# (titlu, latime, format, calculat, si_in_light, ajutor)
COL_CONCURENTA = [
    ("Cine", 26, None, False, True,
     "Primul rând este firma analizată („NOI”). Sub el, concurenții — cei la care clientul "
     "chiar pierde vânzări, nu toți din județ."),
    ("Site", 22, None, False, True, "Adresa site-ului, ca să poți reveni la el."),
    ("Zonă / oraș", 15, None, False, False, "Unde activează. Contează dacă piața e locală."),
    ("Angajați", 10, F_INT, False, True,
     "Aproximativ. Se vede pe LinkedIn, pe site sau din datele publice."),
    ("Cifră de afaceri", 15, F_MONEY, False, True,
     "Ultima cifră publică. În România se găsește gratuit în datele publice ale firmelor. "
     "Pentru „NOI”, aceeași cifră ca în profilul firmei."),
    ("Anul cifrei", 10, F_INT, False, False, "Din ce an e cifra. Una din 2021 nu se compară cu una din 2025."),
    ("De unde e cifra", 17, None, False, True,
     "Ca să știi cât te poți baza pe ea: date publice, estimare proprie, sau ce ți-a spus clientul."),
    ("Notă publică (1-5)", 12, F_NUM, False, True,
     "Nota medie de pe Google, Facebook sau altă platformă. Scrie-o cu zecimale: 4,3 nu 4."),
    ("Câte recenzii", 11, F_INT, False, True,
     "Numărul total. O notă de 5,0 din 3 recenzii nu înseamnă nimic; 4,3 din 200 înseamnă mult."),
    ("Din care 1-2 stele", 12, F_INT, False, False,
     "Recenziile proaste. De obicei spun mai multe decât media."),
    ("Verificat la data", 13, F_DATE, False, True,
     "Când ai citit tu notele. Se schimbă în timp, iar peste șase luni vrei să știi cât de vechi e."),
    ("Preț față de noi", 14, None, False, True, "Cum se poziționează la preț."),
    ("Cu ce câștigă", 30, None, False, True,
     "Un lucru concret pe care îl face mai bine. Dacă nu găsești niciunul, n-ai căutat destul."),
    ("Unde pierde", 30, None, False, True,
     "Ce reclamă clienții lui. Citește-i recenziile de 1-2 stele — acolo scrie."),
    ("Cât se suprapune", 13, None, False, False,
     "Cât de des vă bateți pe același client: direct, parțial sau rar."),
    ("Cotă în grup", 11, F_PCT, True, True,
     "Cât reprezintă din cifra totală a firmelor din acest tabel. Nu e cotă de piață reală — "
     "e cota în grupul pe care îl urmărești."),
    ("% recenzii proaste", 12, F_PCT, True, False, "Calculat: recenziile de 1-2 stele din total."),
    ("Cum stăm față de el", 26, None, True, True, "Calculat automat din notă și din cifra de afaceri."),
    ("Observații", 22, None, False, False, "Orice altceva util."),
]

SURSE_CIFRA = ["Date publice / bilanț", "Estimare proprie", "Spus de client", "Site / presă", "Nu știu"]
PRET_FATA = ["Mult mai ieftin", "Mai ieftin", "La fel", "Mai scump", "Mult mai scump", "Nu știu"]
SUPRAPUNERE = ["Direct", "Parțial", "Rar"]
SURSE_RECENZII = ["Google", "Facebook", "Site propriu", "Trustpilot", "Booking", "eMAG",
                  "Recomandări directe", "Altă platformă"]

DIMENSIUNI = [
    ("Se înțelege în 10 secunde ce vinde?",
     "Intră pe site ca un client care nu știe nimic. Dacă nu pricepi în 10 secunde, nici el nu pricepe."),
    ("Cât de ușor e de găsit online?",
     "Caută pe Google ce vinde + orașul. Apare? Pe ce poziție? Are fișă Google Business completă?"),
    ("Cât de repede răspunde la o cerere?",
     "Trimite o cerere reală prin formular sau sună. Cronometrează. Sub o oră e bine, peste o zi e o problemă."),
    ("Prima impresie la contact",
     "Cum răspunde la telefon, cum arată magazinul sau biroul, cum arată oferta scrisă."),
    ("Prețul, față de ce ai văzut la concurenți",
     "Nu „scump/ieftin”, ci: prețul e justificat de ceva vizibil?"),
    ("Livrează consecvent ce promite?",
     "Din recenzii și din ce spun clienții: se întâmplă la fel de fiecare dată?"),
    ("Cât de ușor e să reclami ceva?",
     "Există un om, un număr, un proces? Sau reclamația se pierde?"),
    ("Cu ce se diferențiază real?",
     "Un lucru pe care concurenții nu îl au. „Calitate” și „seriozitate” nu sunt răspunsuri."),
    ("Ce se aude despre ei în piață?",
     "Ce spun furnizorii, partenerii, foștii angajați. Informal, dar util."),
    ("Tu, ca specialist, ai recomanda-o unui prieten?",
     "Aceeași întrebare ca la NPS, pusă ție. 1 = în niciun caz, 5 = fără ezitare."),
]


def _antet_sectiune(ws, r, text, ncol):
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(bold=True, size=11, color=WHITE)
    for k in range(1, ncol + 1):
        ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=TEAL)
    ws.row_dimensions[r].height = 20


def _lista_ascunsa(ws, col_start, liste):
    """Scrie listele de validare intr-o zona laterala si returneaza referintele."""
    refs = {}
    for i, (nume, valori) in enumerate(liste.items()):
        j = col_start + i
        L = get_column_letter(j)
        ws.column_dimensions[L].hidden = True
        ws.cell(row=1, column=j, value=nume).font = Font(size=8, color=GREY_TXT)
        for k, v in enumerate(valori, start=2):
            ws.cell(row=k, column=j, value=v).font = Font(size=9)
        refs[nume] = f"'{ws.title}'!${L}$2:${L}${len(valori) + 1}"
    return refs


def adauga_concurenta(wb, nume_foaie, n_randuri, light=False):
    """Tabelul de concurenta. Primul rand de date este intotdeauna firma analizata."""
    coloane = [c for c in COL_CONCURENTA if c[4] or not light]
    ws = wb.create_sheet(nume_foaie)
    numar = nume_foaie.split("_")[0]
    title_block(ws, f"{numar} — CONCURENȚA",
                "Primul rând ești tu (firma analizată). Sub el, concurenții. "
                "Cifrele publice și notele se adună de pe internet într-o oră, fără să întrebi pe nimeni.")
    for j, (t, w, fmt, _calc, _l, ajutor) in enumerate(coloane, start=1):
        c = ws.cell(row=HR, column=j, value=t)
        ws.column_dimensions[get_column_letter(j)].width = w
        if ajutor:
            from openpyxl.comments import Comment
            c.comment = Comment(f"{t}\n\n{ajutor}", "EQUIL")
    style_header(ws, HR, len(coloane), height=36)
    last = HR + n_randuri
    for j, (t, w, fmt, calc, _l, _a) in enumerate(coloane, start=1):
        for rr in range(R0, last + 1):
            cell = ws.cell(row=rr, column=j)
            cell.fill = PatternFill("solid", fgColor=CALC_FILL if calc else INPUT_FILL)
            cell.border = BORDER
            cell.font = Font(size=10, color=GREY_TXT if calc else "000000")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = f"B{R0}"
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(coloane))}{last}"
    ws.sheet_view.showGridLines = False

    col = {t: get_column_letter(j) for j, (t, *_r) in enumerate(coloane, start=1)}
    CINE, CA, NOTA = col["Cine"], col["Cifră de afaceri"], col["Notă publică (1-5)"]
    REC = col["Câte recenzii"]

    R_CINE = f"'{nume_foaie}'!${CINE}${R0}:${CINE}${last}"
    R_CA = f"'{nume_foaie}'!${CA}${R0}:${CA}${last}"
    R_NOTA = f"'{nume_foaie}'!${NOTA}${R0}:${NOTA}${last}"
    R_REC = f"'{nume_foaie}'!${REC}${R0}:${REC}${last}"
    NOI_NOTA = f"'{nume_foaie}'!${NOTA}${R0}"
    NOI_CA = f"'{nume_foaie}'!${CA}${R0}"

    ws.cell(row=R0, column=1, value="NOI (firma analizată)").font = Font(size=10, bold=True, color=INK)
    ws.cell(row=R0, column=1).fill = PatternFill("solid", fgColor=TEAL_LIGHT)

    # coloane calculate
    for rr in range(R0, last + 1):
        ws[f"{col['Cotă în grup']}{rr}"] = (
            f'=IF(${CINE}{rr}="","",IFERROR(${CA}{rr}/SUM({R_CA}),""))')
        if not light:
            NEG = col["Din care 1-2 stele"]
            ws[f"{col['% recenzii proaste']}{rr}"] = (
                f'=IF(OR(${CINE}{rr}="",${REC}{rr}=""),"",IFERROR(${NEG}{rr}/${REC}{rr},""))')
        ws[f"{col['Cum stăm față de el']}{rr}"] = (
            f'=IF(${CINE}{rr}="","",IF(ROW()={R0},"— reper —",'
            f'IF(AND(${NOTA}{rr}<>"",{NOI_NOTA}<>""),'
            f'IF(${NOTA}{rr}>{NOI_NOTA}+0.2,"E mai bine văzut",'
            f'IF(${NOTA}{rr}<{NOI_NOTA}-0.2,"Suntem mai bine văzuți","Văzuți la fel")),"Fără notă")'
            f'&IF(AND(${CA}{rr}<>"",{NOI_CA}<>"",{NOI_CA}>0),'
            f'IF(${CA}{rr}>{NOI_CA}*1.2," · mai mare",'
            f'IF(${CA}{rr}<{NOI_CA}*0.8," · mai mic"," · cam cât noi")),"")))')

    for semnal, culoare in (("E mai bine văzut", "F8D7DA"), ("Suntem mai bine văzuți", "D6F0D6")):
        ws.conditional_formatting.add(
            f"{col['Cum stăm față de el']}{R0}:{col['Cum stăm față de el']}{last}",
            CellIsRule(operator="containsText", formula=[f'NOT(ISERROR(SEARCH("{semnal}",'
                                                        f'{col["Cum stăm față de el"]}{R0})))'],
                       fill=PatternFill("solid", fgColor=culoare)))

    liste = {"surse_cifra": SURSE_CIFRA, "pret": PRET_FATA}
    if not light:
        liste["suprapunere"] = SUPRAPUNERE
    refs_liste = _lista_ascunsa(ws, len(coloane) + 12, liste)
    add_dv(ws, refs_liste["surse_cifra"], f"{col['De unde e cifra']}{R0}:{col['De unde e cifra']}{last}")
    add_dv(ws, refs_liste["pret"], f"{col['Preț față de noi']}{R0}:{col['Preț față de noi']}{last}")
    if not light:
        add_dv(ws, refs_liste["suprapunere"],
               f"{col['Cât se suprapune']}{R0}:{col['Cât se suprapune']}{last}")

    # panou de sinteza, in dreapta tabelului
    sc = len(coloane) + 2
    SL, SV = get_column_letter(sc), get_column_letter(sc + 1)
    ws.column_dimensions[SL].width = 30
    ws.column_dimensions[SV].width = 14
    ws.cell(row=HR, column=sc, value="CUM ARATĂ PIAȚA").font = Font(bold=True, size=10, color=WHITE)
    for k in (sc, sc + 1):
        ws.cell(row=HR, column=k).fill = PatternFill("solid", fgColor=INK)

    NR_CONC = f'COUNTA({R_CINE})-IF({R_CINE.replace(f"${R0}", f"${R0}").split(":")[0]}="",0,1)'
    NR_CONC = f'MAX(0,COUNTA({R_CINE})-1)'
    CA_GRUP = f'SUM({R_CA})'
    COTA = f'IFERROR({NOI_CA}/{CA_GRUP},"")'
    NOTA_CONC = (f'IFERROR(AVERAGE(\'{nume_foaie}\'!${NOTA}${R0 + 1}:${NOTA}${last}),"")')
    REC_TOTAL = f'SUM({R_REC})'

    sinteza = [
        ("Concurenți urmăriți", f"={NR_CONC}", F_INT),
        ("Cifra totală a grupului", f"={CA_GRUP}", F_MONEY),
        ("Cota noastră în grup", f"={COTA}", F_PCT),
        ("Piața totală estimată", 0, F_MONEY),
        ("Cota noastră din piață", None, F_PCT),
        ("Nota noastră publică", f"=IF({NOI_NOTA}=\"\",\"\",{NOI_NOTA})", F_NUM),
        ("Nota medie a concurenților", f"={NOTA_CONC}", F_NUM),
        ("Diferența", None, F_NUM),
    ]
    for k, (eticheta, val, fmt) in enumerate(sinteza, start=1):
        rr = HR + k
        ws.cell(row=rr, column=sc, value=eticheta).font = Font(size=10, bold=True, color=INK)
        c = ws.cell(row=rr, column=sc + 1, value=val)
        c.number_format = fmt
        c.border = BORDER
        editabil = eticheta == "Piața totală estimată"
        c.fill = PatternFill("solid", fgColor=INPUT_FILL if editabil else CALC_FILL)
        c.font = Font(size=10, color="000000" if editabil else GREY_TXT)
    PIATA = f"'{nume_foaie}'!${SV}${HR + 4}"
    ws[f"{SV}{HR + 5}"] = f'=IF(OR({PIATA}=0,{PIATA}=""),"",IFERROR({NOI_CA}/{PIATA},""))'
    ws[f"{SV}{HR + 8}"] = (f'=IF(OR({NOI_NOTA}="",{SV}{HR + 7}=""),"",{NOI_NOTA}-{SV}{HR + 7})')
    ws.cell(row=HR + 4, column=sc + 2,
            value="Opțional. Dacă știi cât valorează toată piața, aici iese cota reală.").font = \
        Font(size=9, italic=True, color=GREY_TXT)
    ws.column_dimensions[get_column_letter(sc + 2)].width = 44

    return dict(foaie=nume_foaie, prim_rand=R0, ultim_rand=last, col=col,
                nr_concurenti=NR_CONC, cota_grup=COTA, nota_noastra=NOI_NOTA,
                nota_concurenti=NOTA_CONC, recenzii_total_grup=REC_TOTAL,
                ref_nume=R_CINE, ref_nota=R_NOTA, ref_ca=R_CA)


def adauga_perceptie(wb, nume_foaie, conc, nps_scor, nps_nr, light=False):
    """Recenzii publice + parerea specialistului + drumul catre NPS.

    nps_scor / nps_nr: formule (fara "=") care dau NPS-ul real si numarul de
    clienti intrebati, ca sa se vada alaturi, nu amestecat.
    """
    ws = wb.create_sheet(nume_foaie)
    numar = nume_foaie.split("_")[0]
    title_block(ws, f"{numar} — CE CRED OAMENII DESPRE FIRMĂ",
                "Trei surse, în ordinea efortului: ce se vede public, ce vezi tu ca specialist, "
                "ce spun clienții întrebați direct. Ultima e NPS-ul și e singura care contează cu adevărat; "
                "primele două țin locul până atunci.")
    latimi = [30, 26, 13, 11, 12, 13, 14, 34, 34, 22]
    for j, w in enumerate(latimi, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.sheet_view.showGridLines = False

    n_surse = 6 if light else 10
    r = HR
    _antet_sectiune(ws, r, "1. CE SE VEDE PUBLIC — recenzii, note, platforme", 10)
    r += 1
    cap_recenzii = ["Sursa", "Link", "Verificat la", "Notă (1-5)", "Câte recenzii",
                    "Din care 1-2 stele", "Cea mai recentă", "Ce laudă cel mai des",
                    "Ce reclamă cel mai des", "Observații"]
    for j, t in enumerate(cap_recenzii, start=1):
        ws.cell(row=r, column=j, value=t)
    style_header(ws, r, 10, height=28)
    rec_head = r
    rec_first = r + 1
    rec_last = r + n_surse
    formate = {3: F_DATE, 4: F_NUM, 5: F_INT, 6: F_INT, 7: F_DATE}
    for rr in range(rec_first, rec_last + 1):
        for j in range(1, 11):
            c = ws.cell(row=rr, column=j)
            c.fill = PatternFill("solid", fgColor=INPUT_FILL)
            c.border = BORDER
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j in formate:
                c.number_format = formate[j]
    ws.cell(row=rec_first, column=1, value="Google")

    liste = _lista_ascunsa(ws, 26, {"surse": SURSE_RECENZII})
    add_dv(ws, liste["surse"], f"A{rec_first}:A{rec_last}")

    R_NOTA = f"'{nume_foaie}'!$D${rec_first}:$D${rec_last}"
    R_NR = f"'{nume_foaie}'!$E${rec_first}:$E${rec_last}"
    NOTA_POND = (f'IFERROR(SUMPRODUCT(({R_NOTA}<>"")*({R_NR}<>"")*{R_NOTA}*{R_NR})/'
                 f'SUMPRODUCT(({R_NOTA}<>"")*({R_NR}<>"")*{R_NR}),"")')
    TOTAL_REC = f'SUMPRODUCT(({R_NOTA}<>"")*{R_NR})'

    r = rec_last + 2
    _antet_sectiune(ws, r, "2. PĂREREA SPECIALISTULUI — zece lucruri pe care le poți verifica singur", 10)
    r += 1
    cap_dim = ["Ce verifici", "Notă 1-5", "Cine a evaluat", "Cum verifici",
               "Ce ai observat concret", "", "", "Ce s-ar putea face", "", ""]
    for j, t in enumerate(cap_dim, start=1):
        ws.cell(row=r, column=j, value=t)
    style_header(ws, r, 10, height=28)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
    ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
    dim_first = r + 1
    for k, (dim, cum) in enumerate(DIMENSIUNI):
        rr = dim_first + k
        c = ws.cell(row=rr, column=1, value=dim)
        c.font = Font(size=10, bold=True, color=INK)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor=TEAL_LIGHT)
        c.border = BORDER
        h = ws.cell(row=rr, column=4, value=cum)
        h.font = Font(size=9, italic=True, color=GREY_TXT)
        h.alignment = Alignment(wrap_text=True, vertical="center")
        h.fill = PatternFill("solid", fgColor=CALC_FILL)
        h.border = BORDER
        for j in (2, 3):
            c = ws.cell(row=rr, column=j)
            c.fill = PatternFill("solid", fgColor=INPUT_FILL)
            c.border = BORDER
            c.font = Font(size=10)
            if j == 2:
                c.number_format = F_INT
                c.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=7)
        ws.merge_cells(start_row=rr, start_column=8, end_row=rr, end_column=10)
        for j in (5, 8):
            c = ws.cell(row=rr, column=j)
            c.fill = PatternFill("solid", fgColor=INPUT_FILL)
            c.border = BORDER
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr].height = 34
    dim_last = dim_first + len(DIMENSIUNI) - 1
    R_DIM = f"'{nume_foaie}'!$B${dim_first}:$B${dim_last}"
    MEDIA_SPEC = f'IFERROR(AVERAGE({R_DIM}),"")'
    NR_DIM = f'COUNT({R_DIM})'
    ULTIMA_DIM = f"'{nume_foaie}'!$B${dim_last}"

    # scoruri
    SCOR_PUBLIC = f'IF({NOTA_POND}="","",({NOTA_POND}-1)/4*100)'
    SCOR_SPEC = f'IF({MEDIA_SPEC}="","",({MEDIA_SPEC}-1)/4*100)'
    r = dim_last + 2
    _antet_sectiune(ws, r, "3. SCORUL DE PERCEPȚIE — cât de bine e văzută firma, fără să întrebi clienții", 10)
    r += 1
    scor_row = {}
    SCORURI = [
        ("Nota publică (medie ponderată)", f'={NOTA_POND}', F_NUM,
         "Media notelor de pe platforme, ponderată cu numărul de recenzii."),
        ("Total recenzii publice", f'={TOTAL_REC}', F_INT,
         "Sub 20 de recenzii, nota publică e orientativă."),
        ("Scor din recenzii (0-100)", f'={SCOR_PUBLIC}', '0',
         "Nota publică adusă pe scara 0-100, ca să se poată compara cu restul."),
        ("Scor specialist (0-100)", f'={SCOR_SPEC}', '0',
         "Media celor zece note de mai sus, adusă pe scara 0-100."),
        ("SCOR DE PERCEPȚIE", None, '0',
         "Media celor două de mai sus. NU este NPS — este ce se poate ști fără să întrebi clienții."),
        ("Cât te poți baza pe el", None, None, "Depinde de câte recenzii sunt și dacă ai făcut evaluarea."),
        ("Ce spune diferența", None, None,
         "Cele trei surse măsoară lucruri diferite. Când nu sunt de acord, dezacordul e informația."),
        ("NPS real (dacă există)", f'={nps_scor}', '0',
         "Din clienții întrebați efectiv. Afișat separat, niciodată amestecat cu scorul de percepție."),
        ("Clienți întrebați", f'={nps_nr}', F_INT, "Câți clienți au dat o notă."),
    ]
    for eticheta, val, fmt, expl in SCORURI:
        scor_row[eticheta] = r
        principal = eticheta == "SCOR DE PERCEPȚIE"
        c = ws.cell(row=r, column=1, value=eticheta)
        c.font = Font(size=11 if principal else 10, bold=True, color=INK)
        v = ws.cell(row=r, column=2, value=val)
        if fmt:
            v.number_format = fmt
        v.fill = PatternFill("solid", fgColor=CALC_FILL)
        v.border = BORDER
        v.font = Font(size=12 if principal else 10, bold=principal, color=INK if principal else GREY_TXT)
        v.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
        e = ws.cell(row=r, column=4, value=expl)
        e.font = Font(size=9, italic=True, color=GREY_TXT)
        e.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 20
        r += 1

    B = lambda eticheta: f"'{nume_foaie}'!$B${scor_row[eticheta]}"
    sp, ss = B("Scor din recenzii (0-100)"), B("Scor specialist (0-100)")
    ws[f"B{scor_row['SCOR DE PERCEPȚIE']}"] = (
        f'=IF(AND({sp}="",{ss}=""),"",IF({sp}="",{ss},IF({ss}="",{sp},({sp}+{ss})/2)))')
    tot_rec, nr_dim = B("Total recenzii publice"), f"{NR_DIM}"
    ws[f"B{scor_row['Cât te poți baza pe el']}"] = (
        f'=IF(AND({tot_rec}=0,{nr_dim}=0),"deloc — nu e completat nimic",'
        f'IF({tot_rec}>=20,IF({nr_dim}>=8,"bine","parțial — fă și evaluarea de specialist"),'
        f'IF({nr_dim}>=8,"parțial — prea puține recenzii publice","slab — completează mai mult")))')
    ws[f"B{scor_row['Cât te poți baza pe el']}"].alignment = Alignment(horizontal="left")
    ws[f"B{scor_row['Cât te poți baza pe el']}"].number_format = "General"

    # Cele trei surse masoara lucruri diferite; dezacordul dintre ele e concluzia.
    # NPS-ul se aduce pe 0-100 doar ca sa fie comparabil ca ordin de marime.
    perc = B("SCOR DE PERCEPȚIE")
    nps_v = B("NPS real (dacă există)")
    nps_n = B("Clienți întrebați")
    npsn = f'IF({nps_v}="","",({nps_v}+100)/2)'
    celula_dif = ws[f"B{scor_row['Ce spune diferența']}"]
    celula_dif.value = (
        f'=IF(AND({sp}="",{ss}=""),"Nu e completată nicio sursă.",'
        f'IF(AND({nps_n}>=5,{nps_v}<>"",{perc}<>"",{perc}-({npsn})>=20),'
        f'"Publicul larg vede firma mai bine decât o văd clienții ei. Diferența se plătește în reveniri: '
        f'oamenii cumpără o dată și nu mai revin.",'
        f'IF(AND({nps_n}>=5,{nps_v}<>"",{perc}<>"",({npsn})-{perc}>=20),'
        f'"Clienții existenți sunt mai mulțumiți decât pare din afară. E o problemă de imagine, nu de livrare — '
        f'cere-le o recenzie celor mulțumiți.",'
        f'IF(AND({sp}<>"",{ss}<>"",{sp}-{ss}>=20),'
        f'"Arată mai bine pe internet decât la o verificare atentă. De obicei: recenzii vechi, puține sau cerute selectiv.",'
        f'IF(AND({sp}<>"",{ss}<>"",{ss}-{sp}>=20),'
        f'"Firma e mai bună decât se vede. Problema e de vizibilitate, nu de calitate — și e mult mai ieftin de rezolvat.",'
        f'"Sursele spun cam același lucru. Te poți baza pe concluzie.")))))')
    celula_dif.alignment = Alignment(horizontal="left", wrap_text=True, vertical="center")
    celula_dif.number_format = "General"
    ws.merge_cells(start_row=scor_row["Ce spune diferența"], start_column=2,
                   end_row=scor_row["Ce spune diferența"], end_column=3)
    ws.row_dimensions[scor_row["Ce spune diferența"]].height = 32

    # drumul catre NPS
    r += 1
    _antet_sectiune(ws, r, "4. DRUMUL CĂTRE NPS — unde e firma acum și ce urmează", 10)
    r += 1
    ws.cell(row=r, column=1, value="Nivel").font = Font(bold=True, size=10, color=WHITE)
    ws.cell(row=r, column=2, value="Ai?").font = Font(bold=True, size=10, color=WHITE)
    ws.cell(row=r, column=3, value="Cât costă").font = Font(bold=True, size=10, color=WHITE)
    ws.cell(row=r, column=4, value="Ce îți dă").font = Font(bold=True, size=10, color=WHITE)
    for k in range(1, 11):
        ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=INK)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
    r += 1
    NIVELE = [
        ("Nivelul 1 — ce se vede public",
         f'=IF({tot_rec}>0,"Da","Nu încă")', "0 lei, 30 de minute",
         "Nota și recenziile de pe Google. Nu întrebi pe nimeni, dar vezi ce văd clienții noi înainte să sune."),
        ("Nivelul 2 — părerea specialistului",
         f'=IF({nr_dim}>=5,"Da","Nu încă")', "0 lei, 1-2 ore",
         "Cele zece verificări de mai sus. Structurate, deci comparabile între firme și în timp."),
        ("Nivelul 3 — NPS real",
         f'=IF({nps_nr}>0,"Da","Nu încă")', "0 lei, 2 ore de telefoane",
         "O singură întrebare, pusă clienților: „de la 0 la 10, cât de probabil ne-ați recomanda?”. "
         "Singura care îți spune ce cred clienții tăi, nu ce cred străinii de pe internet."),
    ]
    nivel_rows = []
    for eticheta, formula, cost, ce_da in NIVELE:
        nivel_rows.append(r)
        c = ws.cell(row=r, column=1, value=eticheta)
        c.font = Font(size=10, bold=True, color=INK)
        c.border = BORDER
        v = ws.cell(row=r, column=2, value=formula)
        v.fill = PatternFill("solid", fgColor=CALC_FILL)
        v.border = BORDER
        v.alignment = Alignment(horizontal="center")
        v.font = Font(size=10, bold=True, color=GREY_TXT)
        cc = ws.cell(row=r, column=3, value=cost)
        cc.font = Font(size=9, color=GREY_TXT)
        cc.border = BORDER
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
        d = ws.cell(row=r, column=4, value=ce_da)
        d.font = Font(size=9, color="333333")
        d.alignment = Alignment(wrap_text=True, vertical="center")
        d.border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1
    for rr in nivel_rows:
        ws.conditional_formatting.add(
            f"B{rr}", CellIsRule(operator="equal", formula=['"Da"'],
                                 fill=PatternFill("solid", fgColor="D6F0D6"),
                                 font=Font(bold=True, color="1B7A3D")))
        ws.conditional_formatting.add(
            f"B{rr}", CellIsRule(operator="equal", formula=['"Nu încă"'],
                                 fill=PatternFill("solid", fgColor="FBE3C2"),
                                 font=Font(bold=True, color="8A6D1F")))

    r += 1
    ws.cell(row=r, column=1, value="Următorul pas").font = Font(bold=True, size=10, color=INK)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
    pas = ws.cell(row=r, column=2, value=(
        f'=IF({tot_rec}=0,"Caută firma pe Google și notează nota și numărul de recenzii, în tabelul de sus. '
        f'Durează cinci minute și îți arată ce vede un client nou.",'
        f'IF({nr_dim}<5,"Fă evaluarea de specialist: cele zece verificări de mai sus. '
        f'Trimite o cerere reală prin formular și cronometrează răspunsul — de obicei acolo apare prima surpriză.",'
        f'IF({nps_nr}=0,"Ai tot ce se putea afla fără să întrebi clienții. Mai departe nu se poate ghici: '
        f'sună zece clienți și pune-le o singură întrebare, de la 0 la 10. Două ore de telefoane.",'
        f'IF({nps_nr}<10,"Ai început să întrebi clienții. Mai sună câțiva: sub zece răspunsuri, cifra încă nu e concluzivă.",'
        f'"Ai toate cele trei niveluri. Compară-le: dacă nota publică e mult peste NPS, firma arată mai bine pe internet '
        f'decât e în relația reală — iar asta se vede în reveniri."))))'))
    pas.fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    pas.border = BORDER
    pas.font = Font(size=10, color="1A2E2E")
    pas.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 42

    return dict(foaie=nume_foaie, nota_publica=NOTA_POND, total_recenzii=TOTAL_REC,
                scor_public=SCOR_PUBLIC, scor_specialist=SCOR_SPEC, nr_dimensiuni=NR_DIM,
                scor_perceptie=B("SCOR DE PERCEPȚIE"), increderea=B("Cât te poți baza pe el"),
                recomanda_specialist=ULTIMA_DIM)
