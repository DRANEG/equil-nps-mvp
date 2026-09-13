# -*- coding: utf-8 -*-
"""
EQUIL Growth Intelligence — generator workbook v0.2

Ruleaza: python3 tools/build_equil_workbook.py
Rezultat: dist/EQUIL_Growth_Intelligence_v0_2.xlsx

Fata de v0.1:
  - date separate pe entitati (clienti / produse / oameni / vanzari / tinte / NPS)
  - foaie de parametri (perioade + praguri) => KPI recalculabili pe perioada
  - comparatie perioada curenta vs. perioada anterioara
  - checklist de colectare date: ce camp, cine il poate da, din ce sistem, ce deblocheaza
  - agenda de contacte/data owners la client
  - semnale si oportunitati generate automat din analiza
  - fara UNIQUE/FILTER (compatibil Excel 2016+, LibreOffice, Google Sheets)
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.comments import Comment

from glosar import DICTIONAR

# ---------------------------------------------------------------- constante

import os as _os


def _n(name, default):
    """Numarul de randuri pre-formatate; se poate reduce prin variabile de mediu
    (folosit de testul automat din tools/test_workbook.py)."""
    return int(_os.environ.get(name, default))


N_CLIENTI, N_PRODUSE, N_OAMENI = _n("N_CLIENTI", 200), _n("N_PRODUSE", 100), _n("N_OAMENI", 60)
N_VANZARI, N_TINTE, N_NPS = _n("N_VANZARI", 1000), _n("N_TINTE", 300), _n("N_NPS", 500)
N_OPP, N_CONTACTE = _n("N_OPP", 60), _n("N_CONTACTE", 30)
OUT = _os.environ.get("OUT", "dist/EQUIL_Growth_Intelligence_v0_2.xlsx")

from istoric import adauga_istoric
from piata import adauga_concurenta, adauga_perceptie
from stil import (verifica, INK, TEAL, TEAL_LIGHT, SAND, INPUT_FILL, CALC_FILL, WHITE, GREY_TXT,
                  C_SER1, C_SER2, C_MUTED, C_ORD, C_GOOD, C_WARN, C_BAD,
                  F_MONEY, F_PCT, F_PCT2, F_DATE, F_NUM, F_INT, BORDER,
                  style_header, title_block, label_value, add_dv, ym, ymd, mon, pct)

wb = Workbook()
wb.remove(wb.active)
# Fara asta, Excel afiseaza celulele calculate (si graficele) goale la prima deschidere.
wb.calculation.fullCalcOnLoad = True


# ---------------------------------------------------------------- helpers







# -------------------------------------------------- definitia foilor de date
# fiecare camp: (tech, ro, kind, width, descriere, obligatoriu, format_excel,
#                exemplu, sursa, rol_furnizor, deblocheaza)
# kind: "in" = se completeaza de client / consultant ; "calc" = formula

def C(tech, ro, kind, width, desc, oblig, fmt, example, sursa, rol, unlocks):
    return dict(tech=tech, ro=ro, kind=kind, width=width, desc=desc, oblig=oblig,
                fmt=fmt, example=example, sursa=sursa, rol=rol, unlocks=unlocks)


SH_CLIENTI, SH_PRODUSE, SH_OAMENI = "10_CLIENTI", "11_PRODUSE", "12_OAMENI"
SH_VANZARI, SH_TINTE, SH_NPS = "13_VANZARI", "14_TINTE", "15_NPS"
SH_PAR = "09_PARAMETRI"

SCHEMA = {}

SCHEMA[SH_CLIENTI] = dict(
    label="Clienți (nomenclator)",
    rows=N_CLIENTI,
    cols=[
        C("Customer_ID", "ID client", "in", 13, "Cod unic de client, exact cum apare în ERP/CRM. Se folosește în foaia de vânzări.", "Obligatoriu", None, "C001", "ERP / Facturare", "IT / ERP", "Toate analizele pe client"),
        C("Customer_Name", "Denumire client", "in", 26, "Denumirea comercială a clientului.", "Obligatoriu", None, "Alfa Distribution SRL", "ERP / Facturare", "Director Vânzări", "Raport client, top clienți"),
        C("Segment", "Segment", "in", 14, "Segmentul de client (ex. Key Account, Mid, Small, Online).", "Recomandat", None, "Key Account", "CRM", "Director Vânzări", "Analiză pe segmente"),
        C("Channel", "Canal", "in", 13, "Canalul prin care se vinde: Retail, Wholesale, Online, Direct, Partener.", "Recomandat", None, "Wholesale", "CRM", "Director Vânzări", "Mix de canal, oportunități cross-sell"),
        C("Region", "Regiune / Țară", "in", 14, "Regiunea sau țara clientului.", "Recomandat", None, "București-Ilfov", "ERP / Facturare", "Sales Ops / CRM", "Analiză geografică"),
        C("Country", "Țară", "in", 11, "Cod țară (RO, HU, BG...).", "Opțional", None, "RO", "ERP / Facturare", "Sales Ops / CRM", "Export vs. intern"),
        C("Account_Owner_ID", "ID responsabil cont", "in", 15, "ID-ul persoanei din 12_OAMENI care răspunde de acest client.", "Recomandat", None, "S001", "CRM", "Director Vânzări", "Analiză pe oameni, plan de acțiune"),
        C("Customer_Since", "Client din data", "in", 14, "Data primei relații comerciale.", "Recomandat", F_DATE, "2021-03-15", "ERP / Facturare", "Sales Ops / CRM", "Vechime, retenție, clienți noi"),
        C("Status", "Status", "in", 12, "Activ / Inactiv / Prospect / Pierdut.", "Recomandat", None, "Activ", "CRM", "Director Vânzări", "Reactivare clienți dormanți"),
        C("Payment_Terms_Days", "Termen plată (zile)", "in", 14, "Termenul de plată agreat, în zile.", "Opțional", F_INT, "30", "Balanță / Contabilitate", "CFO / Contabilitate", "Cash-flow, risc"),
        C("Credit_Limit", "Limită credit", "in", 13, "Plafonul de credit acordat.", "Opțional", F_MONEY, "50000", "Balanță / Contabilitate", "CFO / Contabilitate", "Risc comercial"),
        C("Notes", "Observații", "in", 28, "Context util: istoric, relație, riscuri cunoscute.", "Opțional", None, "Renegociere contract în T3", "Interviu / Workshop", "Director Vânzări", "Context pentru recomandări"),
    ])

SCHEMA[SH_PRODUSE] = dict(
    label="Produse / servicii (nomenclator)",
    rows=N_PRODUSE,
    cols=[
        C("Product_ID", "ID produs", "in", 13, "Cod unic de produs/serviciu (SKU), exact ca în ERP.", "Obligatoriu", None, "P001", "ERP / Facturare", "IT / ERP", "Analiză produs, cross-sell"),
        C("Product_Name", "Denumire produs", "in", 26, "Denumirea produsului sau serviciului.", "Obligatoriu", None, "Serviciu mentenanță Premium", "ERP / Facturare", "Operațiuni / Logistică", "Raport produs"),
        C("Category", "Categorie", "in", 16, "Categoria/linia de business.", "Recomandat", None, "Servicii recurente", "ERP / Facturare", "Operațiuni / Logistică", "Mix de categorie, marjă"),
        C("Subcategory", "Subcategorie", "in", 16, "Nivel suplimentar de grupare, dacă există.", "Opțional", None, "Contracte anuale", "ERP / Facturare", "Operațiuni / Logistică", "Detaliere mix"),
        C("List_Price", "Preț de listă", "in", 13, "Prețul standard de listă, fără TVA.", "Recomandat", F_MONEY, "1200", "Excel intern", "Director Vânzări", "Analiză discount, erodare preț"),
        C("Standard_Cost", "Cost standard", "in", 13, "Costul unitar standard (COGS).", "Recomandat", F_MONEY, "760", "Balanță / Contabilitate", "CFO / Contabilitate", "Marjă pe produs"),
        C("Is_Recurring", "Recurent (Da/Nu)", "in", 14, "Produs recurent/abonament sau vânzare unică.", "Opțional", None, "Da", "Excel intern", "Director Vânzări", "Venit recurent vs. one-off"),
        C("Launch_Date", "Data lansării", "in", 13, "Când a fost lansat produsul.", "Opțional", F_DATE, "2024-09-01", "Excel intern", "Marketing", "Creștere produse noi"),
        C("Status", "Status", "in", 12, "Activ / Retras / În pilot.", "Opțional", None, "Activ", "Excel intern", "Operațiuni / Logistică", "Curățare portofoliu"),
    ])

SCHEMA[SH_OAMENI] = dict(
    label="Oameni (echipa comercială / operațională)",
    rows=N_OAMENI,
    cols=[
        C("Person_ID", "ID persoană", "in", 12, "Cod intern al persoanei (nu CNP, nu date personale sensibile).", "Obligatoriu", None, "S001", "HR / State de plată", "HR / Resurse Umane", "Analiză productivitate"),
        C("Person_Name", "Nume", "in", 22, "Nume și prenume sau inițiale, dacă se preferă anonimizarea.", "Obligatoriu", None, "Ana P.", "HR / State de plată", "HR / Resurse Umane", "Plan de coaching"),
        C("Role", "Rol", "in", 18, "Rolul: Sales, Key Account, Customer Service, Operațiuni, Management.", "Obligatoriu", None, "Key Account Manager", "HR / State de plată", "HR / Resurse Umane", "Venit/FTE pe rol"),
        C("Team", "Echipă", "in", 15, "Echipa sau departamentul.", "Recomandat", None, "Vânzări B2B", "HR / State de plată", "HR / Resurse Umane", "Comparație între echipe"),
        C("FTE", "FTE (normă)", "in", 10, "Fracțiune de normă: 1 = full-time, 0.5 = jumătate de normă.", "Obligatoriu", F_NUM, "1", "HR / State de plată", "HR / Resurse Umane", "Venit / FTE, productivitate"),
        C("Start_Date", "Data angajării", "in", 13, "Data intrării în echipă.", "Recomandat", F_DATE, "2023-02-01", "HR / State de plată", "HR / Resurse Umane", "FTE activ pe perioadă"),
        C("Exit_Date", "Data plecării", "in", 13, "Se lasă gol dacă persoana este activă.", "Recomandat", F_DATE, "", "HR / State de plată", "HR / Resurse Umane", "FTE activ, fluctuație"),
        C("Annual_Cost", "Cost anual total", "in", 14, "Cost total angajator pe an (opțional, se poate da agregat pe echipă).", "Opțional", F_MONEY, "96000", "HR / State de plată", "CFO / Contabilitate", "ROI pe om, cost/venit"),
        C("Is_Quota_Carrying", "Are țintă (Da/Nu)", "in", 14, "Dacă persoana are țintă de vânzări.", "Recomandat", None, "Da", "Director Vânzări", "Director Vânzări", "Atingere țintă pe om"),
    ])

SCHEMA[SH_VANZARI] = dict(
    label="Vânzări / tranzacții (faptic)",
    rows=N_VANZARI,
    cols=[
        C("Doc_Date", "Data document", "in", 13, "Data facturii sau a lunii de raportare. O linie = o combinație client × produs × perioadă.", "Obligatoriu", F_DATE, "2026-01-31", "ERP / Facturare", "CFO / Contabilitate", "Toate KPI-urile și trendurile"),
        C("Period", "Perioadă (AAAA-LL)", "calc", 13, "Se calculează automat din data documentului.", "Auto", None, "2026-01", "Calculat automat", "EQUIL (calculat)", "Grupare lunară"),
        C("Doc_Type", "Tip document", "in", 13, "Factură / Storno / Avans. Stornările se trec cu valori negative.", "Recomandat", None, "Factură", "ERP / Facturare", "CFO / Contabilitate", "Corectitudinea veniturilor"),
        C("Customer_ID", "ID client", "in", 12, "Trebuie să existe în 10_CLIENTI.", "Obligatoriu", None, "C001", "ERP / Facturare", "IT / ERP", "Analiză client, concentrare"),
        C("Product_ID", "ID produs", "in", 12, "Trebuie să existe în 11_PRODUSE.", "Obligatoriu", None, "P001", "ERP / Facturare", "IT / ERP", "Analiză produs, cross-sell"),
        C("Person_ID", "ID persoană vânzări", "in", 14, "Cine a făcut vânzarea; trebuie să existe în 12_OAMENI.", "Recomandat", None, "S001", "CRM", "Sales Ops / CRM", "Productivitate pe om"),
        C("Units", "Cantitate", "in", 11, "Număr de unități / ore / licențe.", "Recomandat", F_NUM, "100", "ERP / Facturare", "IT / ERP", "Preț mediu, volum vs. preț"),
        C("Revenue_Net", "Venit net", "in", 13, "Venit fără TVA, după discount.", "Obligatoriu", F_MONEY, "10000", "ERP / Facturare", "CFO / Contabilitate", "Venit, creștere, țintă"),
        C("Discount_Value", "Valoare discount", "in", 13, "Discountul acordat, în valoare absolută.", "Opțional", F_MONEY, "800", "ERP / Facturare", "CFO / Contabilitate", "Erodare preț, marjă"),
        C("COGS", "Cost direct (COGS)", "in", 14, "Costul direct aferent vânzării.", "Obligatoriu", F_MONEY, "7500", "Balanță / Contabilitate", "CFO / Contabilitate", "Profit brut, marjă"),
        C("Gross_Profit", "Profit brut", "calc", 12, "Se calculează automat: Venit net − COGS.", "Auto", F_MONEY, "2500", "Calculat automat", "EQUIL (calculat)", "Marjă, oportunități de marjă"),
        C("Margin_Pct", "Marjă %", "calc", 10, "Se calculează automat: Profit brut / Venit net.", "Auto", F_PCT, "25%", "Calculat automat", "EQUIL (calculat)", "Semnal marjă sub benchmark"),
        C("Customer_Name", "Denumire client", "calc", 22, "Se aduce automat din 10_CLIENTI.", "Auto", None, "Alfa Distribution SRL", "Calculat automat", "EQUIL (calculat)", "Verificare vizuală"),
        C("Product_Name", "Denumire produs", "calc", 22, "Se aduce automat din 11_PRODUSE.", "Auto", None, "Serviciu mentenanță", "Calculat automat", "EQUIL (calculat)", "Verificare vizuală"),
        C("Currency", "Monedă", "in", 9, "Moneda documentului. Ideal toate liniile în aceeași monedă.", "Recomandat", None, "RON", "ERP / Facturare", "CFO / Contabilitate", "Consistența raportării"),
        C("Notes", "Observații", "in", 20, "Context pe linie, dacă e nevoie.", "Opțional", None, "", "Interviu / Workshop", "Director Vânzări", "Context"),
    ])

SCHEMA[SH_TINTE] = dict(
    label="Ținte / buget",
    rows=N_TINTE,
    cols=[
        C("Period_Date", "Perioadă (data)", "in", 13, "Ultima zi a lunii/trimestrului pentru care e stabilită ținta.", "Obligatoriu", F_DATE, "2026-01-31", "Excel intern", "CFO / Contabilitate", "Atingere țintă"),
        C("Scope", "Nivel țintă", "in", 13, "Companie / Persoană / Produs / Client.", "Obligatoriu", None, "Companie", "Excel intern", "Director Vânzări", "Comparație realizat vs. plan"),
        C("Scope_ID", "ID nivel", "in", 12, "Se lasă gol pentru Companie; altfel ID-ul persoanei/produsului/clientului.", "Recomandat", None, "S001", "Excel intern", "Director Vânzări", "Gap pe om/produs"),
        C("Metric", "Indicator", "in", 13, "Venit / Profit brut / Unități.", "Obligatoriu", None, "Venit", "Excel intern", "CFO / Contabilitate", "Ce se măsoară"),
        C("Target_Value", "Valoare țintă", "in", 13, "Valoarea planificată pentru perioada respectivă.", "Obligatoriu", F_MONEY, "120000", "Excel intern", "CFO / Contabilitate", "Gap față de țintă"),
        C("Notes", "Observații", "in", 22, "Ipoteze, revizuiri de buget.", "Opțional", None, "Buget revizuit în T2", "Interviu / Workshop", "CFO / Contabilitate", "Context"),
    ])

SCHEMA[SH_NPS] = dict(
    label="NPS / feedback clienți",
    rows=N_NPS,
    cols=[
        C("Response_ID", "ID răspuns", "in", 12, "Cod unic al răspunsului.", "Obligatoriu", None, "R0001", "Chestionar NPS", "Marketing", "Trasabilitate"),
        C("Survey_Date", "Data răspunsului", "in", 14, "Data la care s-a primit răspunsul.", "Obligatoriu", F_DATE, "2026-02-10", "Chestionar NPS", "Marketing", "NPS pe perioadă, trend"),
        C("Customer_ID", "ID client", "in", 12, "Clientul care a răspuns; trebuie să existe în 10_CLIENTI.", "Obligatoriu", None, "C001", "Chestionar NPS", "Customer Service", "NPS pe client, corelație cu venitul"),
        C("Respondent_Name", "Nume respondent", "in", 20, "Persoana care a răspuns (dacă e permis de GDPR să fie nominalizată).", "Recomandat", None, "Mihai I.", "Chestionar NPS", "Customer Service", "Acțiuni de recuperare nominale"),
        C("Respondent_Role", "Rol respondent", "in", 18, "Rolul persoanei: decident, utilizator, achiziții.", "Recomandat", None, "Decident", "Chestionar NPS", "Customer Service", "Cine e nemulțumit: user vs. decident"),
        C("Score", "Scor NPS (0-10)", "in", 13, "Răspunsul la întrebarea „Cât de probabil ne-ați recomanda?”, de la 0 la 10.", "Obligatoriu", F_INT, "9", "Chestionar NPS", "Marketing", "NPS, % promoteri/detractori"),
        C("NPS_Category", "Categorie", "calc", 12, "Se calculează automat: Promotor (9-10) / Pasiv (7-8) / Detractor (0-6).", "Auto", None, "Promotor", "Calculat automat", "EQUIL (calculat)", "Segmentare NPS"),
        C("Reason", "Motivul scorului", "in", 34, "Comentariul liber al respondentului — cea mai valoroasă informație calitativă.", "Recomandat", None, "Livrări întârziate în ianuarie", "Chestionar NPS", "Customer Service", "Cauze concrete, plan de acțiune"),
        C("Theme", "Temă", "in", 16, "Tema principală: Preț / Calitate / Livrare / Suport / Produs / Relație.", "Recomandat", None, "Livrare", "Chestionar NPS", "Customer Service", "Gruparea cauzelor"),
        C("Channel", "Canal colectare", "in", 14, "Email / Telefon / Față în față / Formular web.", "Opțional", None, "Email", "Chestionar NPS", "Marketing", "Calitatea eșantionului"),
        C("Survey_Wave", "Val / campanie", "in", 13, "Identificatorul valului de măsurare.", "Opțional", None, "2026-T1", "Chestionar NPS", "Marketing", "Comparație între valuri"),
        C("Follow_Up_Owner", "Responsabil follow-up", "in", 18, "Cine preia clientul după răspuns.", "Recomandat", None, "S001", "CRM", "Director Vânzări", "Recuperare detractori"),
        C("Follow_Up_Status", "Status follow-up", "in", 15, "Neînceput / În lucru / Rezolvat.", "Recomandat", None, "În lucru", "CRM", "Customer Service", "Urmărirea acțiunilor"),
    ])










def cl(sheet, tech):
    """Litera coloanei pentru un camp tehnic dintr-o foaie de date."""
    cols = SCHEMA[sheet]["cols"]
    for i, c in enumerate(cols, start=1):
        if c["tech"] == tech:
            return get_column_letter(i)
    raise KeyError(f"{sheet}.{tech}")


# In foile de date: rand 1-2 = titlu, rand 4 = antet, datele incep pe randul 5.
DATA_FIRST_ROW = 5


def R(sheet, tech, first=DATA_FIRST_ROW, last=None):
    """Referinta absoluta la intervalul de date al unei coloane."""
    last = last or SCHEMA[sheet]["rows"] + DATA_FIRST_ROW - 1
    L = cl(sheet, tech)
    return f"'{sheet}'!${L}${first}:${L}${last}"


# ---------------------------------------------------------------- 90_LISTE

LISTE = [
    ("Roluri", ["CEO / Administrator", "CFO / Contabilitate", "Director Vânzări",
                "Sales Ops / CRM", "Marketing", "HR / Resurse Umane", "IT / ERP",
                "Operațiuni / Logistică", "Customer Service", "Achiziții",
                "Consultant extern", "EQUIL (calculat)"]),
    ("Status_date", ["Nesolicitat", "Solicitat", "Primit parțial", "Primit", "Validat",
                     "Blocat", "Nu există", "Nu se aplică"]),
    ("Obligativitate", ["Obligatoriu", "Recomandat", "Opțional", "Auto"]),
    ("Segment_client", ["Key Account", "Mid-market", "Small", "Online", "Distribuitor", "Altele"]),
    ("Canal", ["Retail", "Wholesale", "Online", "Direct", "Partener", "Marketplace"]),
    ("Tip_oportunitate", ["Recuperare client", "Cross-sell", "Up-sell", "Marjă / preț",
                          "Productivitate vânzări", "Risc concentrare", "Recuperare NPS",
                          "Creștere produs", "Gap față de țintă", "Reactivare client",
                          "Dezvoltare oameni", "Proces / operațional"]),
    ("Prioritate", ["HIGH", "MEDIUM", "LOW"]),
    ("Status_oportunitate", ["Identificată", "Validată cu clientul", "În execuție",
                             "Realizată", "Respinsă", "Amânată"]),
    ("Set_de_date", ["Profil firmă", "Clienți", "Produse", "Oameni", "Vânzări",
                     "Ținte / buget", "NPS / feedback", "Financiar", "Documente / context"]),
    ("Sursa_date", ["ERP / Facturare", "CRM", "Excel intern", "Balanță / Contabilitate",
                    "HR / State de plată", "Chestionar NPS", "Interviu / Workshop",
                    "Site / public", "Calculat automat"]),
    ("Metrica_tinta", ["Venit", "Profit brut", "Unități"]),
    ("Nivel_tinta", ["Companie", "Persoană", "Produs", "Client"]),
    ("Status_generic", ["Activ", "Inactiv", "Prospect", "Pierdut", "În pilot", "Retras"]),
    ("Da_Nu", ["Da", "Nu"]),
    ("Categorie_NPS", ["Promotor", "Pasiv", "Detractor"]),
    ("Tema_NPS", ["Preț", "Calitate", "Livrare", "Suport", "Produs", "Relație",
                  "Proces comandă", "Altele"]),
    ("Tip_document", ["Factură", "Storno", "Avans", "Abonament"]),
    ("Moneda", ["RON", "EUR", "USD", "GBP"]),
    ("Status_followup", ["Neînceput", "În lucru", "Rezolvat", "Escaladat"]),
]

ws = wb.create_sheet("90_LISTE")
for j, (name, vals) in enumerate(LISTE, start=1):
    ws.cell(row=1, column=j, value=name)
    ws.column_dimensions[get_column_letter(j)].width = max(16, len(name) + 4)
    for i, v in enumerate(vals, start=2):
        ws.cell(row=i, column=j, value=v).font = Font(size=10)
style_header(ws, 1, len(LISTE), fill=INK, height=22)
ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False


def lista_ref(name):
    j = [n for n, _ in LISTE].index(name) + 1
    L = get_column_letter(j)
    n = len([v for nm, v in LISTE if nm == name][0])
    return f"'90_LISTE'!${L}$2:${L}${n + 1}"




# ---------------------------------------------------------------- 00_GHID

ws = wb.create_sheet("00_GHID", 0)
title_block(ws, "EQUIL GROWTH INTELLIGENCE — v0.2",
            "Pachet de colectare a datelor + motor de KPI, semnale și oportunități. Completează doar celulele galbene.")
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 96
ws.column_dimensions["C"].width = 30

ghid = [
    ("", ""),
    ("LA CE FOLOSEȘTE", "Transformă datele brute ale unei firme în KPI, semnale de risc și oportunități de creștere prioritizate, plus un raport gata de prezentat."),
    ("CUM SE FOLOSEȘTE", ""),
    ("Pas 1 — Profil", "Completează 01_FIRMA în discuția de kickoff. Durează 20 de minute și îți dă contextul pentru tot restul."),
    ("Pas 2 — Oameni", "Completează 02_CONTACTE: cine, din firmă, deține fiecare set de date. Fără asta, colectarea se blochează."),
    ("Pas 3 — Cerere date", "Trimite 03_CHECKLIST_DATE. Fiecare rând spune ce câmp ceri, din ce sistem se scoate, cine îl poate da și ce deblochează. Filtrează pe „Obligatoriu” pentru setul minim viabil."),
    ("Pas 4 — Parametri", "Setează în 09_PARAMETRI perioada analizată, perioada de comparație și pragurile. Toate calculele se recalculează automat."),
    ("Pas 5 — Import date", "Lipește datele în 10_CLIENTI, 11_PRODUSE, 12_OAMENI, 13_VANZARI, 14_TINTE, 15_NPS. Coloanele gri au formule — nu le suprascrie."),
    ("Pas 6 — Citește", "Verifică 20_KPI, apoi 21/22/23 pentru semnale pe clienți, produse și oameni."),
    ("Pas 7 — Prioritizează", "Preia semnalele în 30_OPORTUNITATI. Scor = Impact × Probabilitate × Ușurință. Prioritatea se calculează din pragurile din 09_PARAMETRI."),
    ("Pas 8 — Prezintă", "40_RAPORT este structura pentru clientul final; se completează singur din foile anterioare."),
    ("", ""),
    ("COD DE CULORI", ""),
    ("Galben", "Celulă de completat (de tine sau de client)."),
    ("Gri", "Celulă calculată automat. Nu se scrie în ea."),
    ("Verde-închis", "Antet de tabel."),
    ("", ""),
    ("SETUL MINIM VIABIL", "Cu doar 3 lucruri se poate produce un prim raport: (1) vânzări pe 24 de luni cu client, produs, venit net și cost; (2) nomenclatorul de clienți; (3) ținta de venit pe perioadă. Restul îmbogățește analiza."),
    ("ISTORIC RECOMANDAT", "Minim 13 luni (pentru comparație an/an), ideal 24-36 de luni."),
    ("", ""),
    ("REGULI DE IGIENĂ A DATELOR", ""),
    ("Un ID = o entitate", "Fiecare client și produs are un singur ID, constant în toate foile."),
    ("Valori fără TVA", "Veniturile și costurile se dau nete, în aceeași monedă."),
    ("Stornări", "Se trec ca linii separate, cu valori negative."),
    ("Fără celule „mixte”", "Nu combina text și cifre în aceeași celulă (ex. „12.000 lei” → 12000)."),
    ("Date personale", "Pentru NPS și oameni, folosește inițiale sau ID-uri dacă firma nu a agreat nominalizarea. Semnează NDA înainte de a primi datele."),
    ("", ""),
    ("UN FIȘIER PENTRU FIECARE FIRMĂ", ""),
    ("Regula", "Nu amesteca două firme în același fișier. Păstrezi acest fișier ca șablon nemodificat și faci o copie pentru fiecare client."),
    ("Cod de firmă", "Fiecare firmă primește un cod în 01_FIRMA (EQ-0001, EQ-0002…). Codul rămâne același pentru totdeauna, chiar dacă firma își schimbă numele."),
    ("Nume fișier", "EQUIL_<cod>_<NumeFirma>_<AAAA-LL>.xlsx — ex. EQUIL_EQ-0001_AlfaSRL_2026-06.xlsx. Perioada din nume îți spune din ce rulaj e fișierul."),
    ("Refolosire lunară", "Pentru luna următoare la aceeași firmă: copiază fișierul, schimbă perioada în 09_PARAMETRI și adaugă noile linii de vânzări. Nu ștergi istoricul."),
    ("Șablonul", "Ține o copie curată, fără date, într-un folder separat („_SABLON”). Din ea pleacă fiecare firmă nouă."),
    ("", ""),
    ("CÂT DUREAZĂ (estimat, prima rulare)", ""),
    ("01_FIRMA", "20-30 min, în discuția de kickoff, împreună cu clientul."),
    ("02_CONTACTE", "10-15 min, tot la kickoff. Cea mai bună investiție de timp din tot procesul."),
    ("03_CHECKLIST_DATE", "15 min să-l trimiți și să-l explici. Apoi 3-10 zile lucrătoare de așteptare, cu 2-3 follow-up-uri scurte."),
    ("Import date", "30-60 min dacă exporturile sunt curate. 2-4 ore dacă trebuie curățate manual (denumiri inconsecvente, valori cu text, ID-uri lipsă)."),
    ("09_PARAMETRI", "5 min. Pragurile se ajustează pe industrie."),
    ("Citit 20-23", "30-45 min de citit semnalele și de verificat dacă au sens."),
    ("30_OPORTUNITATI", "45-60 min: transformi semnalele în oportunități, estimezi impactul, pui responsabili."),
    ("40_RAPORT + grafice", "30-45 min: concluziile automate îți dau schița, tu scrii varianta pentru client."),
    ("TOTAL munca ta", "aprox. 4-6 ore de lucru efectiv, întinse pe 1-2 săptămâni calendaristice (așteptarea datelor domină)."),
    ("Rulările următoare", "1-1,5 ore pe lună la aceeași firmă: adaugi datele noi, schimbi perioada, reciteşti semnalele."),
    ("", ""),
    ("COMPATIBILITATE", "Fără funcții dinamice (UNIQUE/FILTER/LET/IFS). Funcționează în Excel 2016+, Microsoft 365, LibreOffice Calc și Google Sheets. Nu ai nevoie de o versiune mai nouă de Office pentru acest fișier."),
    ("LIMITE ACTUALE", "Pre-formatat pentru 1000 de linii de vânzări, 200 de clienți, 100 de produse, 60 de persoane, 500 de răspunsuri NPS. Pentru volume mai mari, trage formulele în jos."),
    ("URMĂTOAREA VERSIUNE", "Benchmark-uri de industrie, previziune de venit, scor de risc de churn, generare automată a naraţiunii din raport."),
]
r = 3
for a, b in ghid:
    if a == "" and b == "":
        r += 1
        continue
    ca = ws.cell(row=r, column=1, value=a)
    cb = ws.cell(row=r, column=2, value=b)
    is_section = b == ""
    ca.font = Font(bold=True, size=11 if is_section else 10,
                   color=TEAL if is_section else INK)
    cb.font = Font(size=10, color="333333")
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    ca.alignment = Alignment(vertical="top", wrap_text=True)
    if is_section:
        ca.fill = PatternFill("solid", fgColor=TEAL_LIGHT)
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    ws.row_dimensions[r].height = 30 if len(b) > 95 else 16
    r += 1
for key, color in (("Galben", INPUT_FILL), ("Gri", CALC_FILL), ("Verde-închis", TEAL)):
    for rr in range(3, r):
        if ws.cell(row=rr, column=1).value == key:
            ws.cell(row=rr, column=3).fill = PatternFill("solid", fgColor=color)
            ws.cell(row=rr, column=3).border = BORDER


# ---------------------------------------------------------------- 01_FIRMA

ws = wb.create_sheet("01_FIRMA")
title_block(ws, "01 — PROFILUL FIRMEI",
            "Se completează la kickoff, împreună cu clientul. Câmpurile galbene sunt de completat.")
# A = etichete, B/C/D = cei trei ani din istoricul financiar, E = explicații
for _col, _w in (("A", 34), ("B", 17), ("C", 17), ("D", 17), ("E", 58)):
    ws.column_dimensions[_col].width = _w

firma = [
    ("IDENTIFICARE", None, None),
    ("Cod client EQUIL", "", "ID-ul intern al firmei în portofoliul tău. Format recomandat: EQ-0001, EQ-0002…"),
    ("Denumire legală", "", "Exact ca la Registrul Comerțului."),
    ("CUI / Cod fiscal", "", ""),
    ("Nr. Registrul Comerțului", "", ""),
    ("Site web", "", ""),
    ("Industrie / CAEN principal", "", "Determină benchmark-urile folosite."),
    ("Sub-industrie / nișă", "", ""),
    ("An înființare", "", ""),
    ("Model de business", "", "B2B / B2C / Mix / Marketplace."),
    ("Piețe / regiuni acoperite", "", ""),
    ("DIMENSIUNE", None, None),
    ("Număr total angajați", "", "Total, nu doar vânzări."),
    ("Din care în vânzări", "", ""),
    ("Număr clienți activi", "", "Clienți cu cel puțin o achiziție în ultimele 12 luni."),
    ("Număr produse / SKU active", "", ""),
    ("Sezonalitate", "", "Lunile de vârf și lunile slabe."),
    ("SISTEME ȘI DATE", None, None),
    ("ERP / sistem de facturare", "", "Ex. SAGA, WinMentor, SAP, Oracle NetSuite, Smartbill."),
    ("CRM", "", "Ex. HubSpot, Pipedrive, Salesforce, Excel."),
    ("Instrument de raportare", "", "Ex. Excel, Power BI, Looker, nimic."),
    ("Se pot exporta datele în Excel/CSV?", "", "Da / Nu / Parțial — și cine face exportul."),
    ("Istoric disponibil (luni)", "", "Câte luni de date se pot scoate din sisteme."),
    ("Se măsoară deja NPS?", "", "Da / Nu. Dacă da, de când și cum."),
    ("OBIECTIVE ȘI CONTEXT", None, None),
    ("Obiectivul principal pe 12 luni", "", "Exprimat în cifre, dacă se poate."),
    ("Cele 3 provocări majore", "", ""),
    ("Ce au încercat deja și nu a mers", "", ""),
    ("Cine ia decizia finală", "", "Numele și rolul decidentului."),
    ("MANDAT EQUIL", None, None),
    ("Perioada analizată", "", "Trebuie să coincidă cu 09_PARAMETRI."),
    ("Moneda de raportare", "", ""),
    ("Consultant EQUIL responsabil", "", ""),
    ("Data kickoff", "", ""),
    ("Data livrării raportului", "", ""),
    ("NDA semnat", "", "Da / Nu / Data. NDA = acord de confidențialitate; vezi 04_DICTIONAR."),
    ("Nume fișier", "", "Convenție: EQUIL_<cod>_<NumeFirma>_<AAAA-LL>.xlsx — ex. EQUIL_EQ-0001_AlfaSRL_2026-06.xlsx"),
]
r = 4
rand_firma = {}
for label, val, note in firma:
    if val is None:
        c = ws.cell(row=r, column=1, value=label)
        c.font = Font(bold=True, size=11, color=WHITE)
        for k in range(1, 6):
            ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=TEAL)
        ws.row_dimensions[r].height = 20
    else:
        label_value(ws, r, label, val, note, span=3, note_col=5)
        rand_firma[label] = r
    r += 1
ws.sheet_view.showGridLines = False
add_dv(ws, lista_ref("Moneda"), f"B{rand_firma['Moneda de raportare']}")

# 09_PARAMETRI se construiește mai jos; referințele sunt doar text, foaia există la salvare.
ISTORIC = adauga_istoric(ws, r + 1, an_ref=f"'{SH_PAR}'!$B$6",
                         moneda_ref=f"'{SH_PAR}'!$B$10", note_col=5)


# ---------------------------------------------------------------- 02_CONTACTE

CONTACT_COLS = [
    ("Contact_ID", "Cod contact", 11),
    ("Rol_in_firma", "Rol în firmă", 22),
    ("Nume", "Nume și prenume", 24),
    ("Functie_exacta", "Funcția exactă", 24),
    ("Email", "Email", 26),
    ("Telefon", "Telefon", 16),
    ("Seturi_de_date_detinute", "Ce date poate furniza", 40),
    ("Sisteme_la_care_are_acces", "Sisteme la care are acces", 24),
    ("Decident", "Poate aproba transmiterea? (Da/Nu)", 16),
    ("Canal_preferat", "Canal preferat de contact", 18),
    ("Disponibilitate", "Disponibilitate", 18),
    ("Data_solicitarii", "Data solicitării", 14),
    ("Termen_agreat", "Termen agreat", 14),
    ("Status", "Status", 15),
    ("Observatii", "Observații", 30),
]

ws = wb.create_sheet("02_CONTACTE")
title_block(ws, "02 — CINE FURNIZEAZĂ DATELE (data owners)",
            "Harta oamenilor din firma clientului. Fiecare set de date are un om responsabil; fără nume și termen, colectarea se blochează.")
hr = 4
for j, (tech, ro, w) in enumerate(CONTACT_COLS, start=1):
    ws.cell(row=hr, column=j, value=f"{ro}\n{tech}")
    ws.column_dimensions[get_column_letter(j)].width = w
style_header(ws, hr, len(CONTACT_COLS))
ws.freeze_panes = f"C{hr+1}"

CONTACTE_SEED = [
    ("CT01", "CEO / Administrator", "", "", "", "", "Obiective, context strategic, aprobarea transmiterii datelor", "—", "Da", "Întâlnire", "", "", "", "Nesolicitat", "Decide ce date pot ieși din firmă."),
    ("CT02", "CFO / Contabilitate", "", "", "", "", "Venituri, costuri (COGS), balanță, ținte/buget, termene de plată", "ERP, contabilitate", "Da", "Email", "", "", "", "Nesolicitat", "Sursa principală pentru profit brut și marjă."),
    ("CT03", "Director Vânzări", "", "", "", "", "Clienți, segmente, canale, ținte pe oameni, context comercial", "CRM", "Da", "Telefon", "", "", "", "Nesolicitat", "Validează interpretarea semnalelor."),
    ("CT04", "Sales Ops / CRM", "", "", "", "", "Export CRM: clienți, responsabili de cont, pipeline, activitate", "CRM", "Nu", "Email", "", "", "", "Nesolicitat", "De obicei cel care face efectiv exportul."),
    ("CT05", "IT / ERP", "", "", "", "", "Export brut din ERP: facturi, nomenclator produse, clienți", "ERP", "Nu", "Email", "", "", "", "Nesolicitat", "Cere formatul CSV/XLSX, nu PDF."),
    ("CT06", "Marketing", "", "", "", "", "NPS, campanii, lansări de produs, lead-uri", "Survey, web", "Nu", "Email", "", "", "", "Nesolicitat", "Deține de obicei instrumentul de chestionar."),
    ("CT07", "HR / Resurse Umane", "", "", "", "", "Structura echipei, FTE, intrări/ieșiri, cost cu personalul", "HR / salarizare", "Da", "Email", "", "", "", "Nesolicitat", "Atenție la GDPR: cere agregat sau anonimizat."),
    ("CT08", "Customer Service", "", "", "", "", "Reclamații, motive de nemulțumire, follow-up detractori", "Ticketing", "Nu", "Telefon", "", "", "", "Nesolicitat", "Sursa cea mai bogată de cauze calitative."),
    ("CT09", "Operațiuni / Logistică", "", "", "", "", "Livrări, stocuri, termene, capacitate", "ERP / WMS", "Nu", "Email", "", "", "", "Nesolicitat", "Relevant când NPS scade din cauza livrării."),
]
r = hr + 1
for row in CONTACTE_SEED:
    for j, v in enumerate(row, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.border = BORDER
        c.font = Font(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
for rr in range(r, hr + N_CONTACTE + 1):
    for j in range(1, len(CONTACT_COLS) + 1):
        c = ws.cell(row=rr, column=j)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
last_contact = hr + N_CONTACTE
for j in range(1, len(CONTACT_COLS) + 1):
    for rr in range(hr + 1, last_contact + 1):
        ws.cell(row=rr, column=j).fill = PatternFill("solid", fgColor=INPUT_FILL)
for rr in range(hr + 1, last_contact + 1):
    ws.cell(row=rr, column=12).number_format = F_DATE
    ws.cell(row=rr, column=13).number_format = F_DATE

add_dv(ws, lista_ref("Roluri"), f"B{hr+1}:B{last_contact}")
add_dv(ws, lista_ref("Da_Nu"), f"I{hr+1}:I{last_contact}")
add_dv(ws, lista_ref("Status_date"), f"N{hr+1}:N{last_contact}")
ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(CONTACT_COLS))}{last_contact}"
ws.conditional_formatting.add(
    f"N{hr+1}:N{last_contact}",
    CellIsRule(operator="equal", formula=['"Primit"'],
               fill=PatternFill("solid", fgColor="D6F0D6")))
ws.conditional_formatting.add(
    f"N{hr+1}:N{last_contact}",
    CellIsRule(operator="equal", formula=['"Blocat"'],
               fill=PatternFill("solid", fgColor="F8D7DA")))
ws.conditional_formatting.add(
    f"N{hr+1}:N{last_contact}",
    CellIsRule(operator="equal", formula=['"Validat"'],
               fill=PatternFill("solid", fgColor="BCE3BC")))

ws["A2"] = ("Sfat: nu cere „toate datele”. Cere de la fiecare om exact setul din 03_CHECKLIST_DATE "
            "care îi este atribuit, cu un termen concret.")
ws["A2"].font = Font(size=10, italic=True, color=GREY_TXT)

CONTACT_NAMES = f"'02_CONTACTE'!$C${hr+1}:$C${last_contact}"


# ---------------------------------------------------------------- 03_CHECKLIST_DATE

CHK_COLS = [
    ("ID", 7), ("Set de date", 15), ("Foaie destinație", 15), ("Câmp (tehnic)", 20),
    ("Denumire RO", 22), ("Ce înseamnă / cum se completează", 46), ("Obligatoriu", 13),
    ("Format / unitate", 15), ("Exemplu", 20), ("Sursă tipică", 18),
    ("Rol furnizor", 20), ("Persoana (nume)", 20), ("Ce deblochează", 28),
    ("Status", 14), ("Data primirii", 13), ("Observații", 26),
]

ws = wb.create_sheet("03_CHECKLIST_DATE")
title_block(ws, "03 — CE DATE CER ȘI DE LA CINE",
            "Lista completă a câmpurilor. Filtrează pe Obligatoriu = „Obligatoriu” pentru setul minim viabil. "
            "Coloana „Persoana” se leagă de numele din 02_CONTACTE.")
hr = 4
for j, (name, w) in enumerate(CHK_COLS, start=1):
    ws.cell(row=hr, column=j, value=name)
    ws.column_dimensions[get_column_letter(j)].width = w
style_header(ws, hr, len(CHK_COLS))
ws.freeze_panes = f"E{hr+1}"

SET_LABEL = {
    SH_CLIENTI: "Clienți", SH_PRODUSE: "Produse", SH_OAMENI: "Oameni",
    SH_VANZARI: "Vânzări", SH_TINTE: "Ținte / buget", SH_NPS: "NPS / feedback",
}

chk_rows = []
i = 1
# randuri de context (non-camp)
CONTEXT_ROWS = [
    ("Profil firmă", "01_FIRMA", "Profil complet firmă", "Profil firmă",
     "Se completează în discuția de kickoff: identificare, dimensiune, sisteme, obiective.",
     "Obligatoriu", "Text", "vezi 01_FIRMA", "Interviu / Workshop", "CEO / Administrator",
     "Context pentru toate interpretările"),
    ("Documente / context", "—", "NDA semnat", "Documente / context",
     "Acord de confidențialitate semnat înainte de transmiterea oricăror date.",
     "Obligatoriu", "PDF", "NDA_2026.pdf", "Interviu / Workshop", "CEO / Administrator",
     "Permite transmiterea legală a datelor"),
    ("Documente / context", "—", "Organigramă echipă comercială", "Documente / context",
     "Structura echipei: cine cui raportează, cine acoperă ce zonă sau segment.",
     "Recomandat", "PDF / imagine", "organigrama.pdf", "HR / State de plată", "HR / Resurse Umane",
     "Interpretarea corectă a analizei pe oameni"),
    ("Documente / context", "—", "Listă de prețuri și politică de discount", "Documente / context",
     "Prețurile de listă și regulile de discount pe segment sau volum.",
     "Recomandat", "XLSX", "pricelist_2026.xlsx", "Excel intern", "Director Vânzări",
     "Analiza erodării de preț și a marjei"),
    ("Financiar", "—", "Balanță / P&L ultimele 24 de luni", "Financiar",
     "Pentru validarea veniturilor și a costurilor față de contabilitate.",
     "Recomandat", "XLSX / PDF", "balanta_2025.xlsx", "Balanță / Contabilitate", "CFO / Contabilitate",
     "Verificarea consistenței cifrelor"),
    ("Financiar", "—", "Cheltuieli fixe lunare (agregat)", "Financiar",
     "Costuri indirecte lunare, agregat. Nu e nevoie de detaliu pe furnizor.",
     "Opțional", "Valoare / lună", "180000", "Balanță / Contabilitate", "CFO / Contabilitate",
     "Prag de rentabilitate, impact real al oportunităților"),
    ("Documente / context", "—", "Acces la export ERP (recurent)", "Documente / context",
     "Cine face exportul lunar și în ce format, pentru actualizări periodice.",
     "Recomandat", "Proces", "export lunar CSV", "ERP / Facturare", "IT / ERP",
     "Repetabilitatea raportului, nu doar o dată"),
]
for setd, foaie, camp, _s2, desc, oblig, fmt, ex, sursa, rol, unlocks in CONTEXT_ROWS:
    chk_rows.append((f"D{i:03d}", setd, foaie, "", camp, desc, oblig, fmt, ex, sursa, rol, unlocks))
    i += 1

for sheet in (SH_VANZARI, SH_CLIENTI, SH_PRODUSE, SH_TINTE, SH_NPS, SH_OAMENI):
    for c in SCHEMA[sheet]["cols"]:
        if c["kind"] == "calc":
            continue
        chk_rows.append((f"D{i:03d}", SET_LABEL[sheet], sheet, c["tech"], c["ro"], c["desc"],
                         c["oblig"], c["fmt"] or "Text", str(c["example"]), c["sursa"],
                         c["rol"], c["unlocks"]))
        i += 1

r = hr + 1
for row in chk_rows:
    vals = list(row[:6]) + [row[6], row[7], row[8], row[9], row[10], "", row[11], "Nesolicitat", None, ""]
    # reordoneaza: ID, Set, Foaie, Camp, RO, Desc, Oblig, Format, Exemplu, Sursa, Rol, Persoana, Deblocheaza, Status, Data, Obs
    vals = [row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8],
            row[9], row[10], "", row[11], "Nesolicitat", None, ""]
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.border = BORDER
        c.font = Font(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if j in (12, 14, 15, 16):
            c.fill = PatternFill("solid", fgColor=INPUT_FILL)
    ws.cell(row=r, column=15).number_format = F_DATE
    ws.row_dimensions[r].height = 28
    r += 1
last_chk = r - 1

add_dv(ws, lista_ref("Status_date"), f"N{hr+1}:N{last_chk}")
add_dv(ws, CONTACT_NAMES, f"L{hr+1}:L{last_chk}")
ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(CHK_COLS))}{last_chk}"
ws.conditional_formatting.add(
    f"G{hr+1}:G{last_chk}",
    CellIsRule(operator="equal", formula=['"Obligatoriu"'],
               fill=PatternFill("solid", fgColor="FBD7D7"),
               font=Font(bold=True, color="9C2A2A")))
ws.conditional_formatting.add(
    f"N{hr+1}:N{last_chk}",
    CellIsRule(operator="equal", formula=['"Primit"'],
               fill=PatternFill("solid", fgColor="D6F0D6")))
ws.conditional_formatting.add(
    f"N{hr+1}:N{last_chk}",
    CellIsRule(operator="equal", formula=['"Validat"'],
               fill=PatternFill("solid", fgColor="BCE3BC")))
ws.conditional_formatting.add(
    f"N{hr+1}:N{last_chk}",
    CellIsRule(operator="equal", formula=['"Blocat"'],
               fill=PatternFill("solid", fgColor="F8D7DA")))

# mini-dashboard de progres al colectarii
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
prog_col = len(CHK_COLS) + 2
PC = get_column_letter(prog_col)
PC2 = get_column_letter(prog_col + 1)
ws.column_dimensions[PC].width = 30
ws.column_dimensions[PC2].width = 12
ws.cell(row=hr, column=prog_col, value="PROGRES COLECTARE").font = Font(bold=True, color=WHITE, size=10)
ws.cell(row=hr, column=prog_col).fill = PatternFill("solid", fgColor=INK)
ws.cell(row=hr, column=prog_col + 1).fill = PatternFill("solid", fgColor=INK)
prog = [
    ("Total câmpuri cerute", f"=COUNTA($A${hr+1}:$A${last_chk})", F_INT),
    ("Din care obligatorii", f'=COUNTIF($G${hr+1}:$G${last_chk},"Obligatoriu")', F_INT),
    ("Primite sau validate", f'=COUNTIF($N${hr+1}:$N${last_chk},"Primit")+COUNTIF($N${hr+1}:$N${last_chk},"Validat")', F_INT),
    ("Obligatorii încă lipsă", f'=COUNTIFS($G${hr+1}:$G${last_chk},"Obligatoriu",$N${hr+1}:$N${last_chk},"<>Primit",$N${hr+1}:$N${last_chk},"<>Validat")', F_INT),
    ("Blocate", f'=COUNTIF($N${hr+1}:$N${last_chk},"Blocat")', F_INT),
    ("% completare", f'=IFERROR(({PC2}{hr+3})/({PC2}{hr+1}),0)', F_PCT),
    ("% completare obligatorii", f'=IFERROR(1-({PC2}{hr+4})/({PC2}{hr+2}),0)', F_PCT),
]
for k, (lab, f, fmt) in enumerate(prog, start=1):
    ws.cell(row=hr + k, column=prog_col, value=lab).font = Font(size=10, bold=True, color=INK)
    c = ws.cell(row=hr + k, column=prog_col + 1, value=f)
    c.number_format = fmt
    c.fill = PatternFill("solid", fgColor=CALC_FILL)
    c.border = BORDER
    c.font = Font(size=10, color=GREY_TXT)



# ---------------------------------------------------------------- 04_DICTIONAR

DICT_COLS = [("Termen", 30), ("Categorie", 13), ("Ce înseamnă", 62),
             ("De ce contează pentru analiză", 56), ("Cum îl ceri / unde îl găsești", 56)]

ws = wb.create_sheet("04_DICTIONAR")
title_block(ws, "04 — DICȚIONAR DE TERMENI",
            "Fiecare termen din fișier, explicat pe scurt, cu formularea pe care o poți folosi direct în discuția cu clientul. "
            "Filtrează pe „Categorie” ca să vezi doar ce te interesează.")
hr = 4
for j, (h, w) in enumerate(DICT_COLS, start=1):
    ws.cell(row=hr, column=j, value=h)
    ws.column_dimensions[get_column_letter(j)].width = w
style_header(ws, hr, len(DICT_COLS))
ws.freeze_panes = f"A{hr+1}"

CAT_COLORS = {"Proces": "E4EDF6", "Sisteme": "E8F1EC", "Financiar": "FBF0DC",
              "Comercial": "F3E9F2", "Oameni": "EDEFF5", "Măsurare": "E9F2F2",
              "Excel": "EFEFEC"}
r = hr + 1
for termen, cat, ce, dece, cum in DICTIONAR:
    vals = (termen, cat, ce, dece, cum)
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(size=10, bold=(j == 1), color=INK if j == 1 else "333333")
        if j == 2:
            c.fill = PatternFill("solid", fgColor=CAT_COLORS.get(cat, "EFEFEC"))
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            c.font = Font(size=9, color=GREY_TXT)
    ws.row_dimensions[r].height = 46
    r += 1
last_dict = r - 1
ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(DICT_COLS))}{last_dict}"
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- 09_PARAMETRI

ws = wb.create_sheet(SH_PAR)
title_block(ws, "09 — PARAMETRI (panoul de control)",
            "Aici se setează perioada și pragurile. Tot restul fișierului se recalculează automat.")
ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 66

P = {}


def par_section(r, text):
    ws.cell(row=r, column=1, value=text).font = Font(bold=True, size=11, color=WHITE)
    for k in (1, 2, 3):
        ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=TEAL)
    ws.row_dimensions[r].height = 20


import datetime as _dt

par_section(4, "PERIOADA DE ANALIZĂ")
label_value(ws, 5, "Perioada curentă — de la", _dt.date(2026, 1, 1), "Prima zi a perioadei analizate.", F_DATE)
label_value(ws, 6, "Perioada curentă — până la", _dt.date(2026, 6, 30), "Ultima zi a perioadei analizate.", F_DATE)
label_value(ws, 7, "Perioada de comparație — de la", _dt.date(2025, 1, 1), "De regulă aceeași perioadă din anul anterior.", F_DATE)
label_value(ws, 8, "Perioada de comparație — până la", _dt.date(2025, 6, 30), "", F_DATE)
label_value(ws, 9, "Data de referință", "=IF($B$6=\"\",TODAY(),$B$6)",
            "Față de ea se numără zilele de inactivitate. Implicit = ultima zi a perioadei analizate, "
            "ca să nu apară dormanți artificiali când analizezi o perioadă din trecut. "
            "Poți scrie în loc =TODAY() sau o dată fixă.", F_DATE, input_cell=False)
label_value(ws, 10, "Moneda de raportare", "RON", "Toate valorile trebuie să fie în aceeași monedă.")

par_section(12, "PRAGURI PENTRU SEMNALE")
label_value(ws, 13, "Scădere de venit considerată alertă", -0.15, "Ex. −15%: sub acest nivel, clientul intră în „Scădere venit”.", F_PCT)
label_value(ws, 14, "Creștere considerată oportunitate", 0.20, "Peste acest nivel, produsul/clientul e marcat ca „în creștere”.", F_PCT)
label_value(ws, 15, "Marjă brută de referință", 0.30, "Sub acest nivel se semnalează problemă de marjă.", F_PCT)
label_value(ws, 16, "Zile fără achiziție = client dormant", 90, "Numărul de zile de la ultima achiziție.", F_INT)
label_value(ws, 17, "Prag NPS acceptabil", 30, "NPS sub acest nivel declanșează plan de recuperare.", F_INT)
label_value(ws, 18, "Prag de atingere a țintei", 0.90, "Sub acest nivel se semnalează gap față de țintă.", F_PCT)
label_value(ws, 19, "Concentrare: prag pe un singur client", 0.20, "Peste acest procent din venit, clientul e risc de concentrare.", F_PCT)
label_value(ws, 20, "Venit minim pentru a intra în analiză", 0, "Ignoră clienții sub acest venit, ca să nu se aglomereze raportul.", F_MONEY)

par_section(22, "PRAGURI DE PRIORITIZARE A OPORTUNITĂȚILOR")
label_value(ws, 23, "Scor minim pentru prioritate HIGH", 40000, "Scor = Impact × Probabilitate × Ușurință.", F_MONEY)
label_value(ws, 24, "Scor minim pentru prioritate MEDIUM", 20000, "", F_MONEY)

par_section(26, "META")
# referința se calculează din poziția reală a câmpului, ca să nu se strice
# când se adaugă sau se scoate un rând din profilul firmei
_rand_nume = rand_firma["Denumire legală"]
label_value(ws, 27, "Client analizat",
            f'=IF(\'01_FIRMA\'!B{_rand_nume}="","[completează 01_FIRMA]",\'01_FIRMA\'!B{_rand_nume})',
            "Se preia din 01_FIRMA.", None, input_cell=False)
label_value(ws, 28, "Consultant EQUIL", "", "")
label_value(ws, 29, "Versiune fișier", "v0.2", "", None, input_cell=False)
label_value(ws, 30, "Data generării raportului", "", "", F_DATE)
add_dv(ws, lista_ref("Moneda"), "B10")

p = lambda row: f"'{SH_PAR}'!$B${row}"
P = dict(cur_from=p(5), cur_to=p(6), cmp_from=p(7), cmp_to=p(8), today=p(9),
         currency=p(10), decline=p(13), growth=p(14), margin=p(15), dormant=p(16),
         nps=p(17), target=p(18), conc=p(19), minrev=p(20), high=p(23), med=p(24))


# ------------------------------------------------- constructor foi de date

def build_data_sheet(name, formulas=None, dv=None, seed=None):
    spec = SCHEMA[name]
    cols = spec["cols"]
    n = spec["rows"]
    ws = wb.create_sheet(name)
    title_block(ws, f"{name.split('_')[0]} — {spec['label'].upper()}",
                "Galben = se completează. Gri = calculat automat, nu suprascrie.")
    hr = 4
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=hr, column=j, value=f"{c['ro']}\n{c['tech']}")
        cell.comment = Comment(f"{c['ro']} ({c['tech']})\n\n{c['desc']}\n\n"
                               f"Obligatoriu: {c['oblig']}\nSursă tipică: {c['sursa']}\n"
                               f"Cine îl poate da: {c['rol']}\nExemplu: {c['example']}", "EQUIL")
        ws.column_dimensions[get_column_letter(j)].width = c["width"]
    style_header(ws, hr, len(cols), height=34)
    first = hr + 1
    last = hr + n
    for j, c in enumerate(cols, start=1):
        L = get_column_letter(j)
        fill = CALC_FILL if c["kind"] == "calc" else INPUT_FILL
        for rr in range(first, last + 1):
            cell = ws.cell(row=rr, column=j)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = BORDER
            cell.font = Font(size=10, color=GREY_TXT if c["kind"] == "calc" else "000000")
            if c["fmt"]:
                cell.number_format = c["fmt"]
        if formulas and c["tech"] in formulas:
            tmpl = formulas[c["tech"]]
            for rr in range(first, last + 1):
                ws.cell(row=rr, column=j, value=tmpl.format(r=rr))
    if seed:
        for k, row in enumerate(seed):
            for j, v in enumerate(row, start=1):
                if v is not None:
                    ws.cell(row=first + k, column=j, value=v)
    ws.freeze_panes = f"C{first}"
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{last}"
    if dv:
        for tech, lista in dv.items():
            L = cl(name, tech)
            add_dv(ws, lista_ref(lista), f"{L}{first}:{L}{last}")
    ws.sheet_view.showGridLines = False
    return ws, first, last


# 10_CLIENTI
build_data_sheet(SH_CLIENTI,
                 dv={"Segment": "Segment_client", "Channel": "Canal", "Status": "Status_generic"},
                 seed=[["C001", "Alfa Distribution SRL", "Key Account", "Wholesale", "București-Ilfov",
                        "RO", "S001", _dt.date(2021, 3, 15), "Activ", 30, 50000, "Exemplu — șterge înainte de utilizare"],
                       ["C002", "Beta Retail SA", "Mid-market", "Retail", "Cluj",
                        "RO", "S002", _dt.date(2023, 6, 1), "Activ", 15, 20000, "Exemplu — șterge înainte de utilizare"]])

# 11_PRODUSE
build_data_sheet(SH_PRODUSE,
                 dv={"Is_Recurring": "Da_Nu", "Status": "Status_generic"},
                 seed=[["P001", "Produs X", "Categoria A", "", 120, 84, "Nu", None, "Activ"],
                       ["P002", "Serviciu Y", "Servicii recurente", "", 1200, 600, "Da", None, "Activ"]])

# 12_OAMENI
build_data_sheet(SH_OAMENI,
                 dv={"Is_Quota_Carrying": "Da_Nu"},
                 seed=[["S001", "Ana P.", "Key Account Manager", "Vânzări B2B", 1, _dt.date(2023, 2, 1), None, 96000, "Da"],
                       ["S002", "Bogdan M.", "Sales Representative", "Vânzări B2B", 1, _dt.date(2024, 9, 1), None, 72000, "Da"]])

# 13_VANZARI
cust_id_r = R(SH_CLIENTI, "Customer_ID")
cust_nm_r = R(SH_CLIENTI, "Customer_Name")
prod_id_r = R(SH_PRODUSE, "Product_ID")
prod_nm_r = R(SH_PRODUSE, "Product_Name")

A = cl(SH_VANZARI, "Doc_Date")
CUS = cl(SH_VANZARI, "Customer_ID")
PRD = cl(SH_VANZARI, "Product_ID")
REV = cl(SH_VANZARI, "Revenue_Net")
CGS = cl(SH_VANZARI, "COGS")
GPc = cl(SH_VANZARI, "Gross_Profit")

vanzari_formulas = {
    "Period": f'=IF(${A}{{r}}="","",{ym(f"${A}{{r}}")})',
    "Gross_Profit": f'=IF(${REV}{{r}}="","",${REV}{{r}}-N(${CGS}{{r}}))',
    "Margin_Pct": f'=IFERROR(${GPc}{{r}}/${REV}{{r}},"")',
    "Customer_Name": f'=IFERROR(INDEX({cust_nm_r},MATCH(${CUS}{{r}},{cust_id_r},0)),"")',
    "Product_Name": f'=IFERROR(INDEX({prod_nm_r},MATCH(${PRD}{{r}},{prod_id_r},0)),"")',
}
build_data_sheet(SH_VANZARI, formulas=vanzari_formulas,
                 dv={"Doc_Type": "Tip_document", "Currency": "Moneda"},
                 seed=[[_dt.date(2026, 1, 31), None, "Factură", "C001", "P001", "S001", 100, 10000, 0, 7500,
                        None, None, None, None, "RON", "Exemplu — șterge înainte de utilizare"],
                       [_dt.date(2026, 1, 31), None, "Factură", "C002", "P002", "S002", 10, 12000, 500, 6000,
                        None, None, None, None, "RON", "Exemplu — șterge înainte de utilizare"]])

# 14_TINTE
build_data_sheet(SH_TINTE,
                 dv={"Scope": "Nivel_tinta", "Metric": "Metrica_tinta"},
                 seed=[[_dt.date(2026, 1, 31), "Companie", "", "Venit", 120000, "Exemplu — șterge înainte de utilizare"]])

# 15_NPS
SC = cl(SH_NPS, "Score")
nps_formulas = {
    "NPS_Category": f'=IF(${SC}{{r}}="","",IF(${SC}{{r}}>=9,"Promotor",IF(${SC}{{r}}>=7,"Pasiv","Detractor")))',
}
build_data_sheet(SH_NPS, formulas=nps_formulas,
                 dv={"Theme": "Tema_NPS", "Follow_Up_Status": "Status_followup"},
                 seed=[["R0001", _dt.date(2026, 2, 10), "C001", "Mihai I.", "Decident", 9, None,
                        "Relație bună, livrări la timp", "Relație", "Email", "2026-T1", "S001", "Neînceput"],
                       ["R0002", _dt.date(2026, 2, 12), "C002", "Ioana R.", "Utilizator", 5, None,
                        "Întârzieri repetate la livrare", "Livrare", "Email", "2026-T1", "S002", "În lucru"]])


# ---------------------------------------------------------------- refs utile

TX_DATE = R(SH_VANZARI, "Doc_Date")
TX_CUST = R(SH_VANZARI, "Customer_ID")
TX_PROD = R(SH_VANZARI, "Product_ID")
TX_PERS = R(SH_VANZARI, "Person_ID")
TX_REV = R(SH_VANZARI, "Revenue_Net")
TX_GP = R(SH_VANZARI, "Gross_Profit")
TX_UNITS = R(SH_VANZARI, "Units")
TX_DISC = R(SH_VANZARI, "Discount_Value")

CU_ID = R(SH_CLIENTI, "Customer_ID")
CU_NM = R(SH_CLIENTI, "Customer_Name")
CU_SEG = R(SH_CLIENTI, "Segment")
CU_OWN = R(SH_CLIENTI, "Account_Owner_ID")

PR_ID = R(SH_PRODUSE, "Product_ID")
PR_NM = R(SH_PRODUSE, "Product_Name")
PR_CAT = R(SH_PRODUSE, "Category")

PE_ID = R(SH_OAMENI, "Person_ID")
PE_NM = R(SH_OAMENI, "Person_Name")
PE_ROLE = R(SH_OAMENI, "Role")
PE_TEAM = R(SH_OAMENI, "Team")
PE_FTE = R(SH_OAMENI, "FTE")
PE_START = R(SH_OAMENI, "Start_Date")
PE_EXIT = R(SH_OAMENI, "Exit_Date")

TG_DATE = R(SH_TINTE, "Period_Date")
TG_SCOPE = R(SH_TINTE, "Scope")
TG_SID = R(SH_TINTE, "Scope_ID")
TG_METRIC = R(SH_TINTE, "Metric")
TG_VAL = R(SH_TINTE, "Target_Value")

NP_DATE = R(SH_NPS, "Survey_Date")
NP_CUST = R(SH_NPS, "Customer_ID")
NP_SCORE = R(SH_NPS, "Score")


def win(a, b):
    return f'{TX_DATE},">="&{a},{TX_DATE},"<="&{b}'


def sumifs(val, a, b, extra=""):
    return f'SUMIFS({val},{win(a,b)}{extra})'


# ------------------------------------------------- 16 / 17 — piata si perceptie
_cf, _ct = P["cur_from"], P["cur_to"]
_nps_n = f'SUMPRODUCT(({NP_DATE}>={_cf})*({NP_DATE}<={_ct})*({NP_SCORE}<>""))'
_nps_scor = (f'IFERROR((COUNTIFS({NP_DATE},">="&{_cf},{NP_DATE},"<="&{_ct},{NP_SCORE},">=9")'
             f'-COUNTIFS({NP_DATE},">="&{_cf},{NP_DATE},"<="&{_ct},{NP_SCORE},"<=6"))'
             f'/({_nps_n})*100,"")')
CONC = adauga_concurenta(wb, "16_CONCURENTA", n_randuri=25, light=False)
PERC = adauga_perceptie(wb, "17_CE_CRED_OAMENII", CONC,
                        nps_scor=_nps_scor, nps_nr=_nps_n, light=False)

# ---------------------------------------------------------------- 20_KPI

ws = wb.create_sheet("20_KPI")
title_block(ws, "20 — KPI (perioada curentă vs. perioada de comparație)",
            "Toate valorile se recalculează din 09_PARAMETRI. Nu se scrie nimic în această foaie.")
ws["A2"].alignment = Alignment(wrap_text=True)
hdr = ["KPI", "Perioada curentă", "Perioada comparativă", "Δ", "Δ %", "Ce înseamnă"]
widths = [34, 17, 19, 15, 11, 60]
hr = 4
for j, (h, w) in enumerate(zip(hdr, widths), start=1):
    ws.cell(row=hr, column=j, value=h)
    ws.column_dimensions[get_column_letter(j)].width = w
style_header(ws, hr, len(hdr))
ws.cell(row=3, column=2, value=f"={P['cur_from']}").number_format = F_DATE
ws.cell(row=3, column=3, value=f"={P['cmp_from']}").number_format = F_DATE
ws.cell(row=3, column=2).font = Font(size=9, italic=True, color=GREY_TXT)
ws.cell(row=3, column=3).font = Font(size=9, italic=True, color=GREY_TXT)

cf, ct, pf, pt = P["cur_from"], P["cur_to"], P["cmp_from"], P["cmp_to"]


def kpi_pair(fn):
    return "=" + fn(cf, ct), "=" + fn(pf, pt)


def f_rev(a, b): return sumifs(TX_REV, a, b)
def f_gp(a, b): return sumifs(TX_GP, a, b)
def f_units(a, b): return sumifs(TX_UNITS, a, b)
def f_disc(a, b): return sumifs(TX_DISC, a, b)
def f_lines(a, b): return f'SUMPRODUCT(({TX_DATE}>={a})*({TX_DATE}<={b})*({TX_REV}<>""))'
def f_active(a, b): return f'SUMPRODUCT(({CU_ID}<>"")*(SUMIFS({TX_REV},{TX_CUST},{CU_ID},{win(a,b)})<>0))'
def f_new(a, b): return (f'SUMPRODUCT(({CU_ID}<>"")*(SUMIFS({TX_REV},{TX_CUST},{CU_ID},{TX_DATE},"<"&{a})=0)'
                         f'*(SUMIFS({TX_REV},{TX_CUST},{CU_ID},{win(a,b)})<>0))')
def f_fte(a, b): return (f'SUMPRODUCT(({PE_ID}<>"")*({PE_START}<={b})*'
                             f'(({PE_EXIT}="")+({PE_EXIT}>={a}))*{PE_FTE})')
def f_nps_n(a, b): return f'SUMPRODUCT(({NP_DATE}>={a})*({NP_DATE}<={b})*({NP_SCORE}<>""))'
def f_prom(a, b): return f'COUNTIFS({NP_DATE},">="&{a},{NP_DATE},"<="&{b},{NP_SCORE},">=9")'
def f_detr(a, b): return f'COUNTIFS({NP_DATE},">="&{a},{NP_DATE},"<="&{b},{NP_SCORE},"<=6")'
def f_target(a, b): return (f'SUMIFS({TG_VAL},{TG_DATE},">="&{a},{TG_DATE},"<="&{b},'
                            f'{TG_METRIC},"Venit",{TG_SCOPE},"Companie")')
def f_top1(a, b): return f'SUMPRODUCT(MAX(({CU_ID}<>"")*SUMIFS({TX_REV},{TX_CUST},{CU_ID},{win(a,b)})))'


KPIS = [
    ("VENIT ȘI PROFIT", None, None, None, None),
    ("Venit net", f_rev, F_MONEY, "pct_ok", "Venitul total fără TVA din perioada selectată."),
    ("Profit brut", f_gp, F_MONEY, "pct_ok", "Venit net minus costurile directe."),
    ("Marjă brută %", None, F_PCT, "pp", "Profit brut / venit net. Compară cu pragul din 09_PARAMETRI."),
    ("Discount acordat", f_disc, F_MONEY, "pct_ok", "Cât s-a lăsat din preț. Creșterea lui erodează marja."),
    ("Unități vândute", f_units, F_INT, "pct_ok", "Separă creșterea de volum de creșterea de preț."),
    ("Preț mediu / unitate", None, F_NUM, "pct_ok", "Venit net / unități. Arată dacă prețul crește sau se erodează."),
    ("CLIENȚI", None, None, None, None),
    ("Clienți activi", f_active, F_INT, "pct_ok", "Clienți cu cel puțin o tranzacție în perioadă."),
    ("Clienți noi", f_new, F_INT, "pct_ok", "Clienți fără nicio achiziție înainte de perioada curentă."),
    ("Venit mediu / client", None, F_MONEY, "pct_ok", "Arată dacă creșterea vine din clienți noi sau din adâncirea relației."),
    ("Număr de linii / facturi", f_lines, F_INT, "pct_ok", "Frecvența comenzilor."),
    ("Valoare medie per linie", None, F_MONEY, "pct_ok", "Venit net / număr de linii."),
    ("Concentrare — cel mai mare client", None, F_PCT, "pp", "Ce procent din venit vine de la un singur client. Peste pragul setat = risc."),
    ("OAMENI", None, None, None, None),
    ("FTE activ", f_fte, F_NUM, "pct_ok", "Normă întreagă echivalentă, activă în perioadă."),
    ("Venit / FTE", None, F_MONEY, "pct_ok", "Indicatorul principal de productivitate."),
    ("Profit brut / FTE", None, F_MONEY, "pct_ok", "Productivitate ajustată cu marja."),
    ("EXPERIENȚA CLIENTULUI", None, None, None, None),
    ("NPS", None, F_PCT2, "abs", "(% promotori − % detractori) × 100. Scala: −100 la +100."),
    ("Număr răspunsuri NPS", f_nps_n, F_INT, "pct_ok", "Sub 20 de răspunsuri, NPS-ul e orientativ, nu concluziv."),
    ("% Promotori", None, F_PCT, "pp", "Scor 9-10."),
    ("% Detractori", None, F_PCT, "pp", "Scor 0-6. Aceștia intră în planul de recuperare."),
    ("PLAN", None, None, None, None),
    ("Țintă de venit", f_target, F_MONEY, "pct_ok", "Din 14_TINTE, nivel „Companie”, indicator „Venit”."),
    ("Atingere țintă %", None, F_PCT, "pp", "Venit realizat / țintă."),
    ("Gap față de țintă", None, F_MONEY, "pct_ok", "Cât mai lipsește. Acesta e punctul de plecare al oportunităților."),
]

r = hr + 1
ROW = {}
for name, fn, fmt, dtype, interp in KPIS:
    if fn is None and fmt is None:
        ws.cell(row=r, column=1, value=name).font = Font(bold=True, size=10, color=WHITE)
        for k in range(1, 7):
            ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=INK)
        ws.row_dimensions[r].height = 18
        r += 1
        continue
    ROW[name] = r
    ws.cell(row=r, column=1, value=name).font = Font(bold=True, size=10, color=INK)
    if fn is not None:
        b, c = kpi_pair(fn)
        ws.cell(row=r, column=2, value=b)
        ws.cell(row=r, column=3, value=c)
    ws.cell(row=r, column=6, value=interp).font = Font(size=9, color=GREY_TXT)
    ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="center")
    for k in (2, 3, 4, 5):
        cc = ws.cell(row=r, column=k)
        cc.fill = PatternFill("solid", fgColor=CALC_FILL)
        cc.border = BORDER
        cc.font = Font(size=10, bold=(k == 2), color=INK if k == 2 else GREY_TXT)
        cc.number_format = fmt if k in (2, 3) else (fmt if k == 4 else F_PCT)
    if dtype == "pp":
        ws.cell(row=r, column=5, value="—").alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5).number_format = "General"
    ws.row_dimensions[r].height = 22
    r += 1
last_kpi = r - 1

# formule derivate
def setf(name, col, formula, fmt=None):
    rr = ROW[name]
    c = ws.cell(row=rr, column=col, value=formula)
    if fmt:
        c.number_format = fmt

for col, (a, b) in ((2, (cf, ct)), (3, (pf, pt))):
    setf("Marjă brută %", col, f'=IFERROR({sumifs(TX_GP,a,b)}/{sumifs(TX_REV,a,b)},0)')
    setf("Preț mediu / unitate", col, f'=IFERROR({sumifs(TX_REV,a,b)}/{sumifs(TX_UNITS,a,b)},0)')
    setf("Venit mediu / client", col, f'=IFERROR({sumifs(TX_REV,a,b)}/({f_active(a,b)}),0)')
    setf("Valoare medie per linie", col, f'=IFERROR({sumifs(TX_REV,a,b)}/({f_lines(a,b)}),0)')
    setf("Concentrare — cel mai mare client", col, f'=IFERROR({f_top1(a,b)}/{sumifs(TX_REV,a,b)},0)')
    setf("Venit / FTE", col, f'=IFERROR({sumifs(TX_REV,a,b)}/({f_fte(a,b)}),0)')
    setf("Profit brut / FTE", col, f'=IFERROR({sumifs(TX_GP,a,b)}/({f_fte(a,b)}),0)')
    setf("NPS", col, f'=IFERROR((({f_prom(a,b)})-({f_detr(a,b)}))/({f_nps_n(a,b)})*100,"")')
    setf("% Promotori", col, f'=IFERROR(({f_prom(a,b)})/({f_nps_n(a,b)}),"")')
    setf("% Detractori", col, f'=IFERROR(({f_detr(a,b)})/({f_nps_n(a,b)}),"")')
    setf("Atingere țintă %", col, f'=IFERROR({sumifs(TX_REV,a,b)}/({f_target(a,b)}),"")')
    setf("Gap față de țintă", col, f'=IF(({f_target(a,b)})=0,"",({f_target(a,b)})-{sumifs(TX_REV,a,b)})')

for name, rr in ROW.items():
    ws.cell(row=rr, column=4, value=f'=IFERROR(IF(OR(B{rr}="",C{rr}=""),"",B{rr}-C{rr}),"")')
    if ws.cell(row=rr, column=5).value is None:
        ws.cell(row=rr, column=5, value=f'=IFERROR(IF(OR(C{rr}=0,C{rr}=""),"",(B{rr}-C{rr})/ABS(C{rr})),"")')

ws.conditional_formatting.add(f"D{hr+1}:D{last_kpi}",
                              CellIsRule(operator="lessThan", formula=["0"], font=Font(color="B02A37")))
ws.conditional_formatting.add(f"D{hr+1}:D{last_kpi}",
                              CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="1B7A3D")))
ws.conditional_formatting.add(f"E{hr+1}:E{last_kpi}",
                              CellIsRule(operator="lessThan", formula=["0"], font=Font(color="B02A37")))
ws.conditional_formatting.add(f"E{hr+1}:E{last_kpi}",
                              CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="1B7A3D")))
ws.freeze_panes = f"A{hr+1}"


# ------------------------------------------------- foi de analiza (semnale)

def analysis_sheet(name, title, subtitle, key_range, cols, n_rows, dim_first_row=5):
    ws = wb.create_sheet(name)
    title_block(ws, title, subtitle)
    hr = 4
    for j, (h, w, fmt, tmpl) in enumerate(cols, start=1):
        ws.cell(row=hr, column=j, value=h)
        ws.column_dimensions[get_column_letter(j)].width = w
    style_header(ws, hr, len(cols), height=36)
    first = hr + 1
    last = hr + n_rows
    for rr in range(first, last + 1):
        src = dim_first_row + (rr - first)
        for j, (h, w, fmt, tmpl) in enumerate(cols, start=1):
            c = ws.cell(row=rr, column=j, value=tmpl.format(r=rr, s=src))
            c.fill = PatternFill("solid", fgColor=CALC_FILL)
            c.border = BORDER
            c.font = Font(size=10, color=GREY_TXT)
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = f"C{first}"
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{last}"
    ws.sheet_view.showGridLines = False
    return ws, first, last


# --- 21_ANALIZA_CLIENTI
cid = f"'{SH_CLIENTI}'!$A{{s}}"
c_cols = [
    ("ID client", 12, None, f'=IF({cid}="","",{cid})'),
    ("Denumire client", 24, None, f'=IF($A{{r}}="","",IFERROR(INDEX({CU_NM},MATCH($A{{r}},{CU_ID},0)),""))'),
    ("Segment", 13, None, f'=IF($A{{r}}="","",IFERROR(INDEX({CU_SEG},MATCH($A{{r}},{CU_ID},0)),""))'),
    ("Responsabil cont", 13, None, f'=IF($A{{r}}="","",IFERROR(INDEX({CU_OWN},MATCH($A{{r}},{CU_ID},0)),""))'),
    ("Venit — perioada curentă", 15, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_REV},{TX_CUST},$A{{r}},{win(cf,ct)}))'),
    ("Venit — comparativ", 15, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_REV},{TX_CUST},$A{{r}},{win(pf,pt)}))'),
    ("Δ venit %", 11, F_PCT, '=IF($A{r}="","",IFERROR(IF($F{r}=0,"",($E{r}-$F{r})/ABS($F{r})),""))'),
    ("Profit brut", 13, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_GP},{TX_CUST},$A{{r}},{win(cf,ct)}))'),
    ("Marjă %", 10, F_PCT, '=IF($A{r}="","",IFERROR($H{r}/$E{r},""))'),
    ("% din venit total", 12, F_PCT, f'=IF($A{{r}}="","",IFERROR($E{{r}}/{sumifs(TX_REV,cf,ct)},""))'),
    ("Ultima achiziție", 13, F_DATE, f'=IF($A{{r}}="","",IFERROR(IF(SUMPRODUCT(MAX(({TX_CUST}=$A{{r}})*({TX_REV}<>"")*{TX_DATE}))=0,"",SUMPRODUCT(MAX(({TX_CUST}=$A{{r}})*({TX_REV}<>"")*{TX_DATE}))),""))'),
    ("Zile de la ultima achiziție", 13, F_INT, f'=IF(OR($A{{r}}="",$K{{r}}=""),"",{P["today"]}-$K{{r}})'),
    ("Nr. produse cumpărate", 12, F_INT, f'=IF($A{{r}}="","",SUMPRODUCT(({PR_ID}<>"")*(SUMIFS({TX_REV},{TX_CUST},$A{{r}},{TX_PROD},{PR_ID},{win(cf,ct)})<>0)))'),
    ("NPS mediu", 11, F_PCT2, f'=IF($A{{r}}="","",IFERROR(AVERAGEIFS({NP_SCORE},{NP_CUST},$A{{r}},{NP_DATE},">="&{cf},{NP_DATE},"<="&{ct}),""))'),
    ("Semnal", 22, None,
     '=IF($A{r}="","",'
     f'IF(AND($F{{r}}>0,$E{{r}}=0),"Client pierdut",'
     f'IF(AND($L{{r}}<>"",$L{{r}}>{P["dormant"]},$F{{r}}>0),"Client dormant",'
     f'IF(AND($G{{r}}<>"",$G{{r}}<={P["decline"]}),"Scădere de venit",'
     f'IF(AND($N{{r}}<>"",$N{{r}}<7),"Detractor activ",'
     f'IF(AND($E{{r}}>0,$J{{r}}>{P["conc"]}),"Risc concentrare",'
     f'IF(AND($E{{r}}>0,$I{{r}}<>"",$I{{r}}<{P["margin"]}),"Marjă sub referință",'
     f'IF(AND($G{{r}}<>"",$G{{r}}>={P["growth"]}),"În creștere",'
     'IF($E{r}=0,"Fără activitate","Stabil")))))))))'),
    ("Oportunitate sugerată", 24, None,
     '=IF($A{r}="","",'
     'IF($O{r}="Risc concentrare","Risc concentrare",'
     'IF($O{r}="Client pierdut","Recuperare client",'
     'IF($O{r}="Client dormant","Reactivare client",'
     'IF($O{r}="Scădere de venit","Recuperare client",'
     'IF($O{r}="Detractor activ","Recuperare NPS",'
     'IF($O{r}="Marjă sub referință","Marjă / preț",'
     'IF($O{r}="În creștere","Up-sell",'
     'IF($O{r}="Fără activitate","Reactivare client",'
     '"Cross-sell")))))))))'),
    ("Potențial estimat", 13, F_MONEY,
     '=IF($A{r}="","",'
     'IFERROR(IF($O{r}="Scădere de venit",MAX(0,$F{r}-$E{r}),'
     'IF(OR($O{r}="Client pierdut",$O{r}="Client dormant",$O{r}="Fără activitate"),$F{r},'
     f'IF($O{{r}}="Marjă sub referință",MAX(0,$E{{r}}*({P["margin"]}-$I{{r}})),'
     '$E{r}*0.15))),""))'),
]
c_cols.append(("Cheie clasament", 13, F_NUM,
               '=IF(OR($A{r}="",$E{r}="",N($E{r})<=0),"",$E{r}+ROW()/1000000)'))
wsA, fA, lA = analysis_sheet("21_ANALIZA_CLIENTI", "21 — ANALIZA CLIENȚILOR",
                             "Se completează singură din 10_CLIENTI + 13_VANZARI. Filtrează coloana „Semnal” ca să vezi unde e treabă de făcut.",
                             CU_ID, c_cols, N_CLIENTI)
KEY_A = get_column_letter(len(c_cols))
wsA.column_dimensions[KEY_A].hidden = True
for sig, color in (("Risc concentrare", "FDE2C8"), ("Client pierdut", "F8D7DA"),
                   ("Client dormant", "FBE3C2"), ("Scădere de venit", "F8D7DA"),
                   ("Detractor activ", "F8D7DA"), ("Marjă sub referință", "FFF3CD"),
                   ("În creștere", "D6F0D6")):
    wsA.conditional_formatting.add(f"O{fA}:O{lA}",
                                   CellIsRule(operator="equal", formula=[f'"{sig}"'],
                                              fill=PatternFill("solid", fgColor=color)))

# --- 22_ANALIZA_PRODUSE
pid = f"'{SH_PRODUSE}'!$A{{s}}"
p_cols = [
    ("ID produs", 12, None, f'=IF({pid}="","",{pid})'),
    ("Denumire produs", 26, None, f'=IF($A{{r}}="","",IFERROR(INDEX({PR_NM},MATCH($A{{r}},{PR_ID},0)),""))'),
    ("Categorie", 16, None, f'=IF($A{{r}}="","",IFERROR(INDEX({PR_CAT},MATCH($A{{r}},{PR_ID},0)),""))'),
    ("Venit — curent", 14, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_REV},{TX_PROD},$A{{r}},{win(cf,ct)}))'),
    ("Venit — comparativ", 14, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_REV},{TX_PROD},$A{{r}},{win(pf,pt)}))'),
    ("Δ venit %", 11, F_PCT, '=IF($A{r}="","",IFERROR(IF($E{r}=0,"",($D{r}-$E{r})/ABS($E{r})),""))'),
    ("Profit brut", 13, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_GP},{TX_PROD},$A{{r}},{win(cf,ct)}))'),
    ("Marjă %", 10, F_PCT, '=IF($A{r}="","",IFERROR($G{r}/$D{r},""))'),
    ("Unități", 11, F_INT, f'=IF($A{{r}}="","",SUMIFS({TX_UNITS},{TX_PROD},$A{{r}},{win(cf,ct)}))'),
    ("Preț mediu", 11, F_NUM, '=IF($A{r}="","",IFERROR($D{r}/$I{r},""))'),
    ("Nr. clienți care cumpără", 13, F_INT, f'=IF($A{{r}}="","",SUMPRODUCT(({CU_ID}<>"")*(SUMIFS({TX_REV},{TX_PROD},$A{{r}},{TX_CUST},{CU_ID},{win(cf,ct)})<>0)))'),
    ("Penetrare în baza de clienți", 13, F_PCT, f'=IF($A{{r}}="","",IFERROR($K{{r}}/({f_active(cf,ct)}),""))'),
    ("Semnal", 22, None,
     '=IF($A{r}="","",'
     f'IF(AND($D{{r}}>0,$H{{r}}<>"",$H{{r}}<{P["margin"]}),"Marjă sub referință",'
     f'IF(AND($F{{r}}<>"",$F{{r}}>={P["growth"]},$L{{r}}<0.3),"Creștere + penetrare mică",'
     f'IF(AND($F{{r}}<>"",$F{{r}}<={P["decline"]}),"Produs în scădere",'
     'IF($D{r}=0,"Fără vânzări","Stabil")))))'),
    ("Oportunitate sugerată", 24, None,
     '=IF($A{r}="","",'
     'IF($M{r}="Marjă sub referință","Marjă / preț",'
     'IF($M{r}="Creștere + penetrare mică","Creștere produs",'
     'IF($M{r}="Produs în scădere","Cross-sell",'
     'IF($M{r}="Fără vânzări","Proces / operațional",'
     '"Cross-sell")))))'),
    ("Potențial estimat", 13, F_MONEY,
     '=IF($A{r}="","",IFERROR('
     f'IF($M{{r}}="Marjă sub referință",MAX(0,$D{{r}}*({P["margin"]}-$H{{r}})),'
     'IF($M{r}="Creștere + penetrare mică",$D{r}*0.5,'
     'IF($M{r}="Produs în scădere",MAX(0,$E{r}-$D{r}),$D{r}*0.1))),""))'),
]
p_cols.append(("Cheie clasament", 13, F_NUM,
               '=IF(OR($A{r}="",$D{r}="",N($D{r})<=0),"",$D{r}+ROW()/1000000)'))
wsP, fP, lP = analysis_sheet("22_ANALIZA_PRODUSE", "22 — ANALIZA PRODUSELOR",
                             "Arată unde e marjă de recuperat și ce produse merg bine, dar ajung la prea puțini clienți.",
                             PR_ID, p_cols, N_PRODUSE)
KEY_P = get_column_letter(len(p_cols))
wsP.column_dimensions[KEY_P].hidden = True
for sig, color in (("Marjă sub referință", "FFF3CD"), ("Creștere + penetrare mică", "D6F0D6"),
                   ("Produs în scădere", "F8D7DA"), ("Fără vânzări", "EDEDED")):
    wsP.conditional_formatting.add(f"M{fP}:M{lP}",
                                   CellIsRule(operator="equal", formula=[f'"{sig}"'],
                                              fill=PatternFill("solid", fgColor=color)))

# --- 23_ANALIZA_OAMENI
sid = f"'{SH_OAMENI}'!$A{{s}}"
o_cols = [
    ("ID persoană", 12, None, f'=IF({sid}="","",{sid})'),
    ("Nume", 20, None, f'=IF($A{{r}}="","",IFERROR(INDEX({PE_NM},MATCH($A{{r}},{PE_ID},0)),""))'),
    ("Rol", 18, None, f'=IF($A{{r}}="","",IFERROR(INDEX({PE_ROLE},MATCH($A{{r}},{PE_ID},0)),""))'),
    ("Echipă", 15, None, f'=IF($A{{r}}="","",IFERROR(INDEX({PE_TEAM},MATCH($A{{r}},{PE_ID},0)),""))'),
    ("FTE", 8, F_NUM, f'=IF($A{{r}}="","",IFERROR(INDEX({PE_FTE},MATCH($A{{r}},{PE_ID},0)),""))'),
    ("Venit — curent", 14, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_REV},{TX_PERS},$A{{r}},{win(cf,ct)}))'),
    ("Venit — comparativ", 14, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_REV},{TX_PERS},$A{{r}},{win(pf,pt)}))'),
    ("Δ venit %", 11, F_PCT, '=IF($A{r}="","",IFERROR(IF($G{r}=0,"",($F{r}-$G{r})/ABS($G{r})),""))'),
    ("Profit brut", 13, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TX_GP},{TX_PERS},$A{{r}},{win(cf,ct)}))'),
    ("Marjă %", 10, F_PCT, '=IF($A{r}="","",IFERROR($I{r}/$F{r},""))'),
    ("Clienți activi", 11, F_INT, f'=IF($A{{r}}="","",SUMPRODUCT(({CU_ID}<>"")*(SUMIFS({TX_REV},{TX_PERS},$A{{r}},{TX_CUST},{CU_ID},{win(cf,ct)})<>0)))'),
    ("Venit / FTE", 13, F_MONEY, '=IF($A{r}="","",IFERROR($F{r}/$E{r},""))'),
    ("Țintă", 13, F_MONEY, f'=IF($A{{r}}="","",SUMIFS({TG_VAL},{TG_DATE},">="&{cf},{TG_DATE},"<="&{ct},{TG_METRIC},"Venit",{TG_SCOPE},"Persoană",{TG_SID},$A{{r}}))'),
    ("Atingere țintă %", 12, F_PCT, '=IF($A{r}="","",IFERROR(IF($M{r}=0,"",$F{r}/$M{r}),""))'),
    ("Semnal", 22, None,
     '=IF($A{r}="","",'
     f'IF(AND($N{{r}}<>"",$N{{r}}<{P["target"]}),"Sub țintă",'
     f'IF(AND($H{{r}}<>"",$H{{r}}<={P["decline"]}),"Portofoliu în scădere",'
     f'IF(AND($F{{r}}>0,$L{{r}}<>"",$L{{r}}<{f"AVERAGE($L${fA}:$L${lA})".replace("$L$","$L$")}*0.7),"Productivitate sub medie",'
     'IF($F{r}=0,"Fără vânzări înregistrate","În parametri")))))'),
    ("Oportunitate sugerată", 24, None,
     '=IF($A{r}="","",'
     'IF($O{r}="Sub țintă","Gap față de țintă",'
     'IF($O{r}="Portofoliu în scădere","Productivitate vânzări",'
     'IF($O{r}="Productivitate sub medie","Dezvoltare oameni",'
     'IF($O{r}="Fără vânzări înregistrate","Proces / operațional",'
     '"Productivitate vânzări")))))'),
    ("Potențial estimat", 13, F_MONEY,
     '=IF($A{r}="","",IFERROR('
     'IF($O{r}="Sub țintă",MAX(0,$M{r}-$F{r}),'
     'IF($O{r}="Portofoliu în scădere",MAX(0,$G{r}-$F{r}),'
     '$F{r}*0.15)),""))'),
]
# corectie: comparatia de productivitate se face pe media coloanei din aceeasi foaie
o_cols[14] = ("Semnal", 22, None,
              '=IF($A{r}="","",'
              f'IF(AND($N{{r}}<>"",$N{{r}}<{P["target"]},$M{{r}}>0),"Sub țintă",'
              f'IF(AND($H{{r}}<>"",$H{{r}}<={P["decline"]}),"Portofoliu în scădere",'
              f'IF(AND($F{{r}}>0,$L{{r}}<>"",$L{{r}}<IFERROR(AVERAGE($L$5:$L${4+N_OAMENI}),0)*0.7),"Productivitate sub medie",'
              'IF($F{r}=0,"Fără vânzări înregistrate","În parametri")))))')
wsO, fO, lO = analysis_sheet("23_ANALIZA_OAMENI", "23 — ANALIZA ECHIPEI",
                             "Productivitate și atingere de țintă pe om. Nu este un instrument de evaluare a performanței individuale, ci de identificat unde ajută coachingul.",
                             PE_ID, o_cols, N_OAMENI)
for sig, color in (("Sub țintă", "F8D7DA"), ("Portofoliu în scădere", "FBE3C2"),
                   ("Productivitate sub medie", "FFF3CD"), ("Fără vânzări înregistrate", "EDEDED")):
    wsO.conditional_formatting.add(f"O{fO}:O{lO}",
                                   CellIsRule(operator="equal", formula=[f'"{sig}"'],
                                              fill=PatternFill("solid", fgColor=color)))


# ---------------------------------------------------------------- 30_OPORTUNITATI

OPP_COLS = [
    ("ID", 8, None), ("Sursă", 13, None), ("Tip oportunitate", 20, None),
    ("Țintă (client / produs / persoană)", 22, None), ("Declanșator / dovadă din date", 40, None),
    ("Impact estimat", 13, F_MONEY), ("Probabilitate", 11, F_PCT), ("Ușurință", 10, F_PCT),
    ("Scor", 13, F_MONEY), ("Prioritate", 11, None), ("Acțiune recomandată", 42, None),
    ("Responsabil EQUIL", 16, None), ("Responsabil client", 16, None),
    ("Termen", 12, F_DATE), ("Status", 15, None), ("Valoare realizată", 13, F_MONEY),
    ("Observații", 26, None),
]
ws = wb.create_sheet("30_OPORTUNITATI")
title_block(ws, "30 — OPORTUNITĂȚI PRIORITIZATE",
            "Scor = Impact × Probabilitate × Ușurință. Prioritatea se calculează din pragurile din 09_PARAMETRI. "
            "Rândurile de mai jos sunt tipare; înlocuiește-le cu semnalele reale din foile 21-23.")
hr = 4
for j, (h, w, fmt) in enumerate(OPP_COLS, start=1):
    ws.cell(row=hr, column=j, value=h)
    ws.column_dimensions[get_column_letter(j)].width = w
style_header(ws, hr, len(OPP_COLS))
first_opp, last_opp = hr + 1, hr + N_OPP

OPP_SEED = [
    ("Manual", "Recuperare client", "", "Venit în scădere cu peste pragul setat față de perioada comparativă (vezi 21, coloana Semnal)", 42000, 0.8, 0.7, "Contact direct cu decidentul, diagnostic al volumului pierdut, ofertă de revenire"),
    ("Manual", "Cross-sell", "", "Client care cumpără produsul X, dar nu și produsul Y pe care îl iau clienți similari", 30000, 0.7, 0.9, "Propunere pachet, pe baza comportamentului clienților comparabili"),
    ("Manual", "Marjă / preț", "", "Venit stabil sau în creștere, dar marjă sub referința din 09_PARAMETRI (vezi 22)", 25000, 0.7, 0.5, "Revizuire preț, mix și cost; renegociere discount pe volum"),
    ("Manual", "Productivitate vânzări", "", "Venit per om semnificativ sub media echipei (vezi 23)", 20000, 0.6, 0.8, "Coaching pe pipeline, acoperire de cont și rată de conversie"),
    ("Manual", "Risc concentrare", "", "Un client depășește pragul de concentrare din venitul total", 50000, 0.8, 0.5, "Plan de retenție pe clientul mare + diversificare activă a bazei"),
    ("Manual", "Recuperare NPS", "", "Detractori identificați în 15_NPS, cu motive concrete", 15000, 0.6, 0.8, "Apel de recuperare în 10 zile, rezolvarea cauzei, remăsurare după 60 de zile"),
    ("Manual", "Creștere produs", "", "Produs în creștere, dar cu penetrare mică în baza de clienți (vezi 22)", 35000, 0.7, 0.7, "Campanie dedicată pe clienții care nu îl au încă"),
    ("Manual", "Gap față de țintă", "", "Atingerea țintei sub pragul din 09_PARAMETRI", 30000, 0.7, 0.7, "Prioritizarea pipeline-ului pe oportunitățile cu probabilitate mare"),
    ("Manual", "Reactivare client", "", "Client dormant: fără achiziții peste pragul de zile setat", 18000, 0.6, 0.9, "Campanie de reactivare cu ofertă de revenire și termen limitat"),
    ("Manual", "Dezvoltare oameni", "", "Diferență de performanță față de colegii comparabili", 22000, 0.6, 0.8, "Plan de training și shadowing pe cele mai slabe două etape din proces"),
]
for k, (src, tip, tinta, trig, imp, prob, ease, act) in enumerate(OPP_SEED):
    rr = first_opp + k
    ws.cell(row=rr, column=1, value=f"OPP{k+1:03d}")
    ws.cell(row=rr, column=2, value=src)
    ws.cell(row=rr, column=3, value=tip)
    ws.cell(row=rr, column=4, value=tinta)
    ws.cell(row=rr, column=5, value=trig)
    ws.cell(row=rr, column=6, value=imp)
    ws.cell(row=rr, column=7, value=prob)
    ws.cell(row=rr, column=8, value=ease)
    ws.cell(row=rr, column=11, value=act)
    ws.cell(row=rr, column=15, value="Identificată")

for rr in range(first_opp, last_opp + 1):
    for j, (h, w, fmt) in enumerate(OPP_COLS, start=1):
        c = ws.cell(row=rr, column=j)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(size=10)
        if fmt:
            c.number_format = fmt
        c.fill = PatternFill("solid", fgColor=CALC_FILL if j in (9, 10) else INPUT_FILL)
    ws.cell(row=rr, column=9, value=f'=IF($F{rr}="","",N($F{rr})*N($G{rr})*N($H{rr}))')
    ws.cell(row=rr, column=10,
            value=f'=IF($I{rr}="","",IF($I{rr}>={P["high"]},"HIGH",IF($I{rr}>={P["med"]},"MEDIUM","LOW")))')
    ws.cell(row=rr, column=9).font = Font(size=10, bold=True, color=INK)
    ws.row_dimensions[rr].height = 30

add_dv(ws, lista_ref("Tip_oportunitate"), f"C{first_opp}:C{last_opp}")
add_dv(ws, lista_ref("Status_oportunitate"), f"O{first_opp}:O{last_opp}")
ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(OPP_COLS))}{last_opp}"
ws.freeze_panes = f"C{first_opp}"
for pri, color, font_color in (("HIGH", "F8D7DA", "9C2A2A"), ("MEDIUM", "FFF3CD", "8A6D1F"), ("LOW", "EDEDED", "5B6770")):
    ws.conditional_formatting.add(f"J{first_opp}:J{last_opp}",
                                  CellIsRule(operator="equal", formula=[f'"{pri}"'],
                                             fill=PatternFill("solid", fgColor=color),
                                             font=Font(bold=True, color=font_color)))

# sumar oportunitati
sc = len(OPP_COLS) + 2
SL, SL2 = get_column_letter(sc), get_column_letter(sc + 1)
ws.column_dimensions[SL].width = 30
ws.column_dimensions[SL2].width = 14
hr_opp = hr
ws.cell(row=hr, column=sc, value="SUMAR PORTOFOLIU").font = Font(bold=True, color=WHITE, size=10)
ws.cell(row=hr, column=sc).fill = PatternFill("solid", fgColor=INK)
ws.cell(row=hr, column=sc + 1).fill = PatternFill("solid", fgColor=INK)
sumar = [
    ("Oportunități identificate", f"=COUNTA($A${first_opp}:$A${last_opp})", F_INT),
    ("Valoare brută (impact)", f"=SUM($F${first_opp}:$F${last_opp})", F_MONEY),
    ("Valoare ponderată (scor)", f"=SUM($I${first_opp}:$I${last_opp})", F_MONEY),
    ("Din care HIGH", f'=SUMIF($J${first_opp}:$J${last_opp},"HIGH",$I${first_opp}:$I${last_opp})', F_MONEY),
    ("Din care MEDIUM", f'=SUMIF($J${first_opp}:$J${last_opp},"MEDIUM",$I${first_opp}:$I${last_opp})', F_MONEY),
    ("Realizat până acum", f"=SUM($P${first_opp}:$P${last_opp})", F_MONEY),
    ("Acoperire gap față de țintă", f"=IFERROR(SUM($I${first_opp}:$I${last_opp})/'20_KPI'!$B${ROW['Gap față de țintă']},\"\")", F_PCT),
]
for k, (lab, f, fmt) in enumerate(sumar, start=1):
    ws.cell(row=hr + k, column=sc, value=lab).font = Font(size=10, bold=True, color=INK)
    c = ws.cell(row=hr + k, column=sc + 1, value=f)
    c.number_format = fmt
    c.fill = PatternFill("solid", fgColor=CALC_FILL)
    c.border = BORDER


# ---------------------------------------------------------------- 40_RAPORT

ws = wb.create_sheet("40_RAPORT")
title_block(ws, "EQUIL GROWTH INTELLIGENCE — RAPORT", None)
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 20
ws.column_dimensions["D"].width = 54


def rap_section(r, text):
    ws.cell(row=r, column=1, value=text).font = Font(bold=True, size=11, color=WHITE)
    for k in range(1, 5):
        ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=TEAL)
    ws.row_dimensions[r].height = 20


def rap_row(r, label, formula, fmt=None, note=None, calc=True):
    ws.cell(row=r, column=1, value=label).font = Font(bold=True, size=10, color=INK)
    c = ws.cell(row=r, column=2, value=formula)
    c.fill = PatternFill("solid", fgColor=CALC_FILL if calc else INPUT_FILL)
    c.border = BORDER
    c.font = Font(size=10, color=GREY_TXT if calc else "000000")
    if fmt:
        c.number_format = fmt
    if note:
        n = ws.cell(row=r, column=4, value=note)
        n.font = Font(size=9, italic=True, color=GREY_TXT)
        n.alignment = Alignment(wrap_text=True, vertical="center")
    return c


r = 4
rap_section(r, "IDENTIFICARE"); r += 1
rap_row(r, "Client", f"={P['currency'].replace('$B$10','$B$27')}", None, "Se preia din 01_FIRMA."); r += 1
rap_row(r, "Perioada analizată", f'={ymd(cf)}&" — "&{ymd(ct)}'); r += 1
rap_row(r, "Comparat cu", f'={ymd(pf)}&" — "&{ymd(pt)}'); r += 1
rap_row(r, "Monedă", f"={P['currency']}"); r += 1
rap_row(r, "Consultant EQUIL", f"='{SH_PAR}'!$B$28"); r += 1
r += 1


def K(name, col="B"):
    return f"'20_KPI'!${col}${ROW[name]}"


CUR = P["currency"]
_rev, _revd = K("Venit net"), K("Venit net", "E")
_mar = K("Marjă brută %")
_act, _new, _arpu = K("Clienți activi"), K("Clienți noi"), K("Venit mediu / client")
_conc = K("Concentrare — cel mai mare client")
_nps, _nresp, _det = K("NPS"), K("Număr răspunsuri NPS"), K("% Detractori")
_tgt, _ating, _gap = K("Țintă de venit"), K("Atingere țintă %"), K("Gap față de țintă")
_vfte = K("Venit / FTE")
_oppn = f"'30_OPORTUNITATI'!${SL2}${hr_opp+1}"
_oppv = f"'30_OPORTUNITATI'!${SL2}${hr_opp+3}"
_oppa = f"'30_OPORTUNITATI'!${SL2}${hr_opp+7}"

CONCLUZII = [
    ("Venit",
     f'=IF({_rev}=0,"Nu există date de vânzări în perioada selectată. Verifică 13_VANZARI și perioada din 09_PARAMETRI.",'
     f'"Venitul net al perioadei este "&{mon(_rev)}&" "&{CUR}&'
     f'IF(OR({_revd}="",{_revd}=0),", fără perioadă de comparație completată.",'
     f'", "&IF({_revd}>=0,"în creștere cu ","în scădere cu ")&{pct(f"ABS({_revd})")}&" față de perioada comparativă."))'),
    ("Marjă",
     f'=IF({_rev}=0,"—",'
     f'"Marja brută este "&{pct(_mar)}&", "&IF({_mar}>={P["margin"]},"peste","sub")&" referința de "&{pct(P["margin"])}&'
     f'IF({_mar}>={P["margin"]},". Marja se menține.",'
     f'". Un singur punct procentual de marjă recuperat înseamnă "&{mon(f"{_rev}*0.01")}&" "&{CUR}&"."))'),
    ("Clienți",
     f'=IF({_act}=0,"—",'
     f'"Au cumpărat "&{mon(_act)}&" clienți, din care "&{mon(_new)}&" noi. '
     f'Venitul mediu pe client este "&{mon(_arpu)}&" "&{CUR}&".")'),
    ("Concentrare",
     f'=IF({_rev}=0,"—",'
     f'"Cel mai mare client aduce "&{pct(_conc)}&" din venit, "&'
     f'IF({_conc}>{P["conc"]},"peste pragul de risc de "&{pct(P["conc"])}&". Retenția lui este prioritate, nu opțiune.",'
     f'"sub pragul de risc. Baza de clienți este echilibrată."))'),
    ("Productivitate",
     f'=IF({_vfte}=0,"—","Fiecare normă întreagă din echipă aduce "&{mon(_vfte)}&" "&{CUR}&" venit în perioada analizată.")'),
    ("NPS",
     f'=IF({_nresp}=0,"Nu există răspunsuri NPS în perioadă. Fără ele, nemulțumirea clienților se vede abia când scad comenzile.",'
     f'"NPS-ul este "&{mon(_nps)}&", din "&{mon(_nresp)}&" răspunsuri, cu "&{pct(_det,0)}&" detractori. "&'
     f'IF({_nps}<{P["nps"]},"Este sub pragul agreat de "&{mon(P["nps"])}&"; detractorii intră în planul de acțiuni.",'
     f'"Este peste pragul agreat.")&'
     f'IF({_nresp}<20," Atenție: sub 20 de răspunsuri, cifra este orientativă, nu concluzivă.",""))'),
    ("Țintă",
     f'=IF({_tgt}=0,"Nu este completată o țintă de venit pentru perioadă (14_TINTE, nivel „Companie”). Fără ea nu se poate măsura gap-ul.",'
     f'"Atingerea țintei este "&{pct(_ating)}&IF({_gap}>0,", mai lipsesc "&{mon(_gap)}&" "&{CUR}&" până la plan.",'
     f'", ținta este depășită."))'),
    ("Piață și reputație",
     f'=IF(AND({CONC["nr_concurenti"]}=0,{PERC["total_recenzii"]}=0),'
     f'"Nu e completat nimic despre concurență și reputație (foile 16 și 17). '
     f'E cel mai ieftin context din tot fișierul: o oră de căutat public, fără să întrebi pe nimeni.",'
     f'"Urmărim "&{mon(CONC["nr_concurenti"])}&" concurenți"&'
     f'IF({CONC["cota_grup"]}="","",", în care firma are "&{pct(CONC["cota_grup"])}&" din cifra grupului")&". "&'
     f'IF({CONC["nota_noastra"]}="","Firma nu are încă notă publică — verifică dacă are fișă Google Business.",'
     f'"Nota publică este "&FIXED({CONC["nota_noastra"]},1)&" din 5"&'
     f'IF({CONC["nota_concurenti"]}="",".",", față de "&FIXED({CONC["nota_concurenti"]},1)&" media concurenților"&'
     f'IF({CONC["nota_noastra"]}>{CONC["nota_concurenti"]}+0.2," — poziție mai bună decât a lor.",'
     f'IF({CONC["nota_noastra"]}<{CONC["nota_concurenti"]}-0.2," — poziție mai slabă; se rezolvă înainte de a crește volumul.",'
     f'" — aceeași poziție."))))&'
     f'IF({PERC["scor_perceptie"]}="",""," Scor de percepție: "&FIXED({PERC["scor_perceptie"]},0)&"/100 (încredere: "&{PERC["increderea"]}&")."))'),
    ("Oportunități",
     f'=IF({_oppn}=0,"Nu sunt încă oportunități completate în 30_OPORTUNITATI.",'
     f'"Sunt "&{mon(_oppn)}&" oportunități identificate, cu valoare ponderată de "&{mon(_oppv)}&" "&{CUR}&'
     f'IF(AND({_gap}>0,{_oppa}<>""),'
     f'IF({_oppa}>=1,", mai mult decât suficient cât să acopere gap-ul față de țintă ("&{pct(_oppa,0)}&" din el).",'
     f'", adică "&{pct(_oppa,0)}&" din gap-ul față de țintă."),"."))'),
]

rap_section(r, "CONCLUZII AUTOMATE (generate din date)"); r += 1
for lab, formula in CONCLUZII:
    ws.cell(row=r, column=1, value=lab).font = Font(bold=True, size=10, color=INK)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    c = ws.cell(row=r, column=2, value=formula)
    c.fill = PatternFill("solid", fgColor=CALC_FILL)
    c.border = BORDER
    c.font = Font(size=10, color="333333")
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 30
    r += 1
ws.cell(row=r, column=1, value="Cum se folosesc").font = Font(size=9, italic=True, color=GREY_TXT)
cnote = ws.cell(row=r, column=2, value="Sunt materie primă, nu text final: le citești, le verifici și scrii mai jos sinteza în limbajul clientului.")
cnote.font = Font(size=9, italic=True, color=GREY_TXT)
r += 2

rap_section(r, "SINTEZA EXECUTIVĂ"); r += 1
ws.cell(row=r, column=1, value="Concluzii (se scriu manual)").font = Font(bold=True, size=10, color=INK)
c = ws.cell(row=r, column=2, value="[3-5 concluzii, formulate în limbajul clientului, sprijinite de cifrele de mai jos]")
c.fill = PatternFill("solid", fgColor=INPUT_FILL)
c.alignment = Alignment(wrap_text=True, vertical="top")
c.border = BORDER
ws.row_dimensions[r].height = 54
ws.cell(row=r, column=4, value="Regula: fiecare concluzie trebuie să aibă o cifră în spate și o acțiune în față.").font = Font(size=9, italic=True, color=GREY_TXT)
r += 2

rap_section(r, "INDICATORI CHEIE"); r += 1
ws.cell(row=r, column=1, value="Indicator").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=2, value="Perioada curentă").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=3, value="Δ % vs. comparativ").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=4, value="Comentariu").font = Font(bold=True, size=10, color=WHITE)
for k in range(1, 5):
    ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=INK)
    ws.cell(row=r, column=k).alignment = Alignment(horizontal="center")
r += 1

REPORT_KPIS = [
    ("Venit net", F_MONEY), ("Profit brut", F_MONEY), ("Marjă brută %", F_PCT),
    ("Clienți activi", F_INT), ("Venit mediu / client", F_MONEY),
    ("Venit / FTE", F_MONEY), ("NPS", F_PCT2),
    ("Atingere țintă %", F_PCT), ("Gap față de țintă", F_MONEY),
    ("Concentrare — cel mai mare client", F_PCT),
]
for name, fmt in REPORT_KPIS:
    src = ROW[name]
    ws.cell(row=r, column=1, value=name).font = Font(size=10, color=INK, bold=True)
    for col, srccol in ((2, "B"), (3, "E")):
        c = ws.cell(row=r, column=col, value=f"='20_KPI'!${srccol}${src}")
        c.number_format = fmt if col == 2 else F_PCT
        c.fill = PatternFill("solid", fgColor=CALC_FILL)
        c.border = BORDER
        c.font = Font(size=10, bold=(col == 2), color=INK if col == 2 else GREY_TXT)
    cc = ws.cell(row=r, column=4)
    cc.fill = PatternFill("solid", fgColor=INPUT_FILL)
    cc.border = BORDER
    cc.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
r += 1

rap_section(r, "TOP 5 OPORTUNITĂȚI (ordonate după scor)"); r += 1
ws.cell(row=r, column=1, value="#").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=2, value="Oportunitate").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=3, value="Scor").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=4, value="Acțiune recomandată").font = Font(bold=True, size=10, color=WHITE)
for k in range(1, 5):
    ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=INK)
r += 1
SCORE_RNG = f"'30_OPORTUNITATI'!$I${first_opp}:$I${last_opp}"
TYPE_RNG = f"'30_OPORTUNITATI'!$C${first_opp}:$C${last_opp}"
TGT_RNG = f"'30_OPORTUNITATI'!$D${first_opp}:$D${last_opp}"
ACT_RNG = f"'30_OPORTUNITATI'!$K${first_opp}:$K${last_opp}"
for k in range(1, 6):
    ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10, color=INK)
    mt = f'MATCH(LARGE({SCORE_RNG},{k}),{SCORE_RNG},0)'
    ws.cell(row=r, column=2, value=f'=IFERROR(INDEX({TYPE_RNG},{mt})&IF(INDEX({TGT_RNG},{mt})="",""," — "&INDEX({TGT_RNG},{mt})),"")')
    ws.cell(row=r, column=3, value=f'=IFERROR(LARGE({SCORE_RNG},{k}),"")').number_format = F_MONEY
    ws.cell(row=r, column=4, value=f'=IFERROR(INDEX({ACT_RNG},{mt}),"")')
    for j in range(2, 5):
        c = ws.cell(row=r, column=j)
        c.fill = PatternFill("solid", fgColor=CALC_FILL)
        c.border = BORDER
        c.font = Font(size=10, color=GREY_TXT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 1
r += 1

rap_section(r, "VALOAREA TOTALĂ A PLANULUI"); r += 1
rap_row(r, "Valoare ponderată a oportunităților", f"='30_OPORTUNITATI'!${SL2}${hr_opp+3}", F_MONEY,
        "Suma scorurilor: impact ajustat cu probabilitatea și ușurința."); r += 1
rap_row(r, "Gap față de țintă", f"='20_KPI'!$B${ROW['Gap față de țintă']}", F_MONEY,
        "Cât lipsește până la țintă în perioada analizată."); r += 1
rap_row(r, "Acoperirea gapului", f"='30_OPORTUNITATI'!${SL2}${hr_opp+7}", F_PCT,
        "Peste 100% înseamnă că planul acoperă integral deficitul, dacă se execută."); r += 2

rap_section(r, "PRIMELE 30 DE ZILE"); r += 1
ws.cell(row=r, column=1, value="#").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=2, value="Acțiune").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=3, value="Responsabil").font = Font(bold=True, size=10, color=WHITE)
ws.cell(row=r, column=4, value="Termen / rezultat așteptat").font = Font(bold=True, size=10, color=WHITE)
for k in range(1, 5):
    ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=INK)
r += 1
for k in range(1, 6):
    ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10, color=INK)
    for j in range(2, 5):
        c = ws.cell(row=r, column=j)
        c.fill = PatternFill("solid", fgColor=INPUT_FILL)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 22
    r += 1
r += 1
ws.cell(row=r, column=1, value="Notă de metodă").font = Font(bold=True, size=10, color=INK)
nc = ws.cell(row=r, column=2, value=("Cifrele provin exclusiv din datele furnizate de client în perioada indicată. "
                                     "Estimările de impact sunt scenarii, nu garanții, și se recalibrează după primele 30 de zile de execuție."))
nc.font = Font(size=9, italic=True, color=GREY_TXT)
nc.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r].height = 32
ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------- 41_GRAFICE
# Paleta: slot 1 albastru / slot 2 portocaliu din paleta categoriala validata;
# gri pentru linia de referinta (tinta); rampa ordinala de albastru pentru prioritati.

ws = wb.create_sheet("41_GRAFICE")
title_block(ws, "41 — GRAFICE",
            "Se desenează singure din datele introduse. Dacă par goale, apasă Ctrl+Alt+F9 (recalculare completă). "
            "Zona de calcul de sub grafice nu se editează.")
ws.sheet_view.showGridLines = False

DATA_ROW = 60          # zona de calcul care alimenteaza graficele
MONTHS = 12
m0, m1 = DATA_ROW + 1, DATA_ROW + MONTHS

def hdr(cell, text, w=None):
    c = ws[cell]
    c.value = text
    c.font = Font(bold=True, size=9, color=WHITE)
    c.fill = PatternFill("solid", fgColor=INK)
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    if w:
        ws.column_dimensions[cell[0]].width = w

def calc(cell, formula, fmt=None):
    c = ws[cell]
    c.value = formula
    c.fill = PatternFill("solid", fgColor=CALC_FILL)
    c.border = BORDER
    c.font = Font(size=9, color=GREY_TXT)
    if fmt:
        c.number_format = fmt

ws.cell(row=DATA_ROW - 1, column=1,
        value="ZONA DE CALCUL PENTRU GRAFICE — nu se editează, nu se șterge.").font = \
    Font(bold=True, size=10, color="9C2A2A")

# --- T1: evolutia lunara (ultimele 12 luni pana la finalul perioadei analizate)
for col, (t, w) in zip("ABCDEF", [("Lună", 11), ("Etichetă", 11), ("Venit net", 13),
                                  ("Profit brut", 13), ("Țintă", 13), ("Marjă %", 10)]):
    hdr(f"{col}{DATA_ROW}", t, w)
for k in range(1, MONTHS + 1):
    rr = DATA_ROW + k
    ms, me = f"$A{rr}", f'DATE(YEAR($A{rr}),MONTH($A{rr})+1,0)'
    calc(f"A{rr}", f'=DATE(YEAR({ct}),MONTH({ct})+{k}-{MONTHS},1)', "yyyy-mm")
    calc(f"B{rr}", f'={ym(ms)}')
    calc(f"C{rr}", f'=SUMIFS({TX_REV},{TX_DATE},">="&{ms},{TX_DATE},"<="&{me})', F_MONEY)
    calc(f"D{rr}", f'=SUMIFS({TX_GP},{TX_DATE},">="&{ms},{TX_DATE},"<="&{me})', F_MONEY)
    calc(f"E{rr}", f'=SUMIFS({TG_VAL},{TG_DATE},">="&{ms},{TG_DATE},"<="&{me},'
                   f'{TG_METRIC},"Venit",{TG_SCOPE},"Companie")', F_MONEY)
    calc(f"F{rr}", f'=IFERROR($D{rr}/$C{rr},0)', F_PCT)

# --- T2 / T3: clasamente
def clasament(col_nume, col_val, titlu, sheet, col_eticheta, col_valoare, cheie, f, l, n=10):
    hdr(f"{col_nume}{DATA_ROW}", titlu, 26)
    hdr(f"{col_val}{DATA_ROW}", "Venit net", 13)
    key = f"'{sheet}'!${cheie}${f}:${cheie}${l}"
    for k in range(1, n + 1):
        rr = DATA_ROW + k
        mt = f'MATCH(LARGE({key},{k}),{key},0)'
        calc(f"{col_nume}{rr}", f"=IFERROR(INDEX('{sheet}'!${col_eticheta}${f}:${col_eticheta}${l},{mt}),\"\")")
        calc(f"{col_val}{rr}", f"=IFERROR(INDEX('{sheet}'!${col_valoare}${f}:${col_valoare}${l},{mt}),0)", F_MONEY)

clasament("H", "I", "Top 10 clienți", "21_ANALIZA_CLIENTI", "B", "E", KEY_A, fA, lA)
clasament("K", "L", "Top 10 produse", "22_ANALIZA_PRODUSE", "B", "D", KEY_P, fP, lP)

# --- T4: structura NPS
hdr(f"N{DATA_ROW}", "Categorie NPS", 16)
hdr(f"O{DATA_ROW}", "Răspunsuri", 12)
nps_win = f'({NP_DATE}>={cf})*({NP_DATE}<={ct})'
for k, (eticheta, cond) in enumerate([
        ("Promotori (9-10)", f'({NP_SCORE}>=9)'),
        ("Pasivi (7-8)", f'({NP_SCORE}>=7)*({NP_SCORE}<=8)'),
        ("Detractori (0-6)", f'({NP_SCORE}<=6)*({NP_SCORE}<>"")')], start=1):
    rr = DATA_ROW + k
    calc(f"N{rr}", f'="{eticheta}"')
    calc(f"O{rr}", f'=SUMPRODUCT({nps_win}*{cond})', F_INT)

# --- T5: oportunitati pe prioritate
hdr(f"Q{DATA_ROW}", "Prioritate", 14)
hdr(f"R{DATA_ROW}", "Valoare ponderată", 15)
for k, pri in enumerate(("HIGH", "MEDIUM", "LOW"), start=1):
    rr = DATA_ROW + k
    calc(f"Q{rr}", f'="{pri}"')
    calc(f"R{rr}", f'=SUMIF(\'30_OPORTUNITATI\'!$J${first_opp}:$J${last_opp},"{pri}",'
                   f'\'30_OPORTUNITATI\'!$I${first_opp}:$I${last_opp})', F_MONEY)


def etichete_text(ch, ref):
    """Categoriile sunt text; fara asta openpyxl le declara ca numere si Excel
    afiseaza 1, 2, 3... in loc de etichete."""
    for ser in ch.series:
        if ser.cat is not None:
            ser.cat.numRef = None
            ser.cat.strRef = StrRef(f=ref)


def stil_axe(ch, numfmt=F_MONEY, grid=False):
    # La bare verticale ("col") axa de categorii e jos; la bare orizontale ("bar"), la stanga.
    orizontal = getattr(ch, "type", "col") == "bar"
    ch.x_axis.axPos = "l" if orizontal else "b"
    ch.y_axis.axPos = "b" if orizontal else "l"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.numFmt = numfmt
    ch.y_axis.majorGridlines = None if not grid else ch.y_axis.majorGridlines
    for ax in (ch.x_axis, ch.y_axis):
        ax.spPr = GraphicalProperties()
        ax.spPr.ln = LineProperties(solidFill="C3C2B7", w=6350)
    ch.style = None


def culoare(ser, hexa, line=False):
    ser.graphicalProperties = GraphicalProperties(solidFill=hexa)
    if line:
        ser.graphicalProperties.line = LineProperties(solidFill=hexa, w=25400)
        ser.graphicalProperties.solidFill = None
    else:
        ser.graphicalProperties.line = LineProperties(noFill=True)


# 1. Venit si profit brut pe luna (aceeasi unitate => o singura axa) + tinta ca referinta
ch1 = BarChart()
ch1.type, ch1.grouping, ch1.gapWidth, ch1.overlap = "col", "clustered", 60, -10
ch1.title = "Venit net și profit brut, pe lună"
ch1.add_data(Reference(ws, min_col=3, max_col=4, min_row=DATA_ROW, max_row=m1), titles_from_data=True)
ch1.set_categories(Reference(ws, min_col=2, min_row=m0, max_row=m1))
culoare(ch1.series[0], C_SER1)
culoare(ch1.series[1], C_SER2)
ln = LineChart()
ln.add_data(Reference(ws, min_col=5, min_row=DATA_ROW, max_row=m1), titles_from_data=True)
culoare(ln.series[0], C_MUTED, line=True)
ln.series[0].smooth = False
ch1 += ln
etichete_text(ch1, f"'41_GRAFICE'!$B${m0}:$B${m1}")
stil_axe(ch1, grid=True)
ch1.y_axis.title = None
ch1.legend.position = "b"
ch1.width, ch1.height = 17, 8.5
ws.add_chart(ch1, "A4")

# 2. Marja bruta % pe luna (alta unitate => grafic separat, niciodata a doua axa)
ch2 = LineChart()
ch2.title = "Marjă brută %, pe lună"
ch2.add_data(Reference(ws, min_col=6, min_row=DATA_ROW, max_row=m1), titles_from_data=True)
ch2.set_categories(Reference(ws, min_col=2, min_row=m0, max_row=m1))
culoare(ch2.series[0], C_SER1, line=True)
ch2.series[0].smooth = False
etichete_text(ch2, f"'41_GRAFICE'!$B${m0}:$B${m1}")
stil_axe(ch2, numfmt="0.0%", grid=True)
ch2.legend = None
ch2.width, ch2.height = 17, 8.5
ws.add_chart(ch2, "L4")

# 3. Top 10 clienti
ch3 = BarChart()
ch3.type, ch3.gapWidth = "bar", 45
ch3.title = "Top 10 clienți după venit (perioada curentă)"
ch3.add_data(Reference(ws, min_col=9, min_row=DATA_ROW, max_row=DATA_ROW + 10), titles_from_data=True)
ch3.set_categories(Reference(ws, min_col=8, min_row=m0, max_row=DATA_ROW + 10))
culoare(ch3.series[0], C_SER1)
etichete_text(ch3, f"'41_GRAFICE'!$H${m0}:$H${DATA_ROW + 10}")
stil_axe(ch3)
ch3.legend = None
ch3.dataLabels = DataLabelList()
ch3.dataLabels.showVal = True
ch3.dataLabels.numFmt = F_MONEY
ch3.width, ch3.height = 17, 9.5
ws.add_chart(ch3, "A22")

# 4. Top 10 produse
ch4 = BarChart()
ch4.type, ch4.gapWidth = "bar", 45
ch4.title = "Top 10 produse după venit (perioada curentă)"
ch4.add_data(Reference(ws, min_col=12, min_row=DATA_ROW, max_row=DATA_ROW + 10), titles_from_data=True)
ch4.set_categories(Reference(ws, min_col=11, min_row=m0, max_row=DATA_ROW + 10))
culoare(ch4.series[0], C_SER1)
etichete_text(ch4, f"'41_GRAFICE'!$K${m0}:$K${DATA_ROW + 10}")
stil_axe(ch4)
ch4.legend = None
ch4.dataLabels = DataLabelList()
ch4.dataLabels.showVal = True
ch4.dataLabels.numFmt = F_MONEY
ch4.width, ch4.height = 17, 9.5
ws.add_chart(ch4, "L22")

# 5. Structura NPS (culori de stare, dar categoriile sunt scrise, nu doar colorate)
ch5 = BarChart()
ch5.type, ch5.gapWidth = "col", 80
ch5.title = "Structura răspunsurilor NPS"
ch5.add_data(Reference(ws, min_col=15, min_row=DATA_ROW, max_row=DATA_ROW + 3), titles_from_data=True)
ch5.set_categories(Reference(ws, min_col=14, min_row=m0, max_row=DATA_ROW + 3))
culoare(ch5.series[0], C_GOOD)
ch5.series[0].data_points = [
    DataPoint(idx=i, spPr=GraphicalProperties(solidFill=hexa, ln=LineProperties(noFill=True)))
    for i, hexa in enumerate((C_GOOD, C_WARN, C_BAD))]
etichete_text(ch5, f"'41_GRAFICE'!$N${m0}:$N${DATA_ROW + 3}")
stil_axe(ch5, numfmt=F_INT)
ch5.legend = None
ch5.dataLabels = DataLabelList()
ch5.dataLabels.showVal = True
ch5.width, ch5.height = 17, 8.5
ws.add_chart(ch5, "A40")

# 6. Oportunitati pe prioritate (rampa ordinala, o singura nuanta)
ch6 = BarChart()
ch6.type, ch6.gapWidth = "col", 80
ch6.title = "Valoarea ponderată a oportunităților, pe prioritate"
ch6.add_data(Reference(ws, min_col=18, min_row=DATA_ROW, max_row=DATA_ROW + 3), titles_from_data=True)
ch6.set_categories(Reference(ws, min_col=17, min_row=m0, max_row=DATA_ROW + 3))
culoare(ch6.series[0], C_ORD[0])
ch6.series[0].data_points = [
    DataPoint(idx=i, spPr=GraphicalProperties(solidFill=C_ORD[i], ln=LineProperties(noFill=True)))
    for i in range(3)]
etichete_text(ch6, f"'41_GRAFICE'!$Q${m0}:$Q${DATA_ROW + 3}")
stil_axe(ch6)
ch6.legend = None
ch6.dataLabels = DataLabelList()
ch6.dataLabels.showVal = True
ch6.dataLabels.numFmt = F_MONEY
ch6.width, ch6.height = 17, 8.5
ws.add_chart(ch6, "L40")

# 7. Nota publica: firma fata de concurenti (evidentiem bara noastra, restul raman discrete)
CONC_NOTA = get_column_letter(list(CONC["col"]).index("Notă publică (1-5)") + 1)
CONC_CINE = get_column_letter(list(CONC["col"]).index("Cine") + 1)
ws_conc = wb[CONC["foaie"]]
ch7 = BarChart()
ch7.type, ch7.gapWidth = "bar", 45
ch7.title = "Nota publică: firma față de concurenți"
ch7.add_data(Reference(ws_conc, min_col=list(CONC["col"]).index("Notă publică (1-5)") + 1,
                       min_row=CONC["prim_rand"] - 1, max_row=CONC["prim_rand"] + 9),
             titles_from_data=True)
ch7.set_categories(Reference(ws_conc, min_col=list(CONC["col"]).index("Cine") + 1,
                             min_row=CONC["prim_rand"], max_row=CONC["prim_rand"] + 9))
culoare(ch7.series[0], "C3C2B7")
ch7.series[0].data_points = [
    DataPoint(idx=0, spPr=GraphicalProperties(solidFill=C_SER1, ln=LineProperties(noFill=True)))]
etichete_text(ch7, f"'{CONC['foaie']}'!${CONC_CINE}${CONC['prim_rand']}:${CONC_CINE}${CONC['prim_rand'] + 9}")
stil_axe(ch7, numfmt="0.0", grid=False)
ch7.y_axis.scaling.min = 0
ch7.y_axis.scaling.max = 5
ch7.legend = None
ch7.dataLabels = DataLabelList()
ch7.dataLabels.showVal = True
ch7.dataLabels.numFmt = "0.0"
ch7.width, ch7.height = 17, 9.5
ws.add_chart(ch7, "A58")

# ordinea foilor
order = ["00_GHID", "01_FIRMA", "02_CONTACTE", "03_CHECKLIST_DATE", "04_DICTIONAR", SH_PAR,
         SH_CLIENTI, SH_PRODUSE, SH_OAMENI, SH_VANZARI, SH_TINTE, SH_NPS,
         "16_CONCURENTA", "17_CE_CRED_OAMENII",
         "20_KPI", "21_ANALIZA_CLIENTI", "22_ANALIZA_PRODUSE", "23_ANALIZA_OAMENI",
         "30_OPORTUNITATI", "40_RAPORT", "41_GRAFICE", "90_LISTE"]
wb._sheets = [wb[s] for s in order]
wb.active = 0

out = OUT
verifica(wb)
wb.save(out)
print("OK ->", out)
print("Foi:", len(wb.sheetnames))
