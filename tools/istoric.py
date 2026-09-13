# -*- coding: utf-8 -*-
"""Istoricul financiar pe trei ani, cu doua grafice.

O singura perioada nu spune daca firma creste sau se scufunda. Trei ani de
cifra de afaceri, profit brut si profit net spun, si se gasesc gratuit in
datele publice ale firmelor.

Folosit de ambele variante, in foaia 01_FIRMA.
"""
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font, PatternFill

from stil import (INK, TEAL, TEAL_LIGHT, INPUT_FILL, CALC_FILL, WHITE, GREY_TXT,
                  C_SER1, C_SER2, F_MONEY, F_PCT, F_INT, BORDER, mon, pct)

C_SER3 = "1BAF7A"   # slotul 3 din paleta validata (albastru, portocaliu, acvamarin)


def adauga_istoric(ws, rand, an_ref, moneda_ref, note_col=5):
    """Scrie blocul de istoric incepand de la randul dat. `an_ref` e celula cu
    data de sfarsit a perioadei analizate, din care se deduc anii."""

    def sectiune(r, text):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=True, size=11, color=WHITE)
        for k in range(1, note_col + 1):
            ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor=TEAL)
        ws.row_dimensions[r].height = 20

    sectiune(rand, "ISTORIC FINANCIAR — ULTIMII TREI ANI")
    r = rand + 1

    # capul de tabel: anii. Se deduc din perioada analizata, dar se pot rescrie.
    ws.cell(row=r, column=1, value="An").font = Font(bold=True, size=10, color=WHITE)
    for k in range(3):
        c = ws.cell(row=r, column=2 + k, value=f"=YEAR({an_ref})-{2 - k}")
        c.number_format = "0"
        c.font = Font(bold=True, size=10, color="000000")
        c.fill = PatternFill("solid", fgColor=INPUT_FILL)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=INK)
    n = ws.cell(row=r, column=note_col,
                value="Se completează singuri din perioada analizată. Dacă ultimul bilanț depus "
                      "e mai vechi, scrie anii peste.")
    n.font = Font(size=9, italic=True, color=GREY_TXT)
    n.alignment = Alignment(wrap_text=True, vertical="center")
    rand_ani = r
    r += 1

    INTRARI = [
        ("Cifră de afaceri", F_MONEY,
         "Fără TVA. Din bilanțul depus la ANAF — se găsește gratuit, public, după CUI."),
        ("Profit brut", F_MONEY,
         "Profitul înainte de impozit. Tot din bilanț."),
        ("Profit net", F_MONEY,
         "Profitul după impozit. Ăsta rămâne efectiv în firmă."),
        ("Număr angajați", F_INT,
         "La sfârșitul anului. Arată dacă creșterea a venit cu oameni în plus sau fără."),
    ]
    rand_intrare = {}
    for eticheta, fmt, nota in INTRARI:
        ws.cell(row=r, column=1, value=eticheta).font = Font(bold=True, size=10, color=INK)
        for k in range(3):
            c = ws.cell(row=r, column=2 + k)
            c.fill = PatternFill("solid", fgColor=INPUT_FILL)
            c.border = BORDER
            c.number_format = fmt
            c.font = Font(size=10)
        nn = ws.cell(row=r, column=note_col, value=nota)
        nn.font = Font(size=9, italic=True, color=GREY_TXT)
        nn.alignment = Alignment(wrap_text=True, vertical="center")
        rand_intrare[eticheta] = r
        r += 1

    rCA = rand_intrare["Cifră de afaceri"]
    rPB = rand_intrare["Profit brut"]
    rPN = rand_intrare["Profit net"]
    rNR = rand_intrare["Număr angajați"]

    CALCULE = [
        ("Marjă brută", F_PCT, lambda c: f'=IF(OR({c}{rCA}="",{c}{rPB}=""),"",IFERROR({c}{rPB}/{c}{rCA},""))',
         "Profit brut împărțit la cifra de afaceri."),
        ("Marjă netă", F_PCT, lambda c: f'=IF(OR({c}{rCA}="",{c}{rPN}=""),"",IFERROR({c}{rPN}/{c}{rCA},""))',
         "Cât rămâne din fiecare leu vândut, după tot."),
        ("Creștere cifră de afaceri", F_PCT, None,
         "Față de anul anterior. Primul an nu are cu ce se compara."),
        ("Cifră de afaceri / angajat", F_MONEY,
         lambda c: f'=IF(OR({c}{rCA}="",{c}{rNR}="",{c}{rNR}=0),"",IFERROR({c}{rCA}/{c}{rNR},""))',
         "Productivitatea. Dacă scade în timp ce cifra crește, firma angajează mai repede decât vinde."),
    ]
    rand_calcul = {}
    for eticheta, fmt, formula, nota in CALCULE:
        ws.cell(row=r, column=1, value=eticheta).font = Font(bold=True, size=10, color=INK)
        for k, col in enumerate("BCD"):
            c = ws.cell(row=r, column=2 + k)
            c.fill = PatternFill("solid", fgColor=CALC_FILL)
            c.border = BORDER
            c.number_format = fmt
            c.font = Font(size=10, color=GREY_TXT)
            if formula:
                c.value = formula(col)
        if eticheta == "Creștere cifră de afaceri":
            ws.cell(row=r, column=2, value="—").alignment = Alignment(horizontal="center")
            for col, prec in (("C", "B"), ("D", "C")):
                ws[f"{col}{r}"] = (f'=IF(OR({col}{rCA}="",{prec}{rCA}="",{prec}{rCA}=0),"",'
                                   f'IFERROR(({col}{rCA}-{prec}{rCA})/ABS({prec}{rCA}),""))')
        nn = ws.cell(row=r, column=note_col, value=nota)
        nn.font = Font(size=9, italic=True, color=GREY_TXT)
        nn.alignment = Alignment(wrap_text=True, vertical="center")
        rand_calcul[eticheta] = r
        r += 1

    ws.cell(row=r, column=1, value="Sursa cifrelor").font = Font(bold=True, size=10, color=INK)
    for k in range(3):
        c = ws.cell(row=r, column=2 + k)
        c.fill = PatternFill("solid", fgColor=INPUT_FILL)
        c.border = BORDER
        c.font = Font(size=10)
    nn = ws.cell(row=r, column=note_col,
                 value="Bilanț ANAF, balanță internă sau spus de client. Contează, pentru că "
                       "primele două se pot verifica, a treia nu.")
    nn.font = Font(size=9, italic=True, color=GREY_TXT)
    nn.alignment = Alignment(wrap_text=True, vertical="center")
    r += 2

    # citirea automata a celor trei ani
    rMN = rand_calcul["Marjă netă"]
    CA1, CA3 = f"$B${rCA}", f"$D${rCA}"
    MN1, MN3 = f"$B${rMN}", f"$D${rMN}"
    ws.cell(row=r, column=1, value="Ce se vede").font = Font(bold=True, size=10, color=INK)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=note_col)
    citire = ws.cell(row=r, column=2, value=(
        f'=IF({CA3}="","Completează măcar ultimul an. Cifrele sunt publice: se caută după CUI.",'
        f'IF({CA1}="","Ai un singur an. O cifră singură nu spune dacă firma crește sau scade — '
        f'adaugă și anii anteriori.",'
        f'"Cifra de afaceri "&IF({CA3}>={CA1},"a crescut cu ","a scăzut cu ")&'
        f'{pct(f"ABS(({CA3}-{CA1})/{CA1})")}&" în trei ani"&'
        f'IF(OR({MN3}="",{MN1}=""),".",'
        f'", iar marja netă a trecut de la "&{pct(MN1)}&" la "&{pct(MN3)}&". "&'
        f'IF(AND({CA3}>{CA1},{MN3}<{MN1}-0.01),'
        f'"Firma crește vânzând mai mult, dar păstrează mai puțin din fiecare leu. '
        f'De obicei: discount în creștere sau costuri care urcă mai repede decât vânzările.",'
        f'IF(AND({CA3}<{CA1},{MN3}>{MN1}+0.01),'
        f'"Vinde mai puțin, dar mai profitabil. Poate fi o alegere bună, dacă e intenționată.",'
        f'IF(AND({CA3}>{CA1},{MN3}>={MN1}),'
        f'"Crește și își păstrează marja — creșterea e sănătoasă.",'
        f'"Și vânzările, și marja au scăzut. Aici se începe cu marja, nu cu volumul."))))))'))
    citire.fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    citire.border = BORDER
    citire.font = Font(size=10, color="1A2E2E")
    citire.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 44
    r += 2

    # ---- graficele
    def etichete_text(ch):
        ref = f"'{ws.title}'!$B${rand_ani}:$D${rand_ani}"
        for ser in ch.series:
            if ser.cat is not None:
                ser.cat.numRef = None
                ser.cat.strRef = StrRef(f=ref)

    def axe(ch, numfmt):
        ch.x_axis.axPos, ch.y_axis.axPos = "b", "l"
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.y_axis.numFmt = numfmt
        for ax in (ch.x_axis, ch.y_axis):
            ax.spPr = GraphicalProperties()
            ax.spPr.ln = LineProperties(solidFill="C3C2B7", w=6350)
        ch.style = None

    def culoare(ser, hexa, linie=False):
        ser.graphicalProperties = GraphicalProperties()
        if linie:
            ser.graphicalProperties.line = LineProperties(solidFill=hexa, w=25400)
        else:
            ser.graphicalProperties.solidFill = hexa
            ser.graphicalProperties.line = LineProperties(noFill=True)

    g1 = BarChart()
    g1.type, g1.grouping, g1.gapWidth, g1.overlap = "col", "clustered", 70, -10
    g1.title = "Cifră de afaceri, profit brut și profit net"
    g1.add_data(Reference(ws, min_col=1, max_col=4, min_row=rCA, max_row=rPN),
                titles_from_data=True, from_rows=True)
    g1.set_categories(Reference(ws, min_col=2, max_col=4, min_row=rand_ani))
    for ser, hexa in zip(g1.series, (C_SER1, C_SER2, C_SER3)):
        culoare(ser, hexa)
    etichete_text(g1)
    axe(g1, F_MONEY)
    g1.legend.position = "b"
    # etichetele de valoare sunt obligatorii: verdele are contrast sub 3:1 pe alb
    g1.dataLabels = DataLabelList()
    g1.dataLabels.showVal = True
    g1.dataLabels.numFmt = F_MONEY
    g1.width, g1.height = 16, 8
    ws.add_chart(g1, f"A{r}")

    g2 = LineChart()
    g2.title = "Marja brută și marja netă"
    g2.add_data(Reference(ws, min_col=1, max_col=4,
                          min_row=rand_calcul["Marjă brută"], max_row=rMN),
                titles_from_data=True, from_rows=True)
    g2.set_categories(Reference(ws, min_col=2, max_col=4, min_row=rand_ani))
    for ser, hexa in zip(g2.series, (C_SER1, C_SER2)):
        culoare(ser, hexa, linie=True)
        ser.smooth = False
    etichete_text(g2)
    axe(g2, "0.0%")
    g2.legend.position = "b"
    g2.dataLabels = DataLabelList()
    g2.dataLabels.showVal = True
    g2.dataLabels.numFmt = "0.0%"
    g2.width, g2.height = 16, 8
    ws.add_chart(g2, f"A{r + 17}")

    return dict(rand_ani=rand_ani, rand_ca=rCA, rand_pb=rPB, rand_pn=rPN,
                rand_marja_neta=rMN, ultim_rand=r + 34,
                ca_ultim=f"'{ws.title}'!$D${rCA}",
                ca_prim=f"'{ws.title}'!$B${rCA}",
                marja_neta_ultim=f"'{ws.title}'!$D${rMN}")
