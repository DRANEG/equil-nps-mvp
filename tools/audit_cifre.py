# -*- coding: utf-8 -*-
"""Audit numeric complet al variantei complete.

Nu verifica doar cativa indicatori, ci recalculeaza in Python fiecare celula
calculata din foile de analiza si o compara cu ce da Excel. Datele de test sunt
alese sa fie dificile: storno cu valoare negativa, client pierdut, client
adormit, client in scadere, produs fara vanzari, om plecat la mijlocul
perioadei, tinte pe mai multe niveluri, raspunsuri NPS in afara perioadei.

Ruleaza: python3 tools/audit_cifre.py
"""
import datetime as dt
import os
import subprocess
import sys
from collections import defaultdict

MIC = dict(N_CLIENTI="8", N_PRODUSE="5", N_OAMENI="4", N_VANZARI="40",
           N_TINTE="10", N_NPS="12", N_OPP="8", N_CONTACTE="4",
           OUT="/tmp/audit_full.xlsx")
subprocess.run([sys.executable, "tools/build_equil_workbook.py"],
               env=dict(os.environ, **MIC), check=True, stdout=subprocess.DEVNULL)

import openpyxl

CALE = MIC["OUT"]
wb = openpyxl.load_workbook(CALE)
R0 = 5

CUR = (dt.date(2026, 1, 1), dt.date(2026, 6, 30))
CMP = (dt.date(2025, 1, 1), dt.date(2025, 6, 30))
PRAG_SCADERE, PRAG_CRESTERE = -0.15, 0.20
PRAG_MARJA, ZILE_DORMANT = 0.30, 90
PRAG_NPS, PRAG_TINTA, PRAG_CONC = 30, 0.90, 0.20

# ------------------------------------------------------------------ date
CLIENTI = [
    # id, nume, segment, canal, regiune, tara, responsabil, din, status
    ("C1", "Client Pierdut SRL", "Small", "Retail", "Cluj", "RO", "S1", dt.date(2022, 1, 1), "Activ"),
    ("C2", "Client Adormit SRL", "Mid-market", "Online", "Iași", "RO", "S1", dt.date(2022, 1, 1), "Activ"),
    ("C3", "Client In Scadere SA", "Key Account", "Wholesale", "București", "RO", "S2", dt.date(2021, 1, 1), "Activ"),
    ("C4", "Client In Crestere SRL", "Mid-market", "Retail", "Timiș", "RO", "S2", dt.date(2023, 1, 1), "Activ"),
    ("C5", "Client Marja Mica SRL", "Small", "Online", "Brașov", "RO", "S3", dt.date(2023, 6, 1), "Activ"),
    ("C6", "Client Nou SRL", "Small", "Retail", "Cluj", "RO", "S3", dt.date(2026, 2, 1), "Activ"),
]
PRODUSE = [
    ("P1", "Produs Marja Buna", "Categoria A", "", 100, 55, "Nu", None, "Activ"),
    ("P2", "Produs Marja Mica", "Categoria A", "", 100, 88, "Nu", None, "Activ"),
    ("P3", "Produs In Crestere", "Categoria B", "", 200, 120, "Da", None, "Activ"),
    ("P4", "Produs Fara Vanzari", "Categoria B", "", 50, 30, "Nu", None, "Activ"),
]
OAMENI = [
    # id, nume, rol, echipa, fte, start, exit, cost, are_tinta
    ("S1", "Ana", "KAM", "B2B", 1.0, dt.date(2023, 1, 1), None, 90000, "Da"),
    ("S2", "Bogdan", "Sales", "B2B", 0.5, dt.date(2024, 1, 1), None, 45000, "Da"),
    ("S3", "Carmen", "Sales", "B2B", 1.0, dt.date(2023, 1, 1), dt.date(2026, 3, 31), 80000, "Da"),
]
# (data, tip, client, produs, persoana, unitati, venit, discount, cost)
VANZARI = [
    # --- perioada comparativa 2025 H1
    (dt.date(2025, 2, 28), "Factură", "C1", "P1", "S1", 100, 10000, 0, 5500),
    (dt.date(2025, 3, 31), "Factură", "C2", "P2", "S1", 80, 8000, 0, 7040),
    (dt.date(2025, 4, 30), "Factură", "C3", "P1", "S2", 200, 20000, 500, 11000),
    (dt.date(2025, 5, 31), "Factură", "C4", "P3", "S2", 25, 5000, 0, 3000),
    (dt.date(2025, 6, 30), "Factură", "C5", "P2", "S3", 50, 5000, 0, 4400),
    # --- in afara ambelor perioade (nu trebuie sa intre nicaieri)
    (dt.date(2025, 9, 30), "Factură", "C3", "P1", "S2", 10, 1000, 0, 550),
    # --- perioada curenta 2026 H1
    (dt.date(2026, 1, 31), "Factură", "C2", "P2", "S1", 40, 4000, 0, 3520),   # ultima comanda C2
    (dt.date(2026, 2, 28), "Factură", "C3", "P1", "S2", 120, 12000, 300, 6600),
    (dt.date(2026, 3, 31), "Factură", "C4", "P3", "S2", 45, 9000, 0, 5400),
    (dt.date(2026, 3, 31), "Factură", "C6", "P1", "S3", 30, 3000, 0, 1650),
    (dt.date(2026, 4, 30), "Factură", "C5", "P2", "S1", 60, 6000, 0, 5280),
    (dt.date(2026, 5, 31), "Factură", "C4", "P1", "S2", 20, 2000, 0, 1100),
    (dt.date(2026, 6, 30), "Storno", "C3", "P1", "S2", -5, -500, 0, -275),   # storno negativ
]
TINTE = [
    (dt.date(2026, 3, 31), "Companie", "", "Venit", 20000, ""),
    (dt.date(2026, 6, 30), "Companie", "", "Venit", 25000, ""),
    (dt.date(2025, 6, 30), "Companie", "", "Venit", 40000, ""),
    (dt.date(2026, 6, 30), "Persoană", "S1", "Venit", 12000, ""),
    (dt.date(2026, 6, 30), "Persoană", "S2", "Venit", 30000, ""),
    (dt.date(2026, 6, 30), "Companie", "", "Unități", 999, ""),   # alt indicator: nu se ia
]
NPS = [
    # id, data, client, nume, rol, scor
    ("R1", dt.date(2026, 2, 10), "C3", "X", "Decident", 9),
    ("R2", dt.date(2026, 2, 20), "C4", "Y", "User", 10),
    ("R3", dt.date(2026, 3, 15), "C5", "Z", "User", 4),
    ("R4", dt.date(2026, 4, 10), "C2", "W", "Decident", 7),
    ("R5", dt.date(2026, 5, 5), "C6", "V", "User", 8),
    ("R6", dt.date(2025, 3, 1), "C1", "U", "User", 2),      # in perioada comparativa
    ("R7", dt.date(2026, 12, 1), "C4", "T", "User", 0),     # dupa perioada: se ignora
]


def scrie(foaie, randuri, coloane):
    ws = wb[foaie]
    for r in range(R0, ws.max_row + 1):
        for c in coloane:
            v = ws.cell(row=r, column=c).value
            if not (isinstance(v, str) and v.startswith("=")):
                ws.cell(row=r, column=c, value=None)
    for i, row in enumerate(randuri):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=R0 + i, column=coloane[j], value=v)


scrie("10_CLIENTI", CLIENTI, [1, 2, 3, 4, 5, 6, 7, 8, 9])
scrie("11_PRODUSE", PRODUSE, [1, 2, 3, 4, 5, 6, 7, 8, 9])
scrie("12_OAMENI", OAMENI, [1, 2, 3, 4, 5, 6, 7, 8, 9])
scrie("13_VANZARI", VANZARI, [1, 3, 4, 5, 6, 7, 8, 9, 10])
scrie("14_TINTE", TINTE, [1, 2, 3, 4, 5, 6])
scrie("15_NPS", NPS, [1, 2, 3, 4, 5, 6])
wb.save(CALE)

# ------------------------------------------------- calcule independente
def in_fereastra(d, f):
    return f[0] <= d <= f[1]


def venit(f=None, client=None, produs=None, persoana=None):
    t = 0
    for d, _tip, c, p, s, _u, v, _disc, _cost in VANZARI:
        if f and not in_fereastra(d, f):
            continue
        if client and c != client:
            continue
        if produs and p != produs:
            continue
        if persoana and s != persoana:
            continue
        t += v
    return t


def profit(f=None, client=None, produs=None, persoana=None):
    t = 0
    for d, _tip, c, p, s, _u, v, _disc, cost in VANZARI:
        if f and not in_fereastra(d, f):
            continue
        if client and c != client:
            continue
        if produs and p != produs:
            continue
        if persoana and s != persoana:
            continue
        t += v - cost
    return t


def unitati(f, produs=None):
    return sum(u for d, _t, _c, p, _s, u, _v, _d2, _co in VANZARI
               if in_fereastra(d, f) and (produs is None or p == produs))


def discount(f):
    return sum(x for d, _t, _c, _p, _s, _u, _v, x, _co in VANZARI if in_fereastra(d, f))


def linii(f):
    return sum(1 for d, *_r in VANZARI if in_fereastra(d, f))


def clienti_activi(f):
    return sum(1 for c in CLIENTI if venit(f, client=c[0]) != 0)


def clienti_noi(f):
    n = 0
    for c in CLIENTI:
        inainte = sum(v for d, _t, cc, _p, _s, _u, v, _d2, _co in VANZARI
                      if cc == c[0] and d < f[0])
        if inainte == 0 and venit(f, client=c[0]) != 0:
            n += 1
    return n


def fte(f):
    t = 0
    for _id, _n, _r, _e, ftev, start, exitd, *_rest in OAMENI:
        if start <= f[1] and (exitd is None or exitd >= f[0]):
            t += ftev
    return t


def top1(f):
    return max(venit(f, client=c[0]) for c in CLIENTI)


def tinta(f, scope="Companie", sid=""):
    return sum(v for d, sc, si, m, v, _n in TINTE
               if in_fereastra(d, f) and m == "Venit" and sc == scope and (si or "") == sid)


def nps_set(f):
    note = [s for _i, d, _c, _n, _r, s in NPS if in_fereastra(d, f)]
    return note


def nps(f):
    note = nps_set(f)
    if not note:
        return ""
    p = sum(1 for s in note if s >= 9)
    de = sum(1 for s in note if s <= 6)
    return (p - de) / len(note) * 100


def ultima_achizitie(client):
    """Ultima vânzare reală: un storno (venit negativ) nu e o achiziție."""
    date = [d for d, _t, c, _p, _s, _u, v, _d2, _co in VANZARI if c == client and v > 0]
    return max(date) if date else None


def nps_client(client, f):
    note = [s for _i, d, c, _n, _r, s in NPS if c == client and in_fereastra(d, f)]
    return sum(note) / len(note) if note else ""


def produse_cumparate(client, f):
    return sum(1 for p in PRODUSE if venit(f, client=client, produs=p[0]) != 0)


def clienti_produs(produs, f):
    return sum(1 for c in CLIENTI if venit(f, client=c[0], produs=produs) != 0)


def clienti_persoana(persoana, f):
    return sum(1 for c in CLIENTI
               if sum(v for d, _t, cc, _p, s, _u, v, _d2, _co in VANZARI
                      if in_fereastra(d, f) and cc == c[0] and s == persoana) != 0)


def delta(now, before):
    return "" if before == 0 else (now - before) / abs(before)


# ------------------------------------------------- evaluare in Excel
wb2 = openpyxl.load_workbook(CALE)
for w in wb2.worksheets:
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = c.value.replace("[", "(").replace("]", ")")
    w.data_validations.dataValidation = []
ascii_path = CALE.replace(".xlsx", "_a.xlsx")
wb2.save(ascii_path)

print("Se calculează workbook-ul...")
import formulas
sol = formulas.ExcelModel().loads(ascii_path).finish().calculate()
BOOK = os.path.basename(ascii_path)


def V(sh, cell):
    try:
        return sol[f"'[{BOOK}]{sh}'!{cell}"].value[0, 0]
    except KeyError:
        return None


rezultate = []


def cmp(eticheta, excel, astept, toleranta=0.01):
    if isinstance(astept, str) or astept is None:
        ok = (str(excel).strip() == str(astept).strip()) or (astept == "" and excel in (None, "", 0))
        if astept == "":
            ok = excel in (None, "", 0) or str(excel).strip() == ""
    else:
        try:
            ok = abs(float(excel) - float(astept)) <= toleranta
        except (TypeError, ValueError):
            ok = False
    rezultate.append((ok, eticheta, excel, astept))
    return ok


# ---------- 20_KPI
kpi = wb["20_KPI"]
rand = {kpi.cell(row=r, column=1).value: r for r in range(5, kpi.max_row + 1)}
ASTEPT_KPI = {
    "Venit net": (venit(CUR), venit(CMP)),
    "Profit brut": (profit(CUR), profit(CMP)),
    "Marjă brută %": (profit(CUR) / venit(CUR), profit(CMP) / venit(CMP)),
    "Discount acordat": (discount(CUR), discount(CMP)),
    "Unități vândute": (unitati(CUR), unitati(CMP)),
    "Preț mediu / unitate": (venit(CUR) / unitati(CUR), venit(CMP) / unitati(CMP)),
    "Clienți activi": (clienti_activi(CUR), clienti_activi(CMP)),
    "Clienți noi": (clienti_noi(CUR), clienti_noi(CMP)),
    "Venit mediu / client": (venit(CUR) / clienti_activi(CUR), venit(CMP) / clienti_activi(CMP)),
    "Număr de linii / facturi": (linii(CUR), linii(CMP)),
    "Valoare medie per linie": (venit(CUR) / linii(CUR), venit(CMP) / linii(CMP)),
    "Concentrare — cel mai mare client": (top1(CUR) / venit(CUR), top1(CMP) / venit(CMP)),
    "FTE activ": (fte(CUR), fte(CMP)),
    "Venit / FTE": (venit(CUR) / fte(CUR), venit(CMP) / fte(CMP)),
    "Profit brut / FTE": (profit(CUR) / fte(CUR), profit(CMP) / fte(CMP)),
    "NPS": (nps(CUR), nps(CMP)),
    "Număr răspunsuri NPS": (len(nps_set(CUR)), len(nps_set(CMP))),
    "% Promotori": (sum(1 for s in nps_set(CUR) if s >= 9) / len(nps_set(CUR)),
                    sum(1 for s in nps_set(CMP) if s >= 9) / len(nps_set(CMP))),
    "% Detractori": (sum(1 for s in nps_set(CUR) if s <= 6) / len(nps_set(CUR)),
                     sum(1 for s in nps_set(CMP) if s <= 6) / len(nps_set(CMP))),
    "Țintă de venit": (tinta(CUR), tinta(CMP)),
    "Atingere țintă %": (venit(CUR) / tinta(CUR), venit(CMP) / tinta(CMP)),
    "Gap față de țintă": (tinta(CUR) - venit(CUR), tinta(CMP) - venit(CMP)),
}
print("\n--- 20_KPI (curent + comparativ) ---")
for nume, (ec, ep) in ASTEPT_KPI.items():
    r = rand[nume]
    cmp(f"20_KPI!{nume} (curent)", V("20_KPI", f"B{r}"), ec)
    cmp(f"20_KPI!{nume} (comparativ)", V("20_KPI", f"C{r}"), ep)

# ---------- 21_ANALIZA_CLIENTI
print("--- 21_ANALIZA_CLIENTI (toate coloanele, toți clienții) ---")
for i, c in enumerate(CLIENTI):
    r = R0 + i
    cid = c[0]
    vc, vp = venit(CUR, client=cid), venit(CMP, client=cid)
    pb = profit(CUR, client=cid)
    ua = ultima_achizitie(cid)
    zile = (CUR[1] - ua).days if ua else ""
    notac = nps_client(cid, CUR)
    cmp(f"21!{cid} venit curent", V("21_ANALIZA_CLIENTI", f"E{r}"), vc)
    cmp(f"21!{cid} venit comparativ", V("21_ANALIZA_CLIENTI", f"F{r}"), vp)
    cmp(f"21!{cid} delta %", V("21_ANALIZA_CLIENTI", f"G{r}"), delta(vc, vp))
    cmp(f"21!{cid} profit brut", V("21_ANALIZA_CLIENTI", f"H{r}"), pb)
    cmp(f"21!{cid} marjă", V("21_ANALIZA_CLIENTI", f"I{r}"), pb / vc if vc else "")
    cmp(f"21!{cid} % din total", V("21_ANALIZA_CLIENTI", f"J{r}"), vc / venit(CUR))
    cmp(f"21!{cid} zile de la ultima", V("21_ANALIZA_CLIENTI", f"L{r}"), zile)
    cmp(f"21!{cid} nr. produse", V("21_ANALIZA_CLIENTI", f"M{r}"), produse_cumparate(cid, CUR))
    cmp(f"21!{cid} NPS mediu", V("21_ANALIZA_CLIENTI", f"N{r}"), notac)

# ---------- 22_ANALIZA_PRODUSE
print("--- 22_ANALIZA_PRODUSE ---")
for i, p in enumerate(PRODUSE):
    r = R0 + i
    pid = p[0]
    vc, vp = venit(CUR, produs=pid), venit(CMP, produs=pid)
    pb = profit(CUR, produs=pid)
    u = unitati(CUR, produs=pid)
    nrc = clienti_produs(pid, CUR)
    cmp(f"22!{pid} venit curent", V("22_ANALIZA_PRODUSE", f"D{r}"), vc)
    cmp(f"22!{pid} venit comparativ", V("22_ANALIZA_PRODUSE", f"E{r}"), vp)
    cmp(f"22!{pid} delta %", V("22_ANALIZA_PRODUSE", f"F{r}"), delta(vc, vp))
    cmp(f"22!{pid} profit brut", V("22_ANALIZA_PRODUSE", f"G{r}"), pb)
    cmp(f"22!{pid} marjă", V("22_ANALIZA_PRODUSE", f"H{r}"), pb / vc if vc else "")
    cmp(f"22!{pid} unități", V("22_ANALIZA_PRODUSE", f"I{r}"), u)
    cmp(f"22!{pid} preț mediu", V("22_ANALIZA_PRODUSE", f"J{r}"), vc / u if u else "")
    cmp(f"22!{pid} nr. clienți", V("22_ANALIZA_PRODUSE", f"K{r}"), nrc)
    cmp(f"22!{pid} penetrare", V("22_ANALIZA_PRODUSE", f"L{r}"), nrc / clienti_activi(CUR))


# ---------- semnalele și potențialul din 21_ANALIZA_CLIENTI
print("--- 21: semnale și potențial ---")


def semnal_client(cid):
    vc, vp = venit(CUR, client=cid), venit(CMP, client=cid)
    pb = profit(CUR, client=cid)
    marja = pb / vc if vc else ""
    g = delta(vc, vp)
    ua = ultima_achizitie(cid)
    zile = (CUR[1] - ua).days if ua else ""
    nota = nps_client(cid, CUR)
    cota = vc / venit(CUR)
    if vp > 0 and vc == 0:
        return "Client pierdut"
    if zile != "" and zile > ZILE_DORMANT and vp > 0:
        return "Client dormant"
    if g != "" and g <= PRAG_SCADERE:
        return "Scădere de venit"
    if nota != "" and nota < 7:
        return "Detractor activ"
    if vc > 0 and cota > PRAG_CONC:
        return "Risc concentrare"
    if vc > 0 and marja != "" and marja < PRAG_MARJA:
        return "Marjă sub referință"
    if g != "" and g >= PRAG_CRESTERE:
        return "În creștere"
    if vc == 0:
        return "Fără activitate"
    return "Stabil"


OPORTUNITATE = {
    "Risc concentrare": "Risc concentrare", "Client pierdut": "Recuperare client",
    "Client dormant": "Reactivare client", "Scădere de venit": "Recuperare client",
    "Detractor activ": "Recuperare NPS", "Marjă sub referință": "Marjă / preț",
    "În creștere": "Up-sell", "Fără activitate": "Reactivare client", "Stabil": "Cross-sell",
}
for i, c in enumerate(CLIENTI):
    r = R0 + i
    cid = c[0]
    sem = semnal_client(cid)
    vc, vp = venit(CUR, client=cid), venit(CMP, client=cid)
    pb = profit(CUR, client=cid)
    marja = pb / vc if vc else 0
    if sem == "Scădere de venit":
        pot = max(0, vp - vc)
    elif sem in ("Client pierdut", "Client dormant", "Fără activitate"):
        pot = vp
    elif sem == "Marjă sub referință":
        pot = max(0, vc * (PRAG_MARJA - marja))
    else:
        pot = vc * 0.15
    cmp(f"21!{cid} semnal", V("21_ANALIZA_CLIENTI", f"O{r}"), sem)
    cmp(f"21!{cid} oportunitate", V("21_ANALIZA_CLIENTI", f"P{r}"), OPORTUNITATE[sem])
    cmp(f"21!{cid} potențial", V("21_ANALIZA_CLIENTI", f"Q{r}"), pot)
    ua = ultima_achizitie(cid)
    cmp(f"21!{cid} ultima achiziție (serial)", V("21_ANALIZA_CLIENTI", f"K{r}"),
        (ua - dt.date(1899, 12, 30)).days if ua else "")

# ---------- semnalele din 22_ANALIZA_PRODUSE
print("--- 22: semnale ---")
for i, p_ in enumerate(PRODUSE):
    r = R0 + i
    pid = p_[0]
    vc, vp = venit(CUR, produs=pid), venit(CMP, produs=pid)
    pb = profit(CUR, produs=pid)
    marja = pb / vc if vc else ""
    g = delta(vc, vp)
    pen = clienti_produs(pid, CUR) / clienti_activi(CUR)
    if vc > 0 and marja != "" and marja < PRAG_MARJA:
        sem = "Marjă sub referință"
    elif g != "" and g >= PRAG_CRESTERE and pen < 0.3:
        sem = "Creștere + penetrare mică"
    elif g != "" and g <= PRAG_SCADERE:
        sem = "Produs în scădere"
    elif vc == 0:
        sem = "Fără vânzări"
    else:
        sem = "Stabil"
    cmp(f"22!{pid} semnal", V("22_ANALIZA_PRODUSE", f"M{r}"), sem)

# ---------- 23_ANALIZA_OAMENI
print("--- 23_ANALIZA_OAMENI ---")
for i, o in enumerate(OAMENI):
    r = R0 + i
    sid, ftev = o[0], o[4]
    vc, vp = venit(CUR, persoana=sid), venit(CMP, persoana=sid)
    pb = profit(CUR, persoana=sid)
    t = tinta(CUR, "Persoană", sid)
    cmp(f"23!{sid} venit curent", V("23_ANALIZA_OAMENI", f"F{r}"), vc)
    cmp(f"23!{sid} venit comparativ", V("23_ANALIZA_OAMENI", f"G{r}"), vp)
    cmp(f"23!{sid} delta %", V("23_ANALIZA_OAMENI", f"H{r}"), delta(vc, vp))
    cmp(f"23!{sid} profit brut", V("23_ANALIZA_OAMENI", f"I{r}"), pb)
    cmp(f"23!{sid} marjă", V("23_ANALIZA_OAMENI", f"J{r}"), pb / vc if vc else "")
    cmp(f"23!{sid} clienți activi", V("23_ANALIZA_OAMENI", f"K{r}"), clienti_persoana(sid, CUR))
    cmp(f"23!{sid} venit / FTE", V("23_ANALIZA_OAMENI", f"L{r}"), vc / ftev)
    cmp(f"23!{sid} țintă", V("23_ANALIZA_OAMENI", f"M{r}"), t)
    cmp(f"23!{sid} atingere țintă", V("23_ANALIZA_OAMENI", f"N{r}"), vc / t if t else "")

# ---------- 41_GRAFICE
print("--- 41_GRAFICE (tabel lunar, clasamente, NPS) ---")
for k in range(1, 13):
    rr = 60 + k
    an, luna = 2026, 6 - 12 + k
    while luna <= 0:
        an, luna = an - 1, luna + 12
    inceput = dt.date(an, luna, 1)
    sfarsit = dt.date(an + (luna == 12), luna % 12 + 1, 1) - dt.timedelta(days=1)
    f = (inceput, sfarsit)
    cmp(f"41!lună {an}-{luna:02d} venit", V("41_GRAFICE", f"C{rr}"), venit(f))
    cmp(f"41!lună {an}-{luna:02d} profit", V("41_GRAFICE", f"D{rr}"), profit(f))

clasament_cl = sorted(((venit(CUR, client=c[0]), c[1]) for c in CLIENTI
                       if venit(CUR, client=c[0]) > 0), reverse=True)
for k, (v, nume) in enumerate(clasament_cl[:5], start=1):
    cmp(f"41!top client {k} nume", V("41_GRAFICE", f"H{60+k}"), nume)
    cmp(f"41!top client {k} venit", V("41_GRAFICE", f"I{60+k}"), v)

clasament_pr = sorted(((venit(CUR, produs=p[0]), p[1]) for p in PRODUSE
                       if venit(CUR, produs=p[0]) > 0), reverse=True)
for k, (v, nume) in enumerate(clasament_pr[:3], start=1):
    cmp(f"41!top produs {k} nume", V("41_GRAFICE", f"K{60+k}"), nume)
    cmp(f"41!top produs {k} venit", V("41_GRAFICE", f"L{60+k}"), v)

note_cur = nps_set(CUR)
cmp("41!promotori", V("41_GRAFICE", "O61"), sum(1 for s in note_cur if s >= 9))
cmp("41!pasivi", V("41_GRAFICE", "O62"), sum(1 for s in note_cur if 7 <= s <= 8))
cmp("41!detractori", V("41_GRAFICE", "O63"), sum(1 for s in note_cur if s <= 6))

# ---------- 40_RAPORT (oglindirea indicatorilor)
print("--- 40_RAPORT ---")
rap = wb["40_RAPORT"]
for r in range(1, rap.max_row + 1):
    et = rap.cell(row=r, column=1).value
    if et in ("Venit net", "Profit brut", "Clienți activi", "Venit / FTE", "Gap față de țintă"):
        cmp(f"40_RAPORT!{et}", V("40_RAPORT", f"B{r}"), ASTEPT_KPI[et][0])

# ------------------------------------------------- raport final
gresite = [x for x in rezultate if not x[0]]
print(f"\n{'='*70}")
print(f"VERIFICATE: {len(rezultate)} valori calculate")
print(f"CORECTE   : {len(rezultate) - len(gresite)}")
print(f"GREȘITE   : {len(gresite)}")
if gresite:
    print("\nDiferențe:")
    for _ok, eticheta, excel, astept in gresite:
        print(f"  {eticheta}\n      Excel={excel!r}  așteptat={astept!r}")
sys.exit(1 if gresite else 0)
