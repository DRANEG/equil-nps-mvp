# -*- coding: utf-8 -*-
"""Stilul comun al fisierelor EQUIL: culori, formate, antete si helperi de formule.

Folosit de build_equil_workbook.py (versiunea completa) si de
build_equil_light.py (versiunea pentru firme mici), ca sa arate la fel.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- culori
INK        = "1F3A3D"
TEAL       = "12706E"
TEAL_LIGHT = "D7EDEC"
SAND       = "FDF6E3"
INPUT_FILL = "FFF8DC"   # galben: se completeaza
CALC_FILL  = "EEF1F3"   # gri: calculat automat
WHITE      = "FFFFFF"
GREY_TXT   = "5B6770"

# Paleta graficelor (validata pentru daltonism; vezi comentariul din 41_GRAFICE)
C_SER1, C_SER2, C_MUTED = "2A78D6", "EB6834", "898781"
C_ORD = ("1C5CAB", "2A78D6", "86B6EF")
C_GOOD, C_WARN, C_BAD = "0CA30C", "FAB219", "D03B3B"

# ---------------------------------------------------------------- formate
F_MONEY = '#,##0'
F_PCT   = '0.0%'
F_PCT2  = '0.00'
F_DATE  = 'yyyy-mm-dd'
F_NUM   = '#,##0.00'
F_INT   = '#,##0'

thin = Side(style="thin", color="C9D2D6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


# ---------------------------------------------------------------- antete
def style_header(ws, row=1, ncols=None, fill=TEAL, height=32):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def title_block(ws, title, subtitle=None, width_cols=6):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    ws.row_dimensions[1].height = 26
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=10, italic=True, color=GREY_TXT)
        ws.row_dimensions[2].height = 16
    ws.sheet_view.showGridLines = False


def label_value(ws, row, label, value=None, note=None, fmt=None, input_cell=True,
                span=1, note_col=3):
    """Un rand eticheta / valoare. span>1 imbina celulele de valoare, ca sa incapa
    text lung cand coloanele sunt inguste; note_col spune unde merge explicatia."""
    ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10, color=INK)
    umplere = PatternFill("solid", fgColor=INPUT_FILL if input_cell else CALC_FILL)
    for j in range(2, 2 + span):
        cc = ws.cell(row=row, column=j)
        cc.border = BORDER
        cc.fill = umplere
    if span > 1:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=1 + span)
    c = ws.cell(row=row, column=2, value=value)
    c.font = Font(size=10, color="000000" if input_cell else GREY_TXT)
    c.alignment = Alignment(vertical="center", wrap_text=span > 1)
    if fmt:
        c.number_format = fmt
    if note:
        n = ws.cell(row=row, column=note_col, value=note)
        n.font = Font(size=9, italic=True, color=GREY_TXT)
        n.alignment = Alignment(wrap_text=True, vertical="center")
    return c


def add_dv(ws, formula_range, cells, allow_blank=True):
    """Lista derulanta care citeste valorile dintr-un interval.

    Formula se scrie FARA semnul egal: in XML-ul unui fisier Excel, formula1
    a unei validari nu incepe cu "=". Cu el, Excel poate cere repararea fisierului.
    """
    dv = DataValidation(type="list", formula1=formula_range.lstrip("="), allow_blank=allow_blank)
    dv.error = "Valoare în afara listei permise. Alege din listă sau completează lista de valori."
    dv.errorTitle = "Valoare neacceptată"
    ws.add_data_validation(dv)
    dv.add(cells)
    return dv


# ------------------------------------------- helperi de formule, fara TEXT()
# Codurile de format din TEXT() depind de limba interfetei Excel
# ("yyyy-mm" nu functioneaza intr-un Excel romanesc), deci nu le folosim nicaieri.

def ym(d):
    """AAAA-LL."""
    return f'YEAR({d})&"-"&RIGHT("0"&MONTH({d}),2)'


def ymd(d):
    """AAAA-LL-ZZ."""
    return f'{ym(d)}&"-"&RIGHT("0"&DAY({d}),2)'


def mon(v, dec=0):
    """Numar cu separator de mii, conform setarilor locale."""
    return f'FIXED({v},{dec})'


def pct(v, dec=1):
    """Procent scris ca text."""
    return f'FIXED(({v})*100,{dec})&"%"'


def verifica(wb):
    """Verificare la construire: paranteze si ghilimele echilibrate in fiecare formula.
    Ridica exceptie inainte de salvare, ca sa nu ajunga un fisier stricat la client."""
    erori = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    if v.count("(") != v.count(")"):
                        erori.append(f"{ws.title}!{c.coordinate}: paranteze "
                                     f"({v.count('(')} deschise, {v.count(')')} inchise)\n    {v[:160]}")
                    elif v.count('"') % 2:
                        erori.append(f"{ws.title}!{c.coordinate}: ghilimele impare\n    {v[:160]}")
    if erori:
        unice = sorted({e.split(":", 1)[1] for e in erori})
        raise AssertionError(f"{len(erori)} formule stricate ({len(unice)} tipare distincte):\n  "
                             + "\n  ".join(erori[:3]))
