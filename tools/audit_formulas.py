# -*- coding: utf-8 -*-
"""Audit static al formulelor din workbook: paranteze, ghilimele, foi inexistente,
functii care nu exista in Excel 2016 si referinte de interval invalide.

Ruleaza: python3 tools/audit_formulas.py [cale.xlsx]
"""
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

PATH = sys.argv[1] if len(sys.argv) > 1 else "dist/EQUIL_Growth_Intelligence_v0_2.xlsx"

# Functii introduse dupa Excel 2016 sau dependente de limba interfetei.
INTERZISE = {"IFS", "MAXIFS", "MINIFS", "TEXTJOIN", "XLOOKUP", "UNIQUE", "FILTER",
             "LET", "SORT", "SEQUENCE", "CONCAT", "SWITCH", "LAMBDA",
             "TEXT"}  # TEXT: codurile de format depind de limba Excel-ului
PERMISE = {"IF", "IFERROR", "SUM", "SUMIF", "SUMIFS", "COUNT", "COUNTA", "COUNTIF",
           "COUNTIFS", "AVERAGE", "AVERAGEIFS", "SUMPRODUCT", "MAX", "MIN", "INDEX",
           "MATCH", "LARGE", "SMALL", "ABS", "N", "TODAY", "OR", "AND", "NOT",
           "ROUND", "FIXED", "YEAR", "MONTH", "DAY", "DATE", "RIGHT", "LEFT", "ROW"}

wb = openpyxl.load_workbook(PATH)
sheets = set(wb.sheetnames)
probleme, nformule, functii = [], 0, {}

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            nformule += 1
            loc = f"{ws.title}!{c.coordinate}"
            if v.count("(") != v.count(")"):
                probleme.append(f"PARANTEZE          {loc}: {v[:110]}")
            if v.count('"') % 2:
                probleme.append(f"GHILIMELE          {loc}: {v[:110]}")
            if "{r}" in v or "{s}" in v:
                probleme.append(f"PLACEHOLDER        {loc}: {v[:110]}")
            for fn in re.findall(r"([A-Z_][A-Z0-9_.]*)\s*\(", v):
                functii[fn] = functii.get(fn, 0) + 1
                if fn in INTERZISE:
                    probleme.append(f"FUNCTIE INTERZISA  {fn} in {loc}")
                elif fn not in PERMISE:
                    probleme.append(f"FUNCTIE NECUNOSCUTA {fn} in {loc}")
            for sh in re.findall(r"'([^']+)'!", v):
                if sh not in sheets:
                    probleme.append(f"FOAIE INEXISTENTA  '{sh}' in {loc}")
            for sh, col, _r1, _r2 in re.findall(
                    r"'([^']+)'!\$([A-Z]{1,2})\$(\d+):\$[A-Z]{1,2}\$(\d+)", v):
                if sh in sheets and column_index_from_string(col) > wb[sh].max_column:
                    probleme.append(f"COLOANA LIPSA      {sh}!{col} referita din {loc}")

print(f"Fisier: {PATH}")
print(f"Formule: {nformule}")
print("Functii:", ", ".join(f"{k}({v})" for k, v in sorted(functii.items())))
print(f"Probleme: {len(probleme)}")
for x in probleme[:30]:
    print("  -", x)
sys.exit(1 if probleme else 0)
